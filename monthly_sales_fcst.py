"""Transactional monthly sales / FCST management for Global MAPS."""

import io
import json
import re
import sqlite3
import uuid
from calendar import monthrange
from collections import defaultdict
from datetime import date, datetime, timedelta, timezone
from zoneinfo import ZoneInfo

from flask import g, jsonify, request, send_file
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from google_sheets_sync import GoogleSheetsMirror, SHEET_SCHEMAS, sheets_status
from maps_taxonomy import PARENT_BUSINESS_AREAS
from commercial_context import (
    apply_monthly_sale_context,
    inherit_monthly_sale_context,
    resolve_customer_order_context,
)


SEOUL = ZoneInfo("Asia/Seoul")
BUSINESS_UNITS = dict(PARENT_BUSINESS_AREAS)
SALES_STATUSES = {"confirmed": "확정", "scheduled": "예정", "pipeline": "추진", "undecided": "미정"}
TIMING_TYPES = {"current_new": "당월 추진", "previous_carryover": "전월 이월"}
CUSTOMER_HISTORIES = {"new": "신규 거래처", "existing": "기존 거래처"}
CONFIRMATION_BASES = {"payment_complete", "lc_open", "approved_credit", "other"}
PLAN_CURRENCIES = ("USD", "EUR", "JPY", "CNH", "KRW")
# Existing SQLite databases have a legacy CHECK constraint that only accepts CNY
# in monthly_sales.transaction_currency.  New operational records use CNH in the
# additive transaction_currency_standard column while the legacy value is kept
# untouched for backwards compatibility and auditability.
LEGACY_CURRENCY_ALIASES = {"CNY": "CNH"}
CURRENCY_UNITS = {"USD": 1, "EUR": 1, "JPY": 1, "CNH": 1, "KRW": 1}
MONTHLY_SALES_OWNERS = ("박정현", "김경태", "장윤선", "최령", "이인경", "김예원")
JPY_RATE_UNIT_ACTION_KEY = "monthly-sales-jpy-rate-unit:1-jpy-v1"
ROUND_TIMING_RULE_ACTION_KEY = "monthly-sales-timing-from-original-month-v1"
ROUND_DEFINITIONS = (
    ("initial", "최초 FCST", "initial_cutoff_date"),
    ("round1", "1차", "round1_cutoff_date"),
    ("round2", "2차", "round2_cutoff_date"),
    ("round3", "3차", "round3_cutoff_date"),
    ("final", "최종마감", "final_cutoff_date"),
)
MILESTONE_FIELDS = (
    "order_agreed_at", "po_received_at", "pi_no", "pi_sent_at", "pi_confirmed_at",
    "payment_expected_at", "payment_actual_at", "shipment_expected_at", "shipment_actual_at",
    "shipping_expected_at", "shipping_actual_at", "exception_reason",
)
DATE_FIELDS = tuple(field for field in MILESTONE_FIELDS if field.endswith("_at"))


SCHEMA_SQL = """
CREATE TABLE IF NOT EXISTS monthly_sales (
  id TEXT PRIMARY KEY,
  sales_no TEXT NOT NULL UNIQUE,
  customer_name TEXT NOT NULL,
  country TEXT NOT NULL,
  owner_name TEXT NOT NULL,
  business_unit TEXT NOT NULL CHECK(business_unit IN ('aesthetic','medical','dental')),
  currency TEXT NOT NULL DEFAULT 'USD' CHECK(currency = 'USD'),
  transaction_currency TEXT NOT NULL DEFAULT 'USD' CHECK(transaction_currency IN ('USD','EUR','JPY','CNY','KRW')),
  transaction_currency_standard TEXT NOT NULL DEFAULT 'USD'
    CHECK(transaction_currency_standard IN ('USD','EUR','JPY','CNH','KRW')),
  transaction_amount REAL NOT NULL DEFAULT 0 CHECK(transaction_amount >= 0),
  customer_history TEXT NOT NULL DEFAULT 'existing' CHECK(customer_history IN ('new','existing')),
  original_month TEXT NOT NULL,
  target_month TEXT NOT NULL,
  timing_type TEXT NOT NULL DEFAULT 'current_new' CHECK(timing_type IN ('current_new','previous_carryover')),
  amount_usd REAL NOT NULL CHECK(amount_usd >= 0),
  plan_rate REAL NOT NULL DEFAULT 0,
  plan_usd_rate REAL NOT NULL DEFAULT 0,
  actual_rate REAL,
  applied_rate REAL NOT NULL DEFAULT 0,
  rate_type TEXT NOT NULL DEFAULT 'plan' CHECK(rate_type IN ('plan','shipment','pending')),
  rate_base_date TEXT,
  rate_source TEXT,
  krw_amount REAL NOT NULL DEFAULT 0,
  sales_status TEXT NOT NULL CHECK(sales_status IN ('confirmed','scheduled','pipeline','undecided')),
  confirmation_basis TEXT NOT NULL DEFAULT '',
  confirmation_note TEXT NOT NULL DEFAULT '',
  preliminary_actual_currency TEXT,
  preliminary_actual_amount REAL,
  preliminary_actual_krw_amount REAL,
  preliminary_reference_no TEXT NOT NULL DEFAULT '',
  preliminary_note TEXT NOT NULL DEFAULT '',
  record_status TEXT NOT NULL DEFAULT 'active' CHECK(record_status IN ('active','cancelled','carried_over')),
  split_role TEXT NOT NULL DEFAULT 'none' CHECK(split_role IN ('none','original','split','carryover_split')),
  carryover_role TEXT NOT NULL DEFAULT 'none' CHECK(carryover_role IN ('none','inflow','outflow')),
  carryover_decision TEXT NOT NULL DEFAULT 'no' CHECK(carryover_decision IN ('yes','no')),
  carryover_decided_at TEXT,
  carryover_target_month TEXT,
  carryover_decision_reason TEXT NOT NULL DEFAULT '',
  original_sales_id TEXT,
  split_group_id TEXT,
  notes TEXT NOT NULL DEFAULT '',
  cancel_reason TEXT,
  version INTEGER NOT NULL DEFAULT 1,
  created_by INTEGER NOT NULL,
  updated_by INTEGER NOT NULL,
  owner_user_id INTEGER,
  created_at TEXT NOT NULL,
  updated_at TEXT NOT NULL,
  cancelled_at TEXT,
  FOREIGN KEY(created_by) REFERENCES users(id),
  FOREIGN KEY(updated_by) REFERENCES users(id),
  FOREIGN KEY(owner_user_id) REFERENCES users(id)
);
CREATE INDEX IF NOT EXISTS idx_monthly_sales_month ON monthly_sales(target_month, record_status, sales_status);
CREATE INDEX IF NOT EXISTS idx_monthly_sales_owner ON monthly_sales(owner_name, target_month);
CREATE INDEX IF NOT EXISTS idx_monthly_sales_original ON monthly_sales(original_sales_id, split_group_id);

CREATE TABLE IF NOT EXISTS monthly_sales_customer_aliases (
  id TEXT PRIMARY KEY,
  customer_master_id TEXT NOT NULL,
  alias_name TEXT NOT NULL,
  alias_country TEXT NOT NULL,
  normalized_name TEXT NOT NULL,
  normalized_country TEXT NOT NULL,
  source TEXT NOT NULL DEFAULT 'manual' CHECK(source IN ('auto','manual','temp')),
  created_by INTEGER,
  created_at TEXT NOT NULL,
  UNIQUE(normalized_name, normalized_country),
  FOREIGN KEY(customer_master_id) REFERENCES customer_master(id),
  FOREIGN KEY(created_by) REFERENCES users(id)
);
CREATE INDEX IF NOT EXISTS idx_monthly_sales_customer_alias_master
  ON monthly_sales_customer_aliases(customer_master_id);

CREATE TABLE IF NOT EXISTS monthly_sales_milestones (
  sales_id TEXT PRIMARY KEY,
  order_agreed_at TEXT,
  po_received_at TEXT,
  pi_no TEXT,
  pi_sent_at TEXT,
  pi_confirmed_at TEXT,
  payment_expected_at TEXT,
  payment_actual_at TEXT,
  shipment_expected_at TEXT,
  shipment_actual_at TEXT,
  shipping_expected_at TEXT,
  shipping_actual_at TEXT,
  exception_reason TEXT NOT NULL DEFAULT '',
  updated_at TEXT NOT NULL,
  FOREIGN KEY(sales_id) REFERENCES monthly_sales(id) ON DELETE CASCADE
);

CREATE TABLE IF NOT EXISTS monthly_sales_split_links (
  id TEXT PRIMARY KEY,
  operation_key TEXT NOT NULL UNIQUE,
  split_group_id TEXT NOT NULL,
  original_sales_id TEXT NOT NULL,
  child_sales_id TEXT NOT NULL,
  amount_before REAL NOT NULL,
  proceed_amount REAL NOT NULL,
  split_amount REAL NOT NULL,
  split_date TEXT NOT NULL,
  reason TEXT NOT NULL,
  new_pi_no TEXT,
  original_version_after INTEGER NOT NULL,
  child_version_after INTEGER NOT NULL,
  status TEXT NOT NULL DEFAULT 'active' CHECK(status IN ('active','cancelled')),
  created_by INTEGER NOT NULL,
  created_at TEXT NOT NULL,
  cancelled_by INTEGER,
  cancelled_at TEXT,
  cancel_reason TEXT,
  FOREIGN KEY(original_sales_id) REFERENCES monthly_sales(id),
  FOREIGN KEY(child_sales_id) REFERENCES monthly_sales(id)
);
CREATE INDEX IF NOT EXISTS idx_monthly_split_group ON monthly_sales_split_links(split_group_id, status);

CREATE TABLE IF NOT EXISTS monthly_sales_carryovers (
  id TEXT PRIMARY KEY,
  operation_key TEXT NOT NULL UNIQUE,
  source_sales_id TEXT NOT NULL,
  target_sales_id TEXT NOT NULL,
  carryover_type TEXT NOT NULL CHECK(carryover_type IN ('partial','full')),
  source_month TEXT NOT NULL,
  target_month TEXT NOT NULL,
  amount_usd REAL NOT NULL,
  status_at_carryover TEXT NOT NULL,
  carryover_label TEXT NOT NULL,
  carryover_date TEXT NOT NULL,
  reason TEXT NOT NULL,
  source_version_after INTEGER NOT NULL,
  target_version_after INTEGER NOT NULL,
  status TEXT NOT NULL DEFAULT 'active' CHECK(status IN ('active','cancelled')),
  created_by INTEGER NOT NULL,
  created_at TEXT NOT NULL,
  cancelled_by INTEGER,
  cancelled_at TEXT,
  cancel_reason TEXT,
  FOREIGN KEY(source_sales_id) REFERENCES monthly_sales(id),
  FOREIGN KEY(target_sales_id) REFERENCES monthly_sales(id)
);
CREATE INDEX IF NOT EXISTS idx_monthly_carryover_source ON monthly_sales_carryovers(source_sales_id, status);

CREATE TABLE IF NOT EXISTS monthly_fcst_snapshots (
  id TEXT PRIMARY KEY,
  target_month TEXT NOT NULL,
  sales_id TEXT NOT NULL,
  sales_no TEXT NOT NULL,
  business_unit TEXT NOT NULL,
  sales_status TEXT NOT NULL,
  amount_usd REAL NOT NULL,
  krw_amount REAL NOT NULL,
  submitted_by INTEGER NOT NULL,
  submitted_at TEXT NOT NULL,
  UNIQUE(target_month, sales_id),
  FOREIGN KEY(sales_id) REFERENCES monthly_sales(id)
);
CREATE INDEX IF NOT EXISTS idx_monthly_fcst_snapshot_month ON monthly_fcst_snapshots(target_month, business_unit);

CREATE TABLE IF NOT EXISTS monthly_sales_history (
  id INTEGER PRIMARY KEY AUTOINCREMENT,
  sales_id TEXT NOT NULL,
  sales_no TEXT NOT NULL,
  event_type TEXT NOT NULL,
  reason TEXT NOT NULL DEFAULT '',
  actor_user_id INTEGER,
  actor_name TEXT NOT NULL,
  occurred_at TEXT NOT NULL,
  effective_at TEXT,
  effective_at_source TEXT NOT NULL DEFAULT 'system_default',
  change_set_json TEXT NOT NULL DEFAULT '{}',
  snapshot_json TEXT NOT NULL,
  FOREIGN KEY(sales_id) REFERENCES monthly_sales(id)
);
CREATE INDEX IF NOT EXISTS idx_monthly_sales_history_asof ON monthly_sales_history(sales_id, occurred_at DESC, id DESC);

CREATE TABLE IF NOT EXISTS monthly_sales_round_imports (
  id TEXT PRIMARY KEY,
  target_month TEXT NOT NULL,
  round_key TEXT NOT NULL CHECK(round_key IN ('initial','round1','round2','round3','final')),
  cutoff_date TEXT NOT NULL,
  source_file TEXT NOT NULL,
  row_count INTEGER NOT NULL CHECK(row_count > 0),
  imported_by INTEGER NOT NULL,
  imported_at TEXT NOT NULL,
  UNIQUE(target_month, round_key),
  FOREIGN KEY(imported_by) REFERENCES users(id)
);
CREATE INDEX IF NOT EXISTS idx_monthly_round_import_cutoff
  ON monthly_sales_round_imports(target_month, cutoff_date, round_key);

CREATE TABLE IF NOT EXISTS monthly_sales_round_rows (
  id TEXT PRIMARY KEY,
  import_id TEXT NOT NULL,
  source_row INTEGER NOT NULL,
  target_month TEXT NOT NULL,
  round_key TEXT NOT NULL,
  cutoff_date TEXT NOT NULL,
  linked_sales_id TEXT,
  sales_no TEXT NOT NULL,
  customer_name TEXT NOT NULL,
  country TEXT NOT NULL,
  owner_name TEXT NOT NULL,
  business_unit TEXT NOT NULL,
  customer_history TEXT NOT NULL,
  original_month TEXT NOT NULL,
  timing_type TEXT NOT NULL,
  transaction_currency TEXT NOT NULL,
  transaction_amount REAL NOT NULL,
  amount_usd REAL NOT NULL,
  krw_amount REAL NOT NULL,
  applied_rate REAL NOT NULL,
  sales_status TEXT NOT NULL,
  record_status TEXT NOT NULL DEFAULT 'active',
  carryover_decision TEXT NOT NULL DEFAULT 'no',
  carryover_target_month TEXT,
  notes TEXT NOT NULL DEFAULT '',
  snapshot_json TEXT NOT NULL,
  UNIQUE(import_id, source_row),
  FOREIGN KEY(import_id) REFERENCES monthly_sales_round_imports(id),
  FOREIGN KEY(linked_sales_id) REFERENCES monthly_sales(id)
);
CREATE INDEX IF NOT EXISTS idx_monthly_round_rows_lookup
  ON monthly_sales_round_rows(target_month, round_key, cutoff_date);
CREATE INDEX IF NOT EXISTS idx_monthly_round_rows_import
  ON monthly_sales_round_rows(import_id, source_row);

CREATE TABLE IF NOT EXISTS monthly_sales_round_import_revisions (
  id TEXT PRIMARY KEY,
  import_id TEXT NOT NULL,
  revision_no INTEGER NOT NULL CHECK(revision_no > 0),
  source_file TEXT NOT NULL,
  row_count INTEGER NOT NULL CHECK(row_count > 0),
  replaced_by INTEGER NOT NULL,
  replaced_at TEXT NOT NULL,
  replacement_reason TEXT NOT NULL,
  snapshot_json TEXT NOT NULL,
  UNIQUE(import_id, revision_no),
  FOREIGN KEY(import_id) REFERENCES monthly_sales_round_imports(id),
  FOREIGN KEY(replaced_by) REFERENCES users(id)
);
CREATE INDEX IF NOT EXISTS idx_monthly_round_import_revisions
  ON monthly_sales_round_import_revisions(import_id, revision_no DESC);

CREATE TABLE IF NOT EXISTS monthly_sales_rates (
  target_month TEXT NOT NULL,
  currency TEXT NOT NULL DEFAULT 'USD',
  plan_rate REAL NOT NULL,
  source TEXT NOT NULL DEFAULT '관리자 입력',
  correction_reason TEXT NOT NULL DEFAULT '',
  updated_by INTEGER NOT NULL,
  updated_at TEXT NOT NULL,
  PRIMARY KEY(target_month, currency)
);

CREATE TABLE IF NOT EXISTS monthly_sales_daily_rates (
  rate_base_date TEXT NOT NULL,
  currency TEXT NOT NULL DEFAULT 'USD',
  rate REAL NOT NULL CHECK(rate > 0),
  source_rate REAL,
  source_unit TEXT NOT NULL DEFAULT '1',
  normalized_unit TEXT NOT NULL DEFAULT 'KRW_PER_1',
  source TEXT NOT NULL,
  fetched_at TEXT,
  uploaded_file TEXT NOT NULL DEFAULT '',
  correction_reason TEXT NOT NULL DEFAULT '',
  updated_by INTEGER NOT NULL,
  updated_at TEXT NOT NULL,
  PRIMARY KEY(rate_base_date, currency)
);

CREATE TABLE IF NOT EXISTS monthly_sales_targets (
  target_month TEXT NOT NULL,
  business_unit TEXT NOT NULL CHECK(business_unit IN ('aesthetic','medical','dental')),
  target_usd REAL NOT NULL DEFAULT 0,
  target_krw REAL NOT NULL DEFAULT 0,
  updated_by INTEGER NOT NULL,
  updated_at TEXT NOT NULL,
  PRIMARY KEY(target_month, business_unit)
);

CREATE TABLE IF NOT EXISTS monthly_sales_master_candidates (
  customer_name TEXT NOT NULL,
  country TEXT NOT NULL,
  owner_name TEXT NOT NULL,
  business_unit TEXT NOT NULL,
  usage_count INTEGER NOT NULL DEFAULT 1,
  first_seen_at TEXT NOT NULL,
  last_seen_at TEXT NOT NULL,
  status TEXT NOT NULL DEFAULT 'candidate' CHECK(status IN ('candidate','confirmed','rejected')),
  PRIMARY KEY(customer_name, country, owner_name, business_unit)
);

CREATE TABLE IF NOT EXISTS monthly_sales_month_settings (
  target_month TEXT PRIMARY KEY,
  month_status TEXT NOT NULL DEFAULT 'open' CHECK(month_status IN ('open','closed')),
  initial_cutoff_date TEXT,
  round1_cutoff_date TEXT,
  round2_cutoff_date TEXT,
  round3_cutoff_date TEXT,
  final_cutoff_date TEXT,
  closed_by INTEGER,
  closed_at TEXT,
  reopened_by INTEGER,
  reopened_at TEXT,
  reopen_reason TEXT,
  updated_at TEXT NOT NULL
);

CREATE TABLE IF NOT EXISTS monthly_sales_sheet_sync (
  id INTEGER PRIMARY KEY CHECK(id = 1),
  dirty INTEGER NOT NULL DEFAULT 1,
  last_attempt_at TEXT,
  last_success_at TEXT,
  last_error TEXT,
  updated_at TEXT NOT NULL
);
"""


def init_monthly_sales_schema(db, now):
    db.executescript(SCHEMA_SQL)
    db.execute(
        "CREATE TABLE IF NOT EXISTS one_time_actions (action_key TEXT PRIMARY KEY, applied_at TEXT NOT NULL)"
    )
    sales_columns = {row["name"] for row in db.execute("PRAGMA table_info(monthly_sales)").fetchall()}
    migrations = {
        "carryover_decision": "TEXT NOT NULL DEFAULT 'no' CHECK(carryover_decision IN ('yes','no'))",
        "carryover_decided_at": "TEXT",
        "carryover_target_month": "TEXT",
        "carryover_decision_reason": "TEXT NOT NULL DEFAULT ''",
        "transaction_currency": "TEXT NOT NULL DEFAULT 'USD' CHECK(transaction_currency IN ('USD','EUR','JPY','CNY','KRW'))",
        "transaction_amount": "REAL NOT NULL DEFAULT 0 CHECK(transaction_amount >= 0)",
        "customer_history": "TEXT NOT NULL DEFAULT 'existing' CHECK(customer_history IN ('new','existing'))",
        "customer_master_id": "TEXT",
        "original_customer_name": "TEXT",
        "customer_link_status": "TEXT NOT NULL DEFAULT 'unlinked' CHECK(customer_link_status IN ('unlinked','review','auto_linked','linked','temp_created','excluded'))",
        "customer_linked_at": "TEXT",
        "customer_linked_by": "INTEGER",
        "transaction_currency_standard": "TEXT NOT NULL DEFAULT 'USD' CHECK(transaction_currency_standard IN ('USD','EUR','JPY','CNH','KRW'))",
        "plan_usd_rate": "REAL NOT NULL DEFAULT 0",
        "owner_user_id": "INTEGER",
        "confirmation_basis": "TEXT NOT NULL DEFAULT ''",
        "confirmation_note": "TEXT NOT NULL DEFAULT ''",
        "preliminary_actual_currency": "TEXT",
        "preliminary_actual_amount": "REAL",
        "preliminary_actual_krw_amount": "REAL",
        "preliminary_reference_no": "TEXT NOT NULL DEFAULT ''",
        "preliminary_note": "TEXT NOT NULL DEFAULT ''",
    }
    for column, definition in migrations.items():
        if column not in sales_columns:
            db.execute(f"ALTER TABLE monthly_sales ADD COLUMN {column} {definition}")
    db.execute(
        """UPDATE monthly_sales
           SET transaction_currency=COALESCE(NULLIF(transaction_currency, ''), 'USD'),
               transaction_currency_standard=CASE
                 WHEN COALESCE(NULLIF(transaction_currency_standard, ''), transaction_currency)='CNY' THEN 'CNH'
                 ELSE COALESCE(NULLIF(transaction_currency_standard, ''), transaction_currency, 'USD') END,
               transaction_amount=CASE WHEN transaction_amount IS NULL OR transaction_amount=0
                                       THEN amount_usd ELSE transaction_amount END,
               customer_history=COALESCE(NULLIF(customer_history, ''), 'existing')"""
    )
    # plan_rate is the immutable selected-currency snapshot.  plan_usd_rate is
    # the companion USD/KRW context required to reproduce non-USD conversions.
    # Existing production FCST rows are USD, so this backfill is exact; the
    # fallback also preserves future legacy rows without changing plan_rate.
    db.execute(
        """UPDATE monthly_sales
           SET plan_usd_rate=CASE
             WHEN transaction_currency_standard='USD' AND plan_rate>0 THEN plan_rate
             WHEN amount_usd>0 AND krw_amount>0 THEN krw_amount/amount_usd
             ELSE COALESCE((SELECT r.plan_rate FROM monthly_sales_rates r
                            WHERE r.target_month=monthly_sales.target_month AND r.currency='USD'),0)
           END
           WHERE COALESCE(plan_usd_rate,0)<=0"""
    )
    user_table_exists = db.execute(
        "SELECT 1 FROM sqlite_master WHERE type='table' AND name='users'"
    ).fetchone()
    if user_table_exists:
        db.execute(
            """UPDATE monthly_sales
               SET owner_user_id=(
                 SELECT MIN(u.id) FROM users u
                 WHERE u.display_name=monthly_sales.owner_name
                   AND u.status='active' AND u.deleted_at IS NULL
                 HAVING COUNT(*)=1
               )
               WHERE owner_user_id IS NULL"""
        )
    db.execute(
        "CREATE INDEX IF NOT EXISTS idx_monthly_sales_owner_user "
        "ON monthly_sales(owner_user_id, target_month)"
    )
    # Preserve the historical CNY row while seeding the new CNH operating key.
    db.execute(
        """INSERT OR IGNORE INTO monthly_sales_rates
           (target_month,currency,plan_rate,source,correction_reason,updated_by,updated_at)
           SELECT target_month,'CNH',plan_rate,source,
                  CASE WHEN correction_reason='' THEN 'CNY 운영키에서 CNH 운영키로 복사' ELSE correction_reason END,
                  updated_by,updated_at
           FROM monthly_sales_rates WHERE currency='CNY'"""
    )
    history_columns = {
        row["name"] for row in db.execute("PRAGMA table_info(monthly_sales_history)").fetchall()
    }
    for column, definition in {
        "effective_at": "TEXT",
        "effective_at_source": "TEXT NOT NULL DEFAULT 'system_default'",
        "change_set_json": "TEXT NOT NULL DEFAULT '{}'",
    }.items():
        if column not in history_columns:
            db.execute(f"ALTER TABLE monthly_sales_history ADD COLUMN {column} {definition}")
    for history in db.execute(
        "SELECT id,occurred_at FROM monthly_sales_history WHERE effective_at IS NULL OR effective_at=''"
    ).fetchall():
        try:
            recorded = datetime.fromisoformat(str(history["occurred_at"]).replace("Z", "+00:00"))
            if recorded.tzinfo is None:
                recorded = recorded.replace(tzinfo=timezone.utc)
            effective_at = recorded.astimezone(SEOUL).date().isoformat()
        except (TypeError, ValueError):
            effective_at = str(history["occurred_at"] or "")[:10]
        db.execute(
            """UPDATE monthly_sales_history
               SET effective_at=?,effective_at_source='legacy_recorded_at',
                   change_set_json=COALESCE(NULLIF(change_set_json,''),'{}')
               WHERE id=?""",
            (effective_at, history["id"]),
        )
    db.execute(
        "CREATE INDEX IF NOT EXISTS idx_monthly_sales_history_effective "
        "ON monthly_sales_history(sales_id,effective_at,id)"
    )
    daily_rate_columns = {
        row["name"] for row in db.execute("PRAGMA table_info(monthly_sales_daily_rates)").fetchall()
    }
    for column, definition in {
        "source_rate": "REAL",
        "source_unit": "TEXT NOT NULL DEFAULT '1'",
        "normalized_unit": "TEXT NOT NULL DEFAULT 'KRW_PER_1'",
        "fetched_at": "TEXT",
    }.items():
        if column not in daily_rate_columns:
            db.execute(f"ALTER TABLE monthly_sales_daily_rates ADD COLUMN {column} {definition}")
    setting_columns = {
        row["name"] for row in db.execute("PRAGMA table_info(monthly_sales_month_settings)").fetchall()
    }
    for column in (
        "initial_cutoff_date", "round1_cutoff_date", "round2_cutoff_date",
        "round3_cutoff_date", "final_cutoff_date",
    ):
        if column not in setting_columns:
            db.execute(f"ALTER TABLE monthly_sales_month_settings ADD COLUMN {column} TEXT")
    db.execute(
        """
        UPDATE monthly_sales
        SET carryover_decision='yes',
            carryover_decided_at=COALESCE(carryover_decided_at, (
              SELECT c.carryover_date FROM monthly_sales_carryovers c
              WHERE c.source_sales_id=monthly_sales.id AND c.status='active'
              ORDER BY c.created_at DESC, c.rowid DESC LIMIT 1
            )),
            carryover_target_month=COALESCE(carryover_target_month, (
              SELECT c.target_month FROM monthly_sales_carryovers c
              WHERE c.source_sales_id=monthly_sales.id AND c.status='active'
              ORDER BY c.created_at DESC, c.rowid DESC LIMIT 1
            )),
            carryover_decision_reason=CASE WHEN carryover_decision_reason='' THEN COALESCE((
              SELECT c.reason FROM monthly_sales_carryovers c
              WHERE c.source_sales_id=monthly_sales.id AND c.status='active'
              ORDER BY c.created_at DESC, c.rowid DESC LIMIT 1
            ), '') ELSE carryover_decision_reason END
        WHERE EXISTS (
          SELECT 1 FROM monthly_sales_carryovers c
          WHERE c.source_sales_id=monthly_sales.id AND c.status='active'
        )
        """
    )
    db.execute(
        "INSERT OR IGNORE INTO monthly_sales_sheet_sync (id, dirty, updated_at) VALUES (1, 1, ?)",
        (now(),),
    )
    db.execute(
        "CREATE INDEX IF NOT EXISTS idx_monthly_sales_customer_master "
        "ON monthly_sales(customer_master_id, target_month)"
    )
    db.execute(
        "UPDATE monthly_sales SET original_customer_name=customer_name "
        "WHERE original_customer_name IS NULL OR original_customer_name=''"
    )
    if not db.execute(
        "SELECT 1 FROM one_time_actions WHERE action_key = ?", (JPY_RATE_UNIT_ACTION_KEY,)
    ).fetchone():
        db.execute("UPDATE monthly_sales_rates SET plan_rate=plan_rate/100.0 WHERE currency='JPY'")
        db.execute("UPDATE monthly_sales_daily_rates SET rate=rate/100.0 WHERE currency='JPY'")
        db.execute(
            """UPDATE monthly_sales
               SET plan_rate=plan_rate/100.0,
                   actual_rate=CASE WHEN actual_rate IS NULL THEN NULL ELSE actual_rate/100.0 END,
                   applied_rate=applied_rate/100.0
               WHERE transaction_currency='JPY'"""
        )
        history_rows = db.execute(
            "SELECT id, snapshot_json FROM monthly_sales_history"
        ).fetchall()
        for history_row in history_rows:
            try:
                snapshot = json.loads(history_row["snapshot_json"])
            except (TypeError, json.JSONDecodeError):
                continue
            if snapshot.get("transaction_currency") != "JPY":
                continue
            for field in ("plan_rate", "actual_rate", "applied_rate"):
                if snapshot.get(field) not in (None, ""):
                    snapshot[field] = round(float(snapshot[field]) / 100.0, 6)
            db.execute(
                "UPDATE monthly_sales_history SET snapshot_json=? WHERE id=?",
                (json.dumps(snapshot, ensure_ascii=False, separators=(",", ":"), default=str), history_row["id"]),
            )
        db.execute(
            "INSERT INTO one_time_actions (action_key, applied_at) VALUES (?, ?)",
            (JPY_RATE_UNIT_ACTION_KEY, now()),
        )
    reconcile_existing_timing_types(db, now)


def valid_month(value):
    return bool(re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", str(value or "")))


def excel_month_value(value):
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m")
    text = str(value or "").strip()
    full_date = re.fullmatch(r"(\d{4})-(0[1-9]|1[0-2])-\d{2}(?:[ T].*)?", text)
    return f"{full_date.group(1)}-{full_date.group(2)}" if full_date else text


def timing_type_from_months(original_month, target_month):
    """Return the business-defined timing when both months can be compared."""
    if not valid_month(original_month) or not valid_month(target_month):
        return None
    if original_month < target_month:
        return "previous_carryover"
    if original_month == target_month:
        return "current_new"
    return None


def inferred_timing_type(value, original_month, target_month):
    """Derive timing from business months and use input only when they cannot decide it."""
    aliases = {
        "current_new": "current_new", "당월 추진": "current_new", "당월 신규": "current_new",
        "previous_carryover": "previous_carryover", "전월 이월": "previous_carryover",
        "이월 유입": "previous_carryover",
    }
    derived = timing_type_from_months(original_month, target_month)
    if derived:
        return derived
    raw = str(value or "").strip()
    if raw:
        return aliases.get(raw, raw)
    return "current_new"


def reconcile_existing_timing_types(db, now):
    """One-time correction for legacy rows whose timing conflicts with their two months."""
    if db.execute(
        "SELECT 1 FROM one_time_actions WHERE action_key = ?",
        (ROUND_TIMING_RULE_ACTION_KEY,),
    ).fetchone():
        return

    timestamp = now()
    system_actor = {"id": None, "display_name": "시스템"}
    live_changed = False
    for sale in db.execute(
        "SELECT id,original_month,target_month,timing_type FROM monthly_sales"
    ).fetchall():
        desired = timing_type_from_months(sale["original_month"], sale["target_month"])
        if not desired or desired == sale["timing_type"]:
            continue
        db.execute(
            """UPDATE monthly_sales
               SET timing_type=?,version=version+1,updated_at=? WHERE id=?""",
            (desired, timestamp, sale["id"]),
        )
        add_history(
            db, sale["id"], "TIMING_RECLASSIFY",
            "최초 추진월 기준 진행 시점 자동 정정", system_actor, lambda: timestamp,
        )
        live_changed = True

    imports = db.execute(
        "SELECT * FROM monthly_sales_round_imports ORDER BY imported_at,id"
    ).fetchall()
    for imported in imports:
        stored_rows = db.execute(
            "SELECT * FROM monthly_sales_round_rows WHERE import_id=? ORDER BY source_row",
            (imported["id"],),
        ).fetchall()
        corrections = []
        for stored in stored_rows:
            desired = timing_type_from_months(stored["original_month"], stored["target_month"])
            if desired and desired != stored["timing_type"]:
                corrections.append((stored, desired))
        if not corrections:
            continue

        revision_no = db.execute(
            "SELECT COALESCE(MAX(revision_no), 0) + 1 FROM monthly_sales_round_import_revisions WHERE import_id=?",
            (imported["id"],),
        ).fetchone()[0]
        archive = {
            "import": row_dict(imported),
            "rows": [row_dict(row) for row in stored_rows],
        }
        db.execute(
            """INSERT INTO monthly_sales_round_import_revisions
               (id,import_id,revision_no,source_file,row_count,replaced_by,replaced_at,
                replacement_reason,snapshot_json)
               VALUES (?,?,?,?,?,?,?,?,?)""",
            (
                uuid.uuid4().hex, imported["id"], revision_no, imported["source_file"],
                imported["row_count"], imported["imported_by"], timestamp,
                "시스템 정정: 최초 추진월 기준 진행 시점 재분류",
                json.dumps(archive, ensure_ascii=False, separators=(",", ":"), default=str),
            ),
        )
        for stored, desired in corrections:
            try:
                snapshot = json.loads(stored["snapshot_json"])
            except (TypeError, json.JSONDecodeError):
                snapshot = {}
            snapshot["timing_type"] = desired
            db.execute(
                """UPDATE monthly_sales_round_rows
                   SET timing_type=?,snapshot_json=? WHERE id=?""",
                (
                    desired,
                    json.dumps(snapshot, ensure_ascii=False, separators=(",", ":"), default=str),
                    stored["id"],
                ),
            )

    if live_changed:
        db.execute(
            "UPDATE monthly_sales_sheet_sync SET dirty=1,updated_at=? WHERE id=1",
            (timestamp,),
        )
    db.execute(
        "INSERT INTO one_time_actions (action_key, applied_at) VALUES (?, ?)",
        (ROUND_TIMING_RULE_ACTION_KEY, timestamp),
    )


def valid_date(value):
    if value in (None, ""):
        return None
    text = str(value).strip()
    try:
        datetime.strptime(text, "%Y-%m-%d")
    except ValueError as exc:
        raise ValueError(f"날짜 형식이 올바르지 않습니다: {text}") from exc
    return text


def number(value, field, minimum=0, required=False):
    if value in (None, ""):
        if required:
            raise ValueError(f"{field}을(를) 입력하세요.")
        return 0.0
    try:
        parsed = float(str(value).replace(",", ""))
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{field}은(는) 숫자로 입력하세요.") from exc
    if parsed < minimum:
        raise ValueError(f"{field}은(는) {minimum:g} 이상이어야 합니다.")
    return round(parsed, 4)


def clean_text(value, field, maximum=300, required=False):
    text = str(value or "").strip()
    if required and not text:
        raise ValueError(f"{field}을(를) 입력하세요.")
    if len(text) > maximum:
        raise ValueError(f"{field}은(는) {maximum}자 이내로 입력하세요.")
    return text


def month_shift(value, offset):
    year, month = map(int, value.split("-"))
    month_index = year * 12 + month - 1 + offset
    return f"{month_index // 12:04d}-{month_index % 12 + 1:02d}"


def cutoff_utc(date_text):
    local = datetime.strptime(date_text, "%Y-%m-%d").replace(
        hour=23, minute=59, second=59, tzinfo=SEOUL
    )
    return local.astimezone(timezone.utc).replace(microsecond=0).isoformat()


def row_dict(row):
    return dict(row) if row is not None else None


def customer_link_key(value):
    return re.sub(r"\s+", " ", str(value or "").strip()).casefold()


def customer_link_group_key(customer_name, country):
    return uuid.uuid5(
        uuid.NAMESPACE_URL,
        f"medpark-monthly-fcst:{customer_link_key(customer_name)}:{customer_link_key(country)}",
    ).hex


def customer_master_candidates(db, customer_name, country):
    """Return strict unique matches and looser review candidates without guessing."""
    name_key, country_key = customer_link_key(customer_name), customer_link_key(country)
    alias = db.execute(
        """SELECT a.customer_master_id,c.customer_id,c.display_name,c.erp_original_name,
                  c.headquarters_country,c.erp_country_name,c.lifecycle_status
           FROM monthly_sales_customer_aliases a
           JOIN customer_master c ON c.id=a.customer_master_id
           WHERE a.normalized_name=? AND a.normalized_country=?""",
        (name_key, country_key),
    ).fetchone()
    if alias:
        return [row_dict(alias)], [row_dict(alias)]
    masters = db.execute(
        """SELECT id AS customer_master_id,customer_id,display_name,erp_original_name,
                  headquarters_country,erp_country_name,lifecycle_status
           FROM customer_master"""
    ).fetchall()
    strict, review = [], []
    loose_name = re.sub(r"[^0-9a-z가-힣]+", "", name_key)
    for row in masters:
        item = row_dict(row)
        names = [customer_link_key(item.get("display_name")), customer_link_key(item.get("erp_original_name"))]
        countries = [customer_link_key(item.get("headquarters_country")), customer_link_key(item.get("erp_country_name"))]
        name_exact = name_key and name_key in names
        country_exact = country_key and country_key in countries
        if name_exact and country_exact:
            strict.append(item)
            continue
        row_loose = {re.sub(r"[^0-9a-z가-힣]+", "", value) for value in names if value}
        if loose_name and loose_name in row_loose:
            review.append(item)
    return strict, strict + review


def resolve_sale_customer(values, db, existing=None):
    requested_id = str(values.get("customer_master_id") or "").strip()
    original_name = (
        existing["original_customer_name"]
        if existing and "original_customer_name" in existing.keys() and existing["original_customer_name"]
        else values["customer_name"]
    )
    if requested_id:
        master = db.execute(
            "SELECT id,display_name,lifecycle_status FROM customer_master WHERE id=?", (requested_id,)
        ).fetchone()
        if not master:
            raise ValueError("선택한 거래처 마스터를 찾을 수 없습니다.")
        if master["lifecycle_status"] != "active" and not (existing and existing["customer_master_id"] == requested_id):
            raise ValueError("거래중단 거래처에는 새 매출을 등록할 수 없습니다.")
        # customer_name/country are the sale-time snapshot.  A later master-name
        # edit must not silently rewrite historical FCST merely because notes or
        # status were updated.
        snapshot_name = (
            existing["customer_name"]
            if existing and existing["customer_master_id"] == requested_id else master["display_name"]
        )
        return requested_id, snapshot_name, original_name, "linked"
    strict, _review = customer_master_candidates(db, values["customer_name"], values["country"])
    if len(strict) == 1:
        master = strict[0]
        return master["customer_master_id"], master["display_name"], original_name, "auto_linked"
    existing_id = existing["customer_master_id"] if existing and "customer_master_id" in existing.keys() else None
    existing_status = existing["customer_link_status"] if existing and "customer_link_status" in existing.keys() else "unlinked"
    return existing_id, values["customer_name"], original_name, existing_status if existing_id else ("review" if len(strict) > 1 else "unlinked")


def apply_sale_customer_link(db, sale_id, values, actor_id, timestamp):
    db.execute(
        """UPDATE monthly_sales
           SET customer_master_id=?,customer_name=?,original_customer_name=?,customer_link_status=?,
               customer_linked_at=CASE WHEN ? IS NULL THEN customer_linked_at ELSE ? END,
               customer_linked_by=CASE WHEN ? IS NULL THEN customer_linked_by ELSE ? END
           WHERE id=?""",
        (values.get("customer_master_id"), values["customer_name"], values.get("original_customer_name"),
         values.get("customer_link_status", "unlinked"), values.get("customer_master_id"), timestamp,
         values.get("customer_master_id"), actor_id, sale_id),
    )


def inherit_sale_customer_link(db, target_sale_id, source_sale, actor_id, timestamp):
    if "customer_master_id" not in source_sale.keys():
        return
    db.execute(
        """UPDATE monthly_sales
           SET customer_master_id=?,original_customer_name=?,customer_link_status=?,
               customer_linked_at=?,customer_linked_by=? WHERE id=?""",
        (source_sale["customer_master_id"], source_sale["original_customer_name"],
         source_sale["customer_link_status"], source_sale["customer_linked_at"],
         source_sale["customer_linked_by"] or actor_id, target_sale_id),
    )


def month_setting(db, target_month, now):
    row = db.execute(
        "SELECT * FROM monthly_sales_month_settings WHERE target_month = ?", (target_month,)
    ).fetchone()
    if row:
        return row_dict(row)
    return {
        "target_month": target_month, "month_status": "open", "closed_by": None,
        "initial_cutoff_date": None, "round1_cutoff_date": None, "round2_cutoff_date": None,
        "round3_cutoff_date": None, "final_cutoff_date": None,
        "closed_at": None, "reopened_by": None, "reopened_at": None,
        "reopen_reason": None, "updated_at": None,
    }


def plan_rates(db, target_month):
    rows = db.execute(
        "SELECT * FROM monthly_sales_rates WHERE target_month=?", (target_month,)
    ).fetchall()
    result = {currency: None for currency in PLAN_CURRENCIES}
    for row in rows:
        currency = LEGACY_CURRENCY_ALIASES.get(row["currency"], row["currency"])
        # A native CNH entry wins over a retained legacy CNY entry.
        if currency in result and (result[currency] is None or row["currency"] != "CNY"):
            result[currency] = float(row["plan_rate"])
    result["KRW"] = 1.0
    return result


def apply_plan_rates(db, target_month, incoming, correction_reason, actor, now):
    rates = {
        currency: (1.0 if currency == "KRW" else number(
            incoming.get(currency, incoming.get("CNY") if currency == "CNH" else None),
            f"{currency} 기준환율", 0
        ))
        for currency in PLAN_CURRENCIES
    }
    if rates["USD"] <= 0:
        raise ValueError("USD 기준환율을 입력하세요.")
    old_rates = plan_rates(db, target_month)
    changed_existing = any(
        old_rates.get(currency) not in (None, 0)
        and abs(float(old_rates[currency]) - rates[currency]) > 0.0001
        for currency in PLAN_CURRENCIES
    )
    if changed_existing and not correction_reason:
        raise ValueError("등록된 기준환율을 변경하려면 변경사유를 입력하세요.")
    timestamp = now()
    for currency, rate in rates.items():
        db.execute(
            """
            INSERT INTO monthly_sales_rates
              (target_month, currency, plan_rate, source, correction_reason, updated_by, updated_at)
            VALUES (?, ?, ?, '대시보드 입력', ?, ?, ?)
            ON CONFLICT(target_month, currency) DO UPDATE SET plan_rate=excluded.plan_rate,
              source=excluded.source, correction_reason=excluded.correction_reason,
              updated_by=excluded.updated_by, updated_at=excluded.updated_at
            """,
            (target_month, currency, rate, correction_reason, actor["id"], timestamp),
        )
    # This table is a default for new/replanned FCST only. Existing sale-level
    # snapshots must never move merely because the monthly default was edited.
    return rates, old_rates


def convert_transaction_amount(transaction_amount, transaction_currency, rates, require_rates=True):
    transaction_currency = LEGACY_CURRENCY_ALIASES.get(
        str(transaction_currency or "").strip().upper(), str(transaction_currency or "").strip().upper()
    )
    if transaction_currency not in PLAN_CURRENCIES:
        raise ValueError("환종은 USD·EUR·JPY·CNH·KRW 중에서 선택하세요.")
    amount = number(transaction_amount, "거래금액", 0.01, True)
    usd_rate = rates.get("USD") or 0
    selected_rate = 1.0 if transaction_currency == "KRW" else (rates.get(transaction_currency) or 0)
    missing = []
    if selected_rate <= 0:
        missing.append(transaction_currency)
    if transaction_currency != "USD" and usd_rate <= 0:
        missing.append("USD")
    if missing and require_rates:
        raise ValueError(f"상단 기준환율에 {', '.join(dict.fromkeys(missing))} 환율을 먼저 입력하세요.")
    if transaction_currency == "KRW":
        krw_amount = amount
    elif selected_rate > 0:
        krw_amount = amount / CURRENCY_UNITS[transaction_currency] * selected_rate
    else:
        krw_amount = 0
    if transaction_currency == "USD":
        amount_usd = amount
    else:
        amount_usd = krw_amount / usd_rate if usd_rate > 0 else 0
    return round(amount_usd, 4), round(krw_amount, 2), float(selected_rate)


def convert_with_rate_context(transaction_amount, transaction_currency, selected_rate, usd_rate):
    """Recalculate an FCST amount from the sale's immutable plan-rate context."""
    transaction_currency = LEGACY_CURRENCY_ALIASES.get(
        str(transaction_currency or "").strip().upper(), str(transaction_currency or "").strip().upper()
    )
    amount = number(transaction_amount, "거래금액", 0.01, True)
    selected_rate = 1.0 if transaction_currency == "KRW" else float(selected_rate or 0)
    usd_rate = float(usd_rate or 0)
    if selected_rate <= 0 or (transaction_currency != "USD" and usd_rate <= 0):
        raise ValueError("이 매출 건의 계획환율 Snapshot이 없습니다. 계획환율 재적용을 진행하세요.")
    krw_amount = amount if transaction_currency == "KRW" else amount * selected_rate
    amount_usd = amount if transaction_currency == "USD" else krw_amount / usd_rate
    return round(amount_usd, 4), round(krw_amount, 2)


def canonical_sale_currency(sale):
    keys = sale.keys() if hasattr(sale, "keys") else sale
    value = (
        sale["transaction_currency_standard"]
        if "transaction_currency_standard" in keys and sale["transaction_currency_standard"]
        else sale["transaction_currency"]
    )
    return LEGACY_CURRENCY_ALIASES.get(str(value or "USD").upper(), str(value or "USD").upper())


def legacy_storage_currency(currency):
    return "CNY" if currency == "CNH" else currency


def sales_owner_options(db):
    placeholders = ",".join("?" for _ in MONTHLY_SALES_OWNERS)
    return [
        row_dict(row) for row in db.execute(
            f"""SELECT id,display_name,role FROM users
                WHERE status='active' AND deleted_at IS NULL
                  AND display_name IN ({placeholders})
                ORDER BY CASE display_name
                  WHEN '박정현' THEN 1 WHEN '김경태' THEN 2 WHEN '장윤선' THEN 3
                  WHEN '최령' THEN 4 WHEN '이인경' THEN 5 WHEN '김예원' THEN 6 ELSE 99 END""",
            MONTHLY_SALES_OWNERS,
        ).fetchall()
    ]


def resolve_owner_user(db, owner_user_id, owner_name, existing=None):
    existing_id = existing["owner_user_id"] if existing and "owner_user_id" in existing.keys() else None
    existing_name = existing["owner_name"] if existing else ""
    requested_id = owner_user_id if owner_user_id not in (None, "") else existing_id
    if requested_id not in (None, ""):
        # owner_name is a sale-time display snapshot. An unrelated edit must not
        # rewrite it merely because the user's current display name changed.
        if existing and str(requested_id) == str(existing_id):
            return existing_id, existing_name
        row = db.execute(
            "SELECT id,display_name,status,deleted_at FROM users WHERE id=?", (requested_id,)
        ).fetchone()
        if not row:
            raise ValueError("선택한 담당자 계정을 찾을 수 없습니다.")
        if (row["status"] != "active" or row["deleted_at"] is not None) and requested_id != existing_id:
            raise ValueError("사용 중인 담당자 계정만 새로 지정할 수 있습니다.")
        if row["display_name"] not in MONTHLY_SALES_OWNERS and requested_id != existing_id:
            raise ValueError("월별 매출 담당자로 지정된 사용자만 선택할 수 있습니다.")
        return row["id"], row["display_name"]
    # Legacy/import callers may still provide a name. Resolve only a unique,
    # active configured user; never guess among duplicate names.
    matches = db.execute(
        """SELECT id,display_name FROM users
           WHERE display_name=? AND status='active' AND deleted_at IS NULL""",
        (owner_name,),
    ).fetchall()
    if len(matches) == 1 and owner_name in MONTHLY_SALES_OWNERS:
        return matches[0]["id"], matches[0]["display_name"]
    if existing and owner_name == existing_name:
        return existing_id, existing_name
    raise ValueError("담당자는 활성 사용자 목록에서 선택하세요.")


def daily_rate_on_or_before(db, rate_date, currency):
    currency = LEGACY_CURRENCY_ALIASES.get(str(currency or "").upper(), str(currency or "").upper())
    if currency == "KRW":
        return {"rate_base_date": rate_date, "currency": "KRW", "rate": 1.0, "source": "KRW 기준통화"}
    aliases = (currency, "CNY") if currency == "CNH" else (currency,)
    placeholders = ",".join("?" for _ in aliases)
    row = db.execute(
        f"""SELECT * FROM monthly_sales_daily_rates
            WHERE currency IN ({placeholders}) AND rate_base_date<=?
            ORDER BY rate_base_date DESC, CASE WHEN currency=? THEN 0 ELSE 1 END LIMIT 1""",
        (*aliases, rate_date, currency),
    ).fetchone()
    return row_dict(row)


def assert_month_open(db, target_month, now):
    if month_setting(db, target_month, now)["month_status"] == "closed":
        raise PermissionError("마감된 월은 수정할 수 없습니다. 관리자가 먼저 재개방해야 합니다.")


def next_sales_no(db, target_month):
    prefix = f"MP-{target_month[2:4]}{target_month[5:7]}-"
    rows = db.execute(
        "SELECT sales_no FROM monthly_sales WHERE sales_no LIKE ?", (f"{prefix}%",)
    ).fetchall()
    values = []
    for row in rows:
        match = re.fullmatch(re.escape(prefix) + r"(\d{3,})", row["sales_no"])
        if match:
            values.append(int(match.group(1)))
    return f"{prefix}{max(values, default=0) + 1:03d}"


def next_child_no(db, root_no, marker):
    rows = db.execute(
        "SELECT sales_no FROM monthly_sales WHERE sales_no LIKE ?", (f"{root_no}-{marker}%",)
    ).fetchall()
    values = []
    for row in rows:
        match = re.fullmatch(re.escape(root_no) + rf"-{marker}(\d{{2,}})", row["sales_no"])
        if match:
            values.append(int(match.group(1)))
    return f"{root_no}-{marker}{max(values, default=0) + 1:02d}"


def root_sales(db, sale):
    root_id = sale["original_sales_id"] or sale["id"]
    root = db.execute("SELECT * FROM monthly_sales WHERE id = ?", (root_id,)).fetchone()
    return root or sale


def milestone_dict(db, sales_id):
    row = db.execute("SELECT * FROM monthly_sales_milestones WHERE sales_id = ?", (sales_id,)).fetchone()
    if not row:
        return {field: None if field != "exception_reason" else "" for field in MILESTONE_FIELDS}
    data = row_dict(row)
    data.pop("sales_id", None)
    data.pop("updated_at", None)
    return data


def carryover_milestones(db, sales_id):
    """Keep the commercial trail but never copy completed shipment facts."""
    milestones = milestone_dict(db, sales_id)
    milestones["shipment_actual_at"] = None
    milestones["shipping_actual_at"] = None
    return milestones


def derive_rate_and_status(values, milestones, apply_suggestion=True):
    status = values["sales_status"]
    suggestion = None
    if milestones.get("payment_actual_at") or milestones.get("shipment_actual_at"):
        suggestion = "confirmed"
    elif milestones.get("pi_confirmed_at"):
        suggestion = "scheduled"
    if suggestion and status != suggestion and apply_suggestion:
        status = suggestion

    actual_rate = values.get("actual_rate")
    applied_rate = values.get("plan_rate", 0)
    if milestones.get("shipment_actual_at"):
        # A salesperson-entered shipment is preliminary.  A cached operating
        # rate may complete its KRW conversion, but it is never labelled ERP Actual.
        rate_type = "shipment" if actual_rate and actual_rate > 0 else "pending"
    else:
        rate_type = "plan"
    return status, applied_rate, rate_type, suggestion


def normalize_sale_input(data, db, now, existing=None, actor=None, allow_legacy_confirmation=False):
    customer_name = clean_text(data.get("customer_name", existing["customer_name"] if existing else ""), "거래처명", 200, True)
    country = clean_text(data.get("country", existing["country"] if existing else ""), "국가", 100, True)
    requested_owner_name = clean_text(
        data.get("owner_name", existing["owner_name"] if existing else ""), "담당자", 100, True
    )
    owner_user_id, owner_name = resolve_owner_user(
        db, data.get("owner_user_id"), requested_owner_name, existing
    )
    business_unit = str(data.get("business_unit", existing["business_unit"] if existing else "")).strip()
    if business_unit not in BUSINESS_UNITS:
        raise ValueError("사업분야는 에스테틱·메디컬·덴탈 중에서 선택하세요.")
    original_month = str(data.get("original_month", existing["original_month"] if existing else "")).strip()
    target_month = str(data.get("target_month", existing["target_month"] if existing else "")).strip()
    if not valid_month(original_month) or not valid_month(target_month):
        raise ValueError("최초 추진월과 매출 대상월을 YYYY-MM 형식으로 입력하세요.")
    timing_value = data.get("timing_type")
    if timing_value is None and existing:
        timing_value = existing["timing_type"]
    timing_type = inferred_timing_type(timing_value, original_month, target_month)
    if timing_type not in TIMING_TYPES:
        raise ValueError("진행 시점 값이 올바르지 않습니다.")
    sales_status = str(data.get("sales_status", existing["sales_status"] if existing else "pipeline")).strip()
    if sales_status not in SALES_STATUSES:
        raise ValueError("매출 진행상태 값이 올바르지 않습니다.")
    customer_history = str(data.get(
        "customer_history", existing["customer_history"] if existing and "customer_history" in existing.keys() else "new"
    )).strip()
    if customer_history not in CUSTOMER_HISTORIES:
        raise ValueError("신규/기존은 신규 거래처 또는 기존 거래처 중에서 선택하세요.")
    carryover_decision = str(data.get(
        "carryover_decision", existing["carryover_decision"] if existing else "no"
    )).strip().lower()
    if carryover_decision not in {"yes", "no"}:
        raise ValueError("차월 구분은 차월 진행 또는 당월 유지 중에서 선택하세요.")
    carryover_decided_at = None
    carryover_target_month = None
    carryover_decision_reason = ""
    if carryover_decision == "yes":
        carryover_decided_at = valid_date(data.get(
            "carryover_decided_at", existing["carryover_decided_at"] if existing else None
        ))
        carryover_target_month = str(data.get(
            "carryover_target_month", existing["carryover_target_month"] if existing else ""
        )).strip()
        carryover_decision_reason = clean_text(data.get(
            "carryover_decision_reason", existing["carryover_decision_reason"] if existing else ""
        ), "차월 사유", 500)
        if not carryover_decided_at or not valid_month(carryover_target_month):
            raise ValueError("차월 진행을 선택한 경우 차월 결정일과 차월 대상월을 입력하세요.")
        if carryover_target_month <= target_month:
            raise ValueError("차월 대상월은 현재 매출 대상월보다 이후여야 합니다.")
    elif existing:
        active_carryover = db.execute(
            "SELECT 1 FROM monthly_sales_carryovers WHERE source_sales_id=? AND status='active' LIMIT 1",
            (existing["id"],),
        ).fetchone()
        if active_carryover:
            raise ValueError("이미 실제 차월 처리된 매출입니다. 당월 유지로 변경하려면 먼저 차월취소를 진행하세요.")
    existing_currency = canonical_sale_currency(existing) if existing else "USD"
    transaction_currency = str(data.get(
        "transaction_currency",
        existing_currency,
    )).strip().upper()
    transaction_currency = LEGACY_CURRENCY_ALIASES.get(transaction_currency, transaction_currency)
    if transaction_currency not in PLAN_CURRENCIES:
        raise ValueError("환종은 USD·EUR·JPY·CNH·KRW 중에서 선택하세요.")
    if existing and transaction_currency != existing_currency:
        raise ValueError("기존 매출의 환종은 일반 수정에서 바꿀 수 없습니다. 취소 후 재등록하거나 관리자에게 문의하세요.")
    transaction_amount = data.get(
        "transaction_amount",
        existing["transaction_amount"] if existing and "transaction_amount" in existing.keys() else data.get("amount_usd"),
    )
    if existing:
        plan_rate = float(existing["plan_rate"] or 0)
        plan_usd_rate = float(existing["plan_usd_rate"] or 0) if "plan_usd_rate" in existing.keys() else 0
        amount_usd, krw_amount = convert_with_rate_context(
            transaction_amount, transaction_currency, plan_rate, plan_usd_rate
        )
    else:
        month_rates = plan_rates(db, target_month)
        amount_usd, krw_amount, plan_rate = convert_transaction_amount(
            transaction_amount, transaction_currency, month_rates
        )
        plan_usd_rate = float(month_rates.get("USD") or 0)
    actual_rate_raw = data.get("actual_rate", existing["actual_rate"] if existing else None)
    actual_rate = number(actual_rate_raw, "출고일 환율", 0.0001) if actual_rate_raw not in (None, "") else None
    actual_rate_changed = actual_rate is not None and (
        not existing or existing["actual_rate"] in (None, "")
        or abs(float(existing["actual_rate"]) - actual_rate) > 0.0001
    )
    if actual_rate_changed and actor is not None and actor["role"] not in {"admin", "manager"}:
        raise ValueError("출고일 운영환율의 수동 입력·정정은 부서장 또는 관리자만 할 수 있습니다.")
    notes = clean_text(data.get("notes", existing["notes"] if existing else ""), "비고", 3000)
    milestones = {}
    current_milestones = milestone_dict(db, existing["id"]) if existing else {}
    incoming = data.get("milestones") or {}
    for field in MILESTONE_FIELDS:
        raw = incoming.get(field, current_milestones.get(field))
        milestones[field] = valid_date(raw) if field in DATE_FIELDS else clean_text(raw, field, 500)
    rate_base_date = valid_date(data.get("rate_base_date", existing["rate_base_date"] if existing else None))
    rate_source = clean_text(data.get("rate_source", existing["rate_source"] if existing else ""), "환율 출처", 200)
    preliminary_currency = str(data.get(
        "preliminary_actual_currency",
        existing["preliminary_actual_currency"] if existing and "preliminary_actual_currency" in existing.keys()
        and existing["preliminary_actual_currency"] else transaction_currency,
    ) or transaction_currency).strip().upper()
    preliminary_currency = LEGACY_CURRENCY_ALIASES.get(preliminary_currency, preliminary_currency)
    if preliminary_currency not in PLAN_CURRENCIES:
        raise ValueError("담당자 출고 환종은 USD·EUR·JPY·CNH·KRW 중에서 선택하세요.")
    preliminary_amount_raw = data.get(
        "preliminary_actual_amount",
        existing["preliminary_actual_amount"] if existing and "preliminary_actual_amount" in existing.keys() else None,
    )
    preliminary_amount = (
        number(preliminary_amount_raw, "담당자 출고금액", 0.01, True)
        if preliminary_amount_raw not in (None, "") else None
    )
    if milestones.get("shipment_actual_at") and preliminary_amount is None:
        preliminary_amount = number(transaction_amount, "담당자 출고금액", 0.01, True)
    if preliminary_amount is not None and not milestones.get("shipment_actual_at"):
        raise ValueError("담당자 출고금액을 입력하려면 실제 출고일을 입력하세요.")
    if milestones.get("shipment_actual_at") and actual_rate is None:
        cached_rate = daily_rate_on_or_before(db, milestones["shipment_actual_at"], preliminary_currency)
        if cached_rate:
            actual_rate = float(cached_rate["rate"])
            rate_base_date = cached_rate["rate_base_date"]
            rate_source = cached_rate["source"]
    if actual_rate is not None:
        if not milestones.get("shipment_actual_at"):
            raise ValueError("출고일 확정환율을 적용하려면 실제 출고일을 입력하세요.")
        if not rate_base_date or not rate_source:
            raise ValueError("출고일 확정환율의 실제 기준일과 출처를 입력하세요.")
        if rate_base_date > milestones["shipment_actual_at"]:
            raise ValueError("환율 기준일은 실제 출고일보다 늦을 수 없습니다.")
        if existing and existing["actual_rate"] not in (None, "") and abs(float(existing["actual_rate"]) - actual_rate) > 0.0001:
            clean_text(data.get("change_reason"), "환율 정정사유", 500, True)
    confirmation_basis = clean_text(data.get(
        "confirmation_basis",
        existing["confirmation_basis"] if existing and "confirmation_basis" in existing.keys() else "",
    ), "영업 확정근거", 100)
    confirmation_note = clean_text(data.get(
        "confirmation_note",
        existing["confirmation_note"] if existing and "confirmation_note" in existing.keys() else "",
    ), "영업 확정근거 메모", 500)
    if not confirmation_basis and milestones.get("payment_actual_at"):
        # The business event itself is the confirmation evidence; do not make
        # the user select the same fact twice.
        confirmation_basis = "payment_complete"
    elif not confirmation_basis and milestones.get("shipment_actual_at"):
        # A preliminary shipment record is stronger evidence than the sales
        # confidence label, while remaining clearly separate from ERP Actual.
        confirmation_basis = "other"
        confirmation_note = confirmation_note or "담당자 출고입력"
    if confirmation_basis and confirmation_basis not in CONFIRMATION_BASES:
        legacy_basis = existing["confirmation_basis"] if existing and "confirmation_basis" in existing.keys() else ""
        if confirmation_basis != legacy_basis:
            raise ValueError("영업 확정근거 값이 올바르지 않습니다.")
    becoming_confirmed = sales_status == "confirmed" and (not existing or existing["sales_status"] != "confirmed")
    if becoming_confirmed and not confirmation_basis and not allow_legacy_confirmation:
        raise ValueError("영업상 확정으로 변경할 때는 확정근거를 선택하세요.")
    if confirmation_basis == "other" and not confirmation_note:
        raise ValueError("기타 확정근거의 내용을 입력하세요.")
    effective_input = data.get("effective_at")
    effective_at = valid_date(effective_input) if effective_input not in (None, "") else datetime.now(SEOUL).date().isoformat()
    effective_at_source = "user_business_date" if effective_input not in (None, "") else "system_default"
    preliminary_krw = None
    if preliminary_amount is not None and actual_rate is not None:
        preliminary_krw = round(
            preliminary_amount if preliminary_currency == "KRW" else preliminary_amount * actual_rate, 2
        )
    values = {
        "customer_name": customer_name, "country": country, "owner_name": owner_name,
        "owner_user_id": owner_user_id,
        "business_unit": business_unit, "original_month": original_month,
        "target_month": target_month, "timing_type": timing_type,
        "customer_history": customer_history,
        "transaction_currency": transaction_currency,
        "transaction_currency_storage": legacy_storage_currency(transaction_currency),
        "transaction_amount": number(transaction_amount, "거래금액", 0.01, True),
        "amount_usd": amount_usd, "plan_rate": plan_rate, "plan_usd_rate": plan_usd_rate,
        "actual_rate": actual_rate,
        "sales_status": sales_status, "notes": notes,
        "confirmation_basis": confirmation_basis, "confirmation_note": confirmation_note,
        "preliminary_actual_currency": preliminary_currency if preliminary_amount is not None else None,
        "preliminary_actual_amount": preliminary_amount,
        "preliminary_actual_krw_amount": preliminary_krw,
        "preliminary_reference_no": clean_text(data.get(
            "preliminary_reference_no",
            existing["preliminary_reference_no"] if existing and "preliminary_reference_no" in existing.keys() else "",
        ), "담당자 출고 참조번호", 150),
        "preliminary_note": clean_text(data.get(
            "preliminary_note",
            existing["preliminary_note"] if existing and "preliminary_note" in existing.keys() else "",
        ), "담당자 출고 메모", 1000),
        "effective_at": effective_at, "effective_at_source": effective_at_source,
        "carryover_decision": carryover_decision,
        "carryover_decided_at": carryover_decided_at,
        "carryover_target_month": carryover_target_month,
        "carryover_decision_reason": carryover_decision_reason,
        "rate_base_date": rate_base_date,
        "rate_source": rate_source,
        "customer_master_id": data.get(
            "customer_master_id",
            existing["customer_master_id"] if existing and "customer_master_id" in existing.keys() else None,
        ),
    }
    customer_master_id, display_name, original_customer_name, customer_link_status = resolve_sale_customer(
        values, db, existing
    )
    values.update({
        "customer_master_id": customer_master_id,
        "customer_name": display_name,
        "original_customer_name": original_customer_name,
        "customer_link_status": customer_link_status,
    })
    apply_status_suggestion = data.get("apply_status_suggestion", True) not in (False, 0, "0", "false", "False")
    status, applied_rate, rate_type, suggestion = derive_rate_and_status(values, milestones, apply_status_suggestion)
    if status == "confirmed" and not confirmation_basis and not allow_legacy_confirmation and (
        not existing or existing["sales_status"] != "confirmed" or status != sales_status
    ):
        raise ValueError("영업상 확정으로 변경할 때는 확정근거를 선택하세요.")
    values.update({
        "sales_status": status, "applied_rate": applied_rate, "rate_type": rate_type,
        # krw_amount remains the Forecast plan value. Preliminary/Actual values
        # live in their own columns and never overwrite the plan snapshot.
        "krw_amount": krw_amount,
        "milestones": milestones,
        "status_suggestion": suggestion if suggestion and suggestion != sales_status else None,
        "status_adjusted_to": suggestion if apply_status_suggestion and suggestion and suggestion != sales_status else None,
    })
    return values


def upsert_milestones(db, sales_id, milestones, timestamp):
    db.execute(
        """
        INSERT INTO monthly_sales_milestones
          (sales_id, order_agreed_at, po_received_at, pi_no, pi_sent_at, pi_confirmed_at,
           payment_expected_at, payment_actual_at, shipment_expected_at, shipment_actual_at,
           shipping_expected_at, shipping_actual_at, exception_reason, updated_at)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(sales_id) DO UPDATE SET
          order_agreed_at=excluded.order_agreed_at, po_received_at=excluded.po_received_at,
          pi_no=excluded.pi_no, pi_sent_at=excluded.pi_sent_at,
          pi_confirmed_at=excluded.pi_confirmed_at, payment_expected_at=excluded.payment_expected_at,
          payment_actual_at=excluded.payment_actual_at, shipment_expected_at=excluded.shipment_expected_at,
          shipment_actual_at=excluded.shipment_actual_at, shipping_expected_at=excluded.shipping_expected_at,
          shipping_actual_at=excluded.shipping_actual_at, exception_reason=excluded.exception_reason,
          updated_at=excluded.updated_at
        """,
        (sales_id, *(milestones.get(field) for field in MILESTONE_FIELDS), timestamp),
    )


def touch_candidate(db, values, timestamp):
    db.execute(
        """
        INSERT INTO monthly_sales_master_candidates
          (customer_name, country, owner_name, business_unit, usage_count, first_seen_at, last_seen_at, status)
        VALUES (?, ?, ?, ?, 1, ?, ?, 'candidate')
        ON CONFLICT(customer_name, country, owner_name, business_unit) DO UPDATE SET
          usage_count = usage_count + 1, last_seen_at = excluded.last_seen_at
        """,
        (values["customer_name"], values["country"], values["owner_name"], values["business_unit"], timestamp, timestamp),
    )


def sale_snapshot(db, sales_id):
    row = db.execute(
        """
        SELECT s.*, creator.display_name AS created_by_name, updater.display_name AS updated_by_name
        FROM monthly_sales s
        LEFT JOIN users creator ON creator.id = s.created_by
        LEFT JOIN users updater ON updater.id = s.updated_by
        WHERE s.id = ?
        """, (sales_id,),
    ).fetchone()
    if not row:
        return None
    data = row_dict(row)
    legacy_currency = data.get("transaction_currency") or "USD"
    data["transaction_currency_legacy"] = legacy_currency
    data["transaction_currency"] = canonical_sale_currency(data)
    data["milestones"] = milestone_dict(db, sales_id)
    try:
        data["commercial_context"] = json.loads(data.get("commercial_context_json") or "{}")
    except (TypeError, json.JSONDecodeError):
        data["commercial_context"] = {}
    return data


def history_change_set(before, after):
    ignored = {"created_by_name", "updated_by_name", "warnings", "progress", "restore_allowed"}
    before = before or {}
    after = after or {}
    changes = {}
    for key in sorted((set(before) | set(after)) - ignored):
        old_value, new_value = before.get(key), after.get(key)
        if old_value != new_value:
            changes[key] = {"before": old_value, "after": new_value}
    return changes


def add_history(
    db, sale_id, event_type, reason, actor, now, *, effective_at=None,
    effective_at_source=None,
):
    snapshot = sale_snapshot(db, sale_id)
    previous_row = db.execute(
        "SELECT snapshot_json FROM monthly_sales_history WHERE sales_id=? ORDER BY id DESC LIMIT 1",
        (sale_id,),
    ).fetchone()
    previous = None
    if previous_row:
        try:
            previous = json.loads(previous_row["snapshot_json"])
        except (TypeError, json.JSONDecodeError):
            previous = None
    if effective_at is None:
        effective_at = datetime.now(SEOUL).date().isoformat()
        effective_at_source = effective_at_source or "system_default"
    else:
        effective_at = valid_date(effective_at)
        effective_at_source = effective_at_source or "user_business_date"
    db.execute(
        """
        INSERT INTO monthly_sales_history
          (sales_id, sales_no, event_type, reason, actor_user_id, actor_name, occurred_at,
           effective_at,effective_at_source,change_set_json,snapshot_json)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (sale_id, snapshot["sales_no"], event_type, reason or "", actor["id"], actor["display_name"], now(),
         effective_at, effective_at_source,
         json.dumps(history_change_set(previous, snapshot), ensure_ascii=False, separators=(",", ":"), default=str),
         json.dumps(snapshot, ensure_ascii=False, separators=(",", ":"), default=str)),
    )


def mark_sheet_dirty(db, now):
    db.execute(
        "UPDATE monthly_sales_sheet_sync SET dirty = 1, updated_at = ? WHERE id = 1", (now(),)
    )


def can_edit_sale(user, sale):
    return user["role"] in {"admin", "manager", "editor"}


def warning_list(sale, today=None, duplicate_suspected=False):
    today = today or datetime.now(SEOUL).date()
    milestones = sale.get("milestones", {})
    warnings = []
    if sale.get("customer_link_status") in {"unlinked", "review"}:
        warnings.append({"level": "warning", "code": "CUSTOMER_LINK_REQUIRED", "message": "거래처 마스터 연결 필요"})
    if sale.get("rate_type") == "pending":
        warnings.append({"level": "error", "code": "RATE_PENDING", "message": "출고일 환율 확인 필요"})
    elif sale.get("rate_type") == "plan" and not sale.get("applied_rate"):
        warnings.append({"level": "error", "code": "PLAN_RATE_MISSING", "message": "기준환율 미등록"})
    for field, label, actual in (
        ("payment_expected_at", "예상 수금일 경과", "payment_actual_at"),
        ("shipment_expected_at", "예상 출고일 경과", "shipment_actual_at"),
        ("shipping_expected_at", "예상 선적일 경과", "shipping_actual_at"),
    ):
        if milestones.get(field) and not milestones.get(actual):
            due = datetime.strptime(milestones[field], "%Y-%m-%d").date()
            if due < today:
                warnings.append({"level": "warning", "code": field.upper(), "message": label})
    inferred = "confirmed" if milestones.get("payment_actual_at") or milestones.get("shipment_actual_at") else (
        "scheduled" if milestones.get("pi_confirmed_at") else None
    )
    if inferred and sale.get("sales_status") != inferred:
        warnings.append({"level": "warning", "code": "STATUS_MISMATCH", "message": f"진행단계상 {SALES_STATUSES[inferred]} 상태 권장"})
    if duplicate_suspected:
        warnings.append({"level": "warning", "code": "DUPLICATE_SUSPECTED", "message": "동일 거래처·월·금액의 중복 매출 의심"})
    later_actual_seen = False
    stage_actuals = (
        "shipping_actual_at", "shipment_actual_at", "payment_actual_at", "pi_confirmed_at",
        "pi_sent_at", "po_received_at", "order_agreed_at",
    )
    for field in stage_actuals:
        if milestones.get(field):
            later_actual_seen = True
        elif later_actual_seen and not milestones.get("exception_reason"):
            warnings.append({"level": "warning", "code": "STAGE_SKIPPED", "message": "진행단계 건너뜀—예외사유 확인"})
            break
    if sale.get("sales_status") == "pipeline":
        try:
            updated = datetime.fromisoformat(str(sale.get("updated_at", "")).replace("Z", "+00:00"))
            updated_date = updated.astimezone(SEOUL).date() if updated.tzinfo else updated.date()
            if (today - updated_date).days >= 14:
                warnings.append({"level": "warning", "code": "STALE_PIPELINE", "message": "14일 이상 업데이트되지 않은 추진 건"})
        except (TypeError, ValueError):
            pass
    try:
        target_year, target_month = map(int, sale.get("target_month", "").split("-"))
        month_end = datetime(target_year, target_month, monthrange(target_year, target_month)[1]).date()
        expected = milestones.get("shipment_expected_at")
        expected_date = datetime.strptime(expected, "%Y-%m-%d").date() if expected else None
        if sale.get("record_status") == "active" and not milestones.get("shipment_actual_at") and (
            (expected_date and expected_date > month_end) or (not expected_date and 0 <= (month_end - today).days <= 7)
        ):
            warnings.append({"level": "warning", "code": "MONTH_END_SHIPMENT_RISK", "message": "월말 이전 출고 어려움 확인"})
    except (AttributeError, TypeError, ValueError):
        pass
    if not all(sale.get(field) for field in ("customer_name", "country", "owner_name", "business_unit")):
        warnings.append({"level": "error", "code": "REQUIRED_MISSING", "message": "기본정보 누락"})
    context = sale.get("commercial_context") or {}
    documents = context.get("required_documents") or []
    if documents:
        warnings.append({
            "level": "info", "code": "REQUIRED_DOCUMENTS",
            "message": "Required Docs: " + " · ".join(item.get("label") or item.get("document_code") or "" for item in documents),
        })
    local_dates = [item.get("holiday_date") for item in (context.get("country_holidays") or [])]
    local_dates += [
        f"{item.get('start_date')}~{item.get('end_date')} {item.get('title')}"
        for item in (context.get("customer_closures") or [])
    ]
    if local_dates:
        warnings.append({"level": "info", "code": "LOCAL_SCHEDULE", "message": "현지 일정: " + " · ".join(filter(None, local_dates))})
    if context.get("logistics_warning"):
        warnings.append({"level": "info", "code": "LOGISTICS_NOTE", "message": context["logistics_warning"]})
    sale["warnings"] = warnings
    return sale


def progress_payload(sale, today=None):
    milestones = sale.get("milestones", {})
    stages = [
        ("order", "오더합의", "order_agreed_at", None),
        ("po", "PO 수신", "po_received_at", None),
        ("pi_sent", "PI 발송", "pi_sent_at", None),
        ("pi_confirm", "PI 컨펌", "pi_confirmed_at", None),
        ("payment", "수금", "payment_actual_at", "payment_expected_at"),
        ("shipment", "출고", "shipment_actual_at", "shipment_expected_at"),
        ("shipping", "선적", "shipping_actual_at", "shipping_expected_at"),
    ]
    result = []
    today = today or datetime.now(SEOUL).date()
    completed_count = 0
    for key, label, actual_field, expected_field in stages:
        actual = milestones.get(actual_field)
        expected = milestones.get(expected_field) if expected_field else None
        overdue = bool(expected and not actual and datetime.strptime(expected, "%Y-%m-%d").date() < today)
        if actual:
            completed_count += 1
        result.append({"key": key, "label": label, "actual": actual, "expected": expected, "completed": bool(actual), "overdue": overdue})
    for index, stage in enumerate(result):
        stage["current"] = index == completed_count and completed_count < len(result)
    return result


def _recorded_timestamp(value):
    if not value:
        return None
    try:
        parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        if parsed.tzinfo is None:
            parsed = parsed.replace(tzinfo=timezone.utc)
        return parsed.astimezone(timezone.utc)
    except (TypeError, ValueError):
        try:
            return datetime.strptime(str(value)[:10], "%Y-%m-%d").replace(tzinfo=timezone.utc)
        except (TypeError, ValueError):
            return None


def _ledger_section(row):
    record_status = row.get("record_status")
    if record_status == "active":
        return row.get("sales_status") or "undecided"
    if record_status == "carried_over":
        return "carried_over"
    if record_status == "cancelled" and row.get("latest_event_type") in {
        "CARRYOVER_TARGET_CANCEL", "SPLIT_CHILD_CANCEL",
    }:
        return "lineage_hidden"
    if record_status == "cancelled":
        return "cancelled"
    return "lineage_hidden"


def enrich_carryover_view(db, rows, recorded_cutoff=None):
    """Attach display-only carryover context without changing either ledger.

    Link timestamps decide whether the operation was known at an as-of cutoff;
    business dates remain available for the labels shown to the user.
    """
    if not rows:
        return rows
    sale_ids = {str(row.get("id")) for row in rows if row.get("id")}
    if not sale_ids:
        return rows
    placeholders = ",".join("?" for _ in sale_ids)
    links = db.execute(
        f"""SELECT * FROM monthly_sales_carryovers
            WHERE source_sales_id IN ({placeholders}) OR target_sales_id IN ({placeholders})
            ORDER BY created_at,id""",
        (*sale_ids, *sale_ids),
    ).fetchall()
    source_ids = {str(link["source_sales_id"]) for link in links}
    cancel_events = defaultdict(list)
    if source_ids:
        event_placeholders = ",".join("?" for _ in source_ids)
        for event in db.execute(
            f"""SELECT sales_id,effective_at,occurred_at,id
                FROM monthly_sales_history
                WHERE event_type='CARRYOVER_CANCEL'
                  AND sales_id IN ({event_placeholders})
                ORDER BY occurred_at,id""",
            tuple(source_ids),
        ).fetchall():
            cancel_events[str(event["sales_id"])].append(event)

    cutoff = _recorded_timestamp(recorded_cutoff)
    outgoing = defaultdict(list)
    incoming = defaultdict(list)
    for link in links:
        created_at = _recorded_timestamp(link["created_at"])
        if cutoff and created_at and created_at > cutoff:
            continue
        cancelled_at = _recorded_timestamp(link["cancelled_at"])
        cancelled_in_view = bool(cancelled_at and (not cutoff or cancelled_at <= cutoff))
        cancel_effective_at = None
        if cancelled_in_view:
            candidates = []
            for event in cancel_events.get(str(link["source_sales_id"]), []):
                event_at = _recorded_timestamp(event["occurred_at"])
                if cutoff and event_at and event_at > cutoff:
                    continue
                if created_at and event_at and event_at < created_at:
                    continue
                candidates.append((event_at, event))
            if candidates:
                if cancelled_at:
                    candidates.sort(key=lambda pair: abs((pair[0] - cancelled_at).total_seconds()) if pair[0] else float("inf"))
                cancel_effective_at = candidates[0][1]["effective_at"]
            if not cancel_effective_at and cancelled_at:
                cancel_effective_at = cancelled_at.astimezone(SEOUL).date().isoformat()
        payload = {
            "id": link["id"],
            "source_sales_id": link["source_sales_id"],
            "target_sales_id": link["target_sales_id"],
            "source_month": link["source_month"],
            "target_month": link["target_month"],
            "amount_usd": link["amount_usd"],
            "carryover_date": link["carryover_date"],
            "status": "cancelled" if cancelled_in_view else "active",
            "cancel_effective_at": cancel_effective_at,
        }
        outgoing[str(link["source_sales_id"])].append(payload)
        incoming[str(link["target_sales_id"])].append(payload)

    for row in rows:
        sale_id = str(row.get("id") or "")
        row["carryover_history"] = outgoing.get(sale_id, [])
        row["carryover_in_history"] = incoming.get(sale_id, [])
        row["ledger_section"] = _ledger_section(row)
    return rows


def current_sales(db, target_month, user):
    sql = """
        SELECT s.*, creator.display_name AS created_by_name, updater.display_name AS updated_by_name
        FROM monthly_sales s
        LEFT JOIN users creator ON creator.id = s.created_by
        LEFT JOIN users updater ON updater.id = s.updated_by
        WHERE s.target_month = ?
    """
    params = [target_month]
    sql += " ORDER BY CASE s.sales_status WHEN 'confirmed' THEN 1 WHEN 'scheduled' THEN 2 WHEN 'pipeline' THEN 3 ELSE 4 END, s.customer_name, s.sales_no"
    latest_events = {
        row["sales_id"]: row["event_type"]
        for row in db.execute(
            """SELECT h.sales_id, h.event_type
               FROM monthly_sales_history h
               JOIN (
                 SELECT sales_id, MAX(id) AS max_id
                 FROM monthly_sales_history
                 GROUP BY sales_id
               ) latest ON latest.max_id=h.id"""
        ).fetchall()
    }
    rows = []
    for row in db.execute(sql, params).fetchall():
        item = row_dict(row)
        item["transaction_currency_legacy"] = item.get("transaction_currency") or "USD"
        item["transaction_currency"] = canonical_sale_currency(item)
        try:
            item["commercial_context"] = json.loads(item.get("commercial_context_json") or "{}")
        except (TypeError, json.JSONDecodeError):
            item["commercial_context"] = {}
        item["milestones"] = milestone_dict(db, item["id"])
        item["progress"] = progress_payload(item)
        item["latest_event_type"] = latest_events.get(item["id"])
        item["restore_allowed"] = (
            item["record_status"] == "cancelled" and latest_events.get(item["id"]) == "CANCEL"
        )
        rows.append(item)
    enrich_carryover_view(db, rows)
    duplicate_counts = defaultdict(int)
    for item in rows:
        if item["record_status"] == "active":
            duplicate_counts[(
                item["customer_name"].strip().lower(), item.get("transaction_currency", "USD"),
                round(item.get("transaction_amount") or item["amount_usd"], 4), item["target_month"],
            )] += 1
    for item in rows:
        signature = (
            item["customer_name"].strip().lower(), item.get("transaction_currency", "USD"),
            round(item.get("transaction_amount") or item["amount_usd"], 4), item["target_month"],
        )
        warning_list(item, duplicate_suspected=item["record_status"] == "active" and duplicate_counts[signature] > 1)
    return rows


def imported_round_rows(db, target_month, round_key=None, cutoff_date=None, latest_before=None):
    """Return an immutable imported round snapshot, or None when no batch exists."""
    clauses = ["target_month=?"]
    params = [target_month]
    if round_key:
        clauses.append("round_key=?")
        params.append(round_key)
    if cutoff_date:
        clauses.append("cutoff_date=?")
        params.append(cutoff_date)
    if latest_before:
        clauses.append("cutoff_date<=?")
        params.append(latest_before)
    batch = db.execute(
        f"""SELECT * FROM monthly_sales_round_imports
            WHERE {' AND '.join(clauses)}
            ORDER BY cutoff_date DESC,
              CASE round_key WHEN 'final' THEN 5 WHEN 'round3' THEN 4 WHEN 'round2' THEN 3
                             WHEN 'round1' THEN 2 ELSE 1 END DESC,
              imported_at DESC
            LIMIT 1""",
        params,
    ).fetchone()
    if not batch:
        return None
    snapshot_day = datetime.strptime(batch["cutoff_date"], "%Y-%m-%d").date()
    rows = []
    for stored in db.execute(
        "SELECT * FROM monthly_sales_round_rows WHERE import_id=? ORDER BY source_row",
        (batch["id"],),
    ).fetchall():
        try:
            item = json.loads(stored["snapshot_json"])
        except (TypeError, json.JSONDecodeError):
            continue
        item.update({
            "id": stored["id"],
            "sales_no": stored["sales_no"],
            "target_month": stored["target_month"],
            "historical_round_key": stored["round_key"],
            "historical_cutoff_date": stored["cutoff_date"],
            "historical_import_id": stored["import_id"],
            "linked_sales_id": stored["linked_sales_id"],
            "restore_allowed": False,
        })
        item["transaction_currency_legacy"] = item.get("transaction_currency_legacy") or item.get("transaction_currency") or "USD"
        item["transaction_currency"] = LEGACY_CURRENCY_ALIASES.get(
            str(item.get("transaction_currency_standard") or item.get("transaction_currency") or "USD").upper(),
            str(item.get("transaction_currency_standard") or item.get("transaction_currency") or "USD").upper(),
        )
        item.setdefault("record_status", "active")
        item.setdefault("split_role", "none")
        item.setdefault("carryover_role", "none")
        item.setdefault("original_sales_id", None)
        item.setdefault("split_group_id", None)
        item.setdefault("customer_link_status", "linked" if stored["linked_sales_id"] else "unlinked")
        item.setdefault("milestones", {field: None if field != "exception_reason" else "" for field in MILESTONE_FIELDS})
        item["latest_event_type"] = None
        item["carryover_history"] = []
        item["carryover_in_history"] = []
        item["ledger_section"] = _ledger_section(item)
        item["progress"] = progress_payload(item, snapshot_day)
        rows.append(item)
    duplicate_counts = defaultdict(int)
    for item in rows:
        duplicate_counts[(
            item["customer_name"].strip().lower(), item.get("transaction_currency", "USD"),
            round(item.get("transaction_amount") or item["amount_usd"], 4), item["target_month"],
        )] += 1
    for item in rows:
        signature = (
            item["customer_name"].strip().lower(), item.get("transaction_currency", "USD"),
            round(item.get("transaction_amount") or item["amount_usd"], 4), item["target_month"],
        )
        warning_list(item, snapshot_day, duplicate_counts[signature] > 1)
    return rows


def round_imports_payload(db, target_month):
    labels = {key: label for key, label, _column in ROUND_DEFINITIONS}
    return [
        {
            **row_dict(row),
            "round_label": labels.get(row["round_key"], row["round_key"]),
        }
        for row in db.execute(
            """SELECT i.*, u.display_name AS imported_by_name,
                      (SELECT COUNT(*) FROM monthly_sales_round_import_revisions r
                       WHERE r.import_id=i.id) AS revision_count
               FROM monthly_sales_round_imports i
               LEFT JOIN users u ON u.id=i.imported_by
               WHERE i.target_month=?
               ORDER BY CASE i.round_key WHEN 'initial' THEN 1 WHEN 'round1' THEN 2
                         WHEN 'round2' THEN 3 WHEN 'round3' THEN 4 ELSE 5 END""",
            (target_month,),
        ).fetchall()
    ]


def apply_history_change_set(state, change_set):
    state = dict(state or {})
    for key, values in (change_set or {}).items():
        if isinstance(values, dict) and "after" in values:
            state[key] = values.get("after")
    return state


def historical_sales(db, target_month, as_of, user, round_key=None):
    imported = imported_round_rows(
        db, target_month, round_key=round_key, cutoff_date=as_of
    ) if round_key else None
    if imported is not None:
        return imported
    # The default as-of view reproduces what the system knew by the selected
    # business day's end.  effective_at remains immutable business metadata and
    # never moves a later entry into an earlier recorded-state view.
    recorded_cutoff = cutoff_utc(as_of)
    rows = db.execute(
        """SELECT sales_id,event_type,occurred_at,effective_at,effective_at_source,
                  change_set_json,snapshot_json,id
           FROM monthly_sales_history
           WHERE occurred_at<=?
           ORDER BY sales_id,occurred_at,id""",
        (recorded_cutoff,),
    ).fetchall()
    states = {}
    latest_events = {}
    for row in rows:
        try:
            snapshot = json.loads(row["snapshot_json"])
        except (TypeError, json.JSONDecodeError):
            continue
        try:
            changes = json.loads(row["change_set_json"] or "{}")
        except (TypeError, json.JSONDecodeError):
            changes = {}
        if row["sales_id"] not in states or not changes:
            states[row["sales_id"]] = snapshot
        else:
            states[row["sales_id"]] = apply_history_change_set(states[row["sales_id"]], changes)
        latest_events[row["sales_id"]] = row["event_type"]
    output = []
    for sale_id, item in states.items():
        if item.get("target_month") != target_month:
            continue
        item["transaction_currency_legacy"] = item.get("transaction_currency_legacy") or item.get("transaction_currency") or "USD"
        item["transaction_currency"] = LEGACY_CURRENCY_ALIASES.get(
            str(item.get("transaction_currency_standard") or item.get("transaction_currency") or "USD").upper(),
            str(item.get("transaction_currency_standard") or item.get("transaction_currency") or "USD").upper(),
        )
        item["progress"] = progress_payload(item, datetime.strptime(as_of, "%Y-%m-%d").date())
        item["latest_event_type"] = latest_events.get(sale_id)
        item["restore_allowed"] = False
        output.append(item)
    enrich_carryover_view(db, output, recorded_cutoff)
    duplicate_counts = defaultdict(int)
    for item in output:
        if item["record_status"] == "active":
            duplicate_counts[(
                item["customer_name"].strip().lower(), item.get("transaction_currency", "USD"),
                round(item.get("transaction_amount") or item["amount_usd"], 4), item["target_month"],
            )] += 1
    for item in output:
        signature = (
            item["customer_name"].strip().lower(), item.get("transaction_currency", "USD"),
            round(item.get("transaction_amount") or item["amount_usd"], 4), item["target_month"],
        )
        warning_list(
            item,
            datetime.strptime(as_of, "%Y-%m-%d").date(),
            item["record_status"] == "active" and duplicate_counts[signature] > 1,
        )
    return output


def initial_fcst_map(db, target_month):
    rows = db.execute(
        "SELECT sales_id, amount_usd, krw_amount FROM monthly_fcst_snapshots WHERE target_month = ?",
        (target_month,),
    ).fetchall()
    return {row["sales_id"]: {"amount_usd": row["amount_usd"], "krw_amount": row["krw_amount"]} for row in rows}


def summary_payload(
    db, rows, target_month, as_of=None, selected_business_unit="", initial_override_rows=None
):
    targets = {
        row["business_unit"]: row_dict(row)
        for row in db.execute("SELECT * FROM monthly_sales_targets WHERE target_month = ?", (target_month,)).fetchall()
    }
    row_ids = {row["id"] for row in rows}
    initial_rows = initial_override_rows if initial_override_rows is not None else []
    if initial_override_rows is None and row_ids:
        placeholders = ",".join("?" for _ in row_ids)
        initial_rows = db.execute(
            f"SELECT sales_id, business_unit, amount_usd, krw_amount FROM monthly_fcst_snapshots WHERE target_month=? AND sales_id IN ({placeholders})",
            (target_month, *row_ids),
        ).fetchall()
    carryover_rows = []
    if row_ids:
        placeholders = ",".join("?" for _ in row_ids)
        if as_of:
            cutoff = cutoff_utc(as_of)
            carryover_rows = db.execute(
                f"""SELECT source_sales_id, amount_usd FROM monthly_sales_carryovers
                    WHERE source_month=? AND source_sales_id IN ({placeholders})
                      AND created_at<=? AND (cancelled_at IS NULL OR cancelled_at>?)""",
                (target_month, *row_ids, cutoff, cutoff),
            ).fetchall()
        else:
            carryover_rows = db.execute(
                f"""SELECT source_sales_id, amount_usd FROM monthly_sales_carryovers
                    WHERE source_month=? AND source_sales_id IN ({placeholders}) AND status='active'""",
                (target_month, *row_ids),
            ).fetchall()
    row_by_id = {row["id"]: row for row in rows}
    scope_units = {selected_business_unit} if selected_business_unit in BUSINESS_UNITS else set(BUSINESS_UNITS)
    units = list(BUSINESS_UNITS) + ["total"]
    summary = {}
    for unit in units:
        unit_rows = rows if unit == "total" else [row for row in rows if row["business_unit"] == unit]
        active = [row for row in unit_rows if row["record_status"] == "active"]
        values = {
            "business_unit": unit,
            "label": "합계" if unit == "total" else BUSINESS_UNITS[unit],
            "target_usd": sum(targets.get(key, {}).get("target_usd", 0) for key in scope_units) if unit == "total" else (targets.get(unit, {}).get("target_usd", 0) if unit in scope_units else 0),
            "target_krw": sum(targets.get(key, {}).get("target_krw", 0) for key in scope_units) if unit == "total" else (targets.get(unit, {}).get("target_krw", 0) if unit in scope_units else 0),
            "initial_fcst_usd": sum(row["amount_usd"] for row in initial_rows if unit == "total" or row["business_unit"] == unit),
            "initial_fcst_krw": sum(row["krw_amount"] for row in initial_rows if unit == "total" or row["business_unit"] == unit),
            "carryover_in_usd": sum(row["amount_usd"] for row in active if row["timing_type"] == "previous_carryover"),
            "carryover_in_krw": sum(row["krw_amount"] for row in active if row["timing_type"] == "previous_carryover"),
            "new_usd": sum(row["amount_usd"] for row in active if row["timing_type"] == "current_new"),
            "new_krw": sum(row["krw_amount"] for row in active if row["timing_type"] == "current_new"),
            "carryover_out_usd": sum(
                link["amount_usd"] for link in carryover_rows
                if unit == "total" or row_by_id.get(link["source_sales_id"], {}).get("business_unit") == unit
            ),
            "carryover_out_krw": sum(
                link["amount_usd"] * row_by_id.get(link["source_sales_id"], {}).get("applied_rate", 0)
                for link in carryover_rows
                if unit == "total" or row_by_id.get(link["source_sales_id"], {}).get("business_unit") == unit
            ),
            "count": len(active),
        }
        for status in SALES_STATUSES:
            values[f"{status}_usd"] = sum(row["amount_usd"] for row in active if row["sales_status"] == status)
            values[f"{status}_krw"] = sum(row["krw_amount"] for row in active if row["sales_status"] == status)
        values["high_confidence_usd"] = values["confirmed_usd"] + values["scheduled_usd"]
        values["current_fcst_usd"] = values["high_confidence_usd"] + values["pipeline_usd"]
        values["current_fcst_krw"] = values["confirmed_krw"] + values["scheduled_krw"] + values["pipeline_krw"]
        values["target_variance_usd"] = values["current_fcst_usd"] - values["target_usd"]
        values["target_variance_krw"] = values["current_fcst_krw"] - values["target_krw"]
        values["initial_variance_usd"] = values["current_fcst_usd"] - values["initial_fcst_usd"]
        values["initial_variance_krw"] = values["current_fcst_krw"] - values["initial_fcst_krw"]
        values["achievement_pct"] = round(values["current_fcst_usd"] / values["target_usd"] * 100, 1) if values["target_usd"] else None
        summary[unit] = values
    return summary


def live_status_payload(rows):
    """Build the three operational slices shown in TODAY SALES STATUS.

    Revenue totals only include active records.  The next-month amount starts as
    soon as a user marks a record for carryover, and continues to include the
    source record after the carryover has been processed.  Cancelled records are
    excluded from every metric.
    """
    scopes = {
        "total": rows,
        "current_new": [row for row in rows if row.get("timing_type") == "current_new"],
        "previous_carryover": [row for row in rows if row.get("timing_type") == "previous_carryover"],
    }
    payload = {}
    for scope, scoped_rows in scopes.items():
        active = [row for row in scoped_rows if row.get("record_status") == "active"]
        not_cancelled = [row for row in scoped_rows if row.get("record_status") != "cancelled"]
        values = {}
        for status in SALES_STATUSES:
            values[f"{status}_usd"] = sum(
                row.get("amount_usd", 0) for row in active if row.get("sales_status") == status
            )
            values[f"{status}_krw"] = sum(
                row.get("krw_amount", 0) for row in active if row.get("sales_status") == status
            )
        values["high_confidence_usd"] = values["confirmed_usd"] + values["scheduled_usd"]
        values["high_confidence_krw"] = values["confirmed_krw"] + values["scheduled_krw"]
        values["current_fcst_usd"] = values["high_confidence_usd"] + values["pipeline_usd"]
        values["current_fcst_krw"] = values["high_confidence_krw"] + values["pipeline_krw"]
        carryover_rows = [
            row for row in not_cancelled
            if row.get("carryover_decision") == "yes"
            or row.get("carryover_role") == "outflow"
            or row.get("record_status") == "carried_over"
        ]
        values["next_month_usd"] = sum(row.get("amount_usd", 0) for row in carryover_rows)
        values["next_month_krw"] = sum(row.get("krw_amount", 0) for row in carryover_rows)
        payload[scope] = values
    return payload


def _local_date_from_timestamp(value):
    if not value:
        return None
    try:
        parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        if parsed.tzinfo is None:
            parsed = parsed.replace(tzinfo=timezone.utc)
        return parsed.astimezone(SEOUL).date().isoformat()
    except (TypeError, ValueError):
        return str(value)[:10] if len(str(value)) >= 10 else None


def round_cards_payload(db, target_month, user, args, settings, displayed_summary=None):
    today = datetime.now(SEOUL).date()
    selected_business_unit = str(args.get("business_unit", "")).strip()
    configured = {
        key: settings.get(column) for key, _label, column in ROUND_DEFINITIONS
    }
    if not configured["initial"]:
        submitted = db.execute(
            "SELECT MAX(submitted_at) AS submitted_at FROM monthly_fcst_snapshots WHERE target_month=?",
            (target_month,),
        ).fetchone()
        configured["initial"] = _local_date_from_timestamp(submitted["submitted_at"] if submitted else None)
    cards = []
    for key, label, _column in ROUND_DEFINITIONS:
        cutoff_date = configured.get(key)
        available = False
        total = {}
        comparison = None
        if cutoff_date:
            try:
                cutoff_day = datetime.strptime(cutoff_date, "%Y-%m-%d").date()
            except ValueError:
                cutoff_day = None
            available = bool(cutoff_day and cutoff_day <= today)
            if available:
                rows = historical_sales(db, target_month, cutoff_date, user, round_key=key)
                rows = filter_rows(rows, args)
                comparison = live_status_payload(rows)
                total = summary_payload(
                    db, rows, target_month, cutoff_date, selected_business_unit
                ).get("total", {})
        if key in {"initial", "round1"}:
            amount_key = "current_fcst"
            formula = "확정+예정+추진"
        elif key in {"round2", "round3"}:
            amount_key = "high_confidence"
            formula = "확정+예정"
        else:
            amount_key = "confirmed"
            formula = "확정"
        if key == "initial" and displayed_summary is not None:
            amount_key = "initial_fcst"
            formula = "최초 제출값"
            total = displayed_summary
        card = {
            "key": key,
            "label": label,
            "cutoff_date": cutoff_date,
            "available": available,
            "formula": formula,
            "amount_usd": total.get(f"{amount_key}_usd") if available else None,
            "amount_krw": (
                total.get("confirmed_krw", 0) + total.get("scheduled_krw", 0)
                if available and amount_key == "high_confidence"
                else total.get(f"{amount_key}_krw") if available else None
            ),
            "comparison": comparison,
        }
        if key == "round2" and available:
            card.update({
                "secondary_label": "추진 포함 시",
                "secondary_usd": total.get("current_fcst_usd", 0),
                "secondary_krw": total.get("current_fcst_krw", 0),
            })
        cards.append(card)
    return cards


def filter_rows(rows, args):
    status = str(args.get("status", "")).strip()
    business_unit = str(args.get("business_unit", "")).strip()
    owner = str(args.get("owner", "")).strip()
    country = str(args.get("country", "")).strip()
    customer = str(args.get("customer", "")).strip()
    original_month = str(args.get("original_month", "")).strip()
    timing = str(args.get("timing_type", "")).strip()
    customer_history = str(args.get("customer_history", "")).strip()
    split_only = str(args.get("split", "")).strip()
    carryover = str(args.get("carryover", "")).strip()
    risk = str(args.get("risk", "")).strip()
    query = str(args.get("search", "")).strip().lower()
    output = []
    for row in rows:
        if status and row["sales_status"] != status:
            continue
        if business_unit and row["business_unit"] != business_unit:
            continue
        if owner and row["owner_name"] != owner:
            continue
        if country and row["country"] != country:
            continue
        if customer and row["customer_name"] != customer:
            continue
        if original_month and row["original_month"] != original_month:
            continue
        if timing and row["timing_type"] != timing:
            continue
        if customer_history and row.get("customer_history", "existing") != customer_history:
            continue
        if split_only == "yes" and row["split_role"] == "none":
            continue
        if split_only == "no" and row["split_role"] != "none":
            continue
        carryover_decision = row.get("carryover_decision") or (
            "yes" if row.get("carryover_role") == "outflow" or row.get("record_status") == "carried_over" else "no"
        )
        if carryover == "yes" and carryover_decision != "yes":
            continue
        if carryover == "no" and carryover_decision != "no":
            continue
        if risk == "yes" and not row.get("warnings"):
            continue
        if query and query not in " ".join((row["sales_no"], row["customer_name"], row["country"], row["owner_name"], row["notes"])).lower():
            continue
        output.append(row)
    return output


def filter_options(db, rows, user):
    candidates_sql = """
        SELECT customer_name, country, owner_name
        FROM monthly_sales_master_candidates
        WHERE status != 'rejected'
    """
    candidates = db.execute(candidates_sql).fetchall()
    return {
        "owners": sorted(
            set(MONTHLY_SALES_OWNERS)
            | {row["owner_name"] for row in rows if row["owner_name"]}
            | {row["owner_name"] for row in candidates if row["owner_name"]}
        ),
        "countries": sorted(
            {row["country"] for row in rows if row["country"]}
            | {row["country"] for row in candidates if row["country"]}
        ),
        "customers": sorted(
            {row["customer_name"] for row in rows if row["customer_name"]}
            | {row["customer_name"] for row in candidates if row["customer_name"]}
        ),
        "original_months": sorted({row["original_month"] for row in rows if row["original_month"]}, reverse=True),
    }


def sheet_sync_state(db):
    state = row_dict(db.execute("SELECT * FROM monthly_sales_sheet_sync WHERE id = 1").fetchone()) or {}
    return {**sheets_status(), **state}


def rows_for_sheets(db):
    users = {row["id"]: row["display_name"] for row in db.execute("SELECT id, display_name FROM users").fetchall()}
    output = {name: [] for name in SHEET_SCHEMAS}
    sales = db.execute("SELECT * FROM monthly_sales ORDER BY target_month, sales_no").fetchall()
    for row in sales:
        item = row_dict(row)
        item["sales_id"] = item.pop("id")
        item["carryover_status"] = "차월 진행" if item.pop("carryover_decision", "no") == "yes" else "당월 유지"
        item["created_by"] = users.get(item.get("created_by"), item.get("created_by"))
        item["updated_by"] = users.get(item.get("updated_by"), item.get("updated_by"))
        output["SALES"].append(item)
        milestones = milestone_dict(db, item["sales_id"])
        output["MILESTONES"].append({"sales_id": item["sales_id"], "sales_no": item["sales_no"], **milestones, "updated_at": row["updated_at"]})
    for row in db.execute("SELECT * FROM monthly_sales_split_links ORDER BY created_at, id").fetchall():
        item = row_dict(row)
        item["created_by"] = users.get(item.get("created_by"), item.get("created_by"))
        item["cancelled_by"] = users.get(item.get("cancelled_by"), item.get("cancelled_by"))
        output["SPLIT_LINKS"].append(item)
    for row in db.execute("SELECT * FROM monthly_sales_carryovers ORDER BY created_at, id").fetchall():
        item = row_dict(row)
        item["created_by"] = users.get(item.get("created_by"), item.get("created_by"))
        item["cancelled_by"] = users.get(item.get("cancelled_by"), item.get("cancelled_by"))
        output["CARRYOVER_HISTORY"].append(item)
    for row in db.execute("SELECT * FROM monthly_fcst_snapshots ORDER BY target_month, sales_no").fetchall():
        item = row_dict(row)
        item["submitted_by"] = users.get(item.get("submitted_by"), item.get("submitted_by"))
        output["FORECAST"].append(item)
    for row in db.execute("SELECT * FROM monthly_sales_history ORDER BY id").fetchall():
        output["HISTORY"].append(row_dict(row))
    for row in db.execute("SELECT * FROM monthly_sales_rates ORDER BY target_month, currency").fetchall():
        item = row_dict(row)
        output["RATES"].append({
            "rate_scope": "plan", "target_month": item["target_month"], "rate_base_date": "",
            "currency": item["currency"], "rate": item["plan_rate"], "source": item["source"],
            "uploaded_file": "", "updated_by": users.get(item.get("updated_by"), item.get("updated_by")),
            "updated_at": item["updated_at"], "correction_reason": item["correction_reason"],
        })
    for row in db.execute("SELECT * FROM monthly_sales_daily_rates ORDER BY rate_base_date, currency").fetchall():
        item = row_dict(row)
        output["RATES"].append({
            "rate_scope": "actual", "target_month": item["rate_base_date"][:7],
            "rate_base_date": item["rate_base_date"], "currency": item["currency"], "rate": item["rate"],
            "source": item["source"], "uploaded_file": item["uploaded_file"],
            "updated_by": users.get(item.get("updated_by"), item.get("updated_by")),
            "updated_at": item["updated_at"], "correction_reason": item["correction_reason"],
        })
    for row in db.execute("SELECT * FROM monthly_sales_targets ORDER BY target_month, business_unit").fetchall():
        item = row_dict(row)
        item["updated_by"] = users.get(item.get("updated_by"), item.get("updated_by"))
        output["TARGETS"].append(item)
    output["MASTER_CANDIDATES"] = [row_dict(row) for row in db.execute("SELECT * FROM monthly_sales_master_candidates ORDER BY customer_name").fetchall()]
    for row in db.execute("SELECT * FROM monthly_sales_month_settings ORDER BY target_month").fetchall():
        item = row_dict(row)
        item["closed_by"] = users.get(item.get("closed_by"), item.get("closed_by"))
        item["reopened_by"] = users.get(item.get("reopened_by"), item.get("reopened_by"))
        output["SETTINGS"].append(item)
    return output


def sync_sheets(db, now):
    status = sheets_status()
    if not status["configured"]:
        return {"status": "unconfigured", "message": "Google Sheets 연동정보를 설정하면 자동 동기화됩니다."}
    timestamp = now()
    db.execute("UPDATE monthly_sales_sheet_sync SET last_attempt_at = ?, updated_at = ? WHERE id = 1", (timestamp, timestamp))
    db.commit()
    try:
        result = GoogleSheetsMirror().replace_all(rows_for_sheets(db))
    except Exception as exc:
        message = str(exc)[:500]
        db.execute("UPDATE monthly_sales_sheet_sync SET dirty = 1, last_error = ?, updated_at = ? WHERE id = 1", (message, now()))
        db.commit()
        return {"status": "failed", "message": message}
    db.execute(
        "UPDATE monthly_sales_sheet_sync SET dirty = 0, last_success_at = ?, last_error = NULL, updated_at = ? WHERE id = 1",
        (result["synced_at"], now()),
    )
    db.commit()
    return {"status": "success", **result}


def duplicate_pi(db, pi_no, exclude_sales_id=None):
    if not pi_no:
        return None
    sql = """
        SELECT s.sales_no, s.customer_name FROM monthly_sales_milestones m
        JOIN monthly_sales s ON s.id = m.sales_id
        WHERE LOWER(m.pi_no) = LOWER(?) AND s.record_status = 'active'
    """
    params = [pi_no]
    if exclude_sales_id:
        sql += " AND s.id != ?"
        params.append(exclude_sales_id)
    return db.execute(sql, params).fetchone()


def sale_duplicate_signature(values):
    return (
        values["customer_name"].strip().lower(), values["country"].strip().lower(),
        values["owner_name"].strip().lower(), values["business_unit"], values["target_month"],
        values["transaction_currency"], round(values["transaction_amount"], 4),
    )


def duplicate_sale(db, values):
    return db.execute(
        """SELECT sales_no FROM monthly_sales
           WHERE LOWER(customer_name)=LOWER(?) AND LOWER(country)=LOWER(?) AND LOWER(owner_name)=LOWER(?)
             AND business_unit=? AND target_month=? AND transaction_currency=?
             AND ABS(transaction_amount-?)<0.0001 AND record_status='active'
           LIMIT 1""",
        (values["customer_name"], values["country"], values["owner_name"], values["business_unit"],
         values["target_month"], values["transaction_currency"], values["transaction_amount"]),
    ).fetchone()


def strict_round_row_link(db, values, supplied_sales_no=""):
    """Link only on an exact, unique identifier or exact row signature."""
    if supplied_sales_no:
        matches = db.execute(
            "SELECT id FROM monthly_sales WHERE sales_no=? AND target_month=?",
            (supplied_sales_no, values["target_month"]),
        ).fetchall()
        if len(matches) == 1:
            return matches[0]["id"]
    pi_no = (values.get("milestones", {}).get("pi_no") or "").strip()
    if pi_no:
        matches = db.execute(
            """SELECT s.id FROM monthly_sales s
               JOIN monthly_sales_milestones m ON m.sales_id=s.id
               WHERE s.target_month=? AND LOWER(m.pi_no)=LOWER(?)""",
            (values["target_month"], pi_no),
        ).fetchall()
        if len(matches) == 1:
            return matches[0]["id"]
    matches = db.execute(
        """SELECT id FROM monthly_sales
           WHERE target_month=? AND business_unit=? AND transaction_currency=?
             AND ABS(transaction_amount-?)<0.0001
             AND LOWER(country)=LOWER(?)
             AND (LOWER(customer_name)=LOWER(?) OR LOWER(COALESCE(original_customer_name,''))=LOWER(?))""",
        (
            values["target_month"], values["business_unit"], values["transaction_currency"],
            values["transaction_amount"], values["country"], values["customer_name"],
            values.get("original_customer_name") or values["customer_name"],
        ),
    ).fetchall()
    return matches[0]["id"] if len(matches) == 1 else None


def parse_round_import_sheet(db, sheet, target_month, cutoff_date, round_key, actor, now):
    header_alias = {
        "관리번호": "sales_no", "거래처명": "customer_name", "국가": "country",
        "담당자": "owner_name", "사업분야": "business_unit", "신규/기존": "customer_history",
        "최초 추진월": "original_month", "매출 대상월": "target_month", "대상월": "target_month",
        "진행 시점": "timing_type", "매출 추진시점": "timing_type", "유형": "timing_type",
        "환종": "transaction_currency", "거래금액": "transaction_amount",
        "매출액(USD)": "transaction_amount", "진행상태": "sales_status", "PI 번호": "pi_no",
        "차월 구분": "carryover_decision", "차월 여부": "carryover_decision",
        "차월 결정일": "carryover_decided_at", "차월 대상월": "carryover_target_month",
        "차월 사유": "carryover_decision_reason", "오더합의일": "order_agreed_at",
        "PO 수신일": "po_received_at", "PI 발송일": "pi_sent_at", "PI 컨펌일": "pi_confirmed_at",
        "예상 수금일": "payment_expected_at", "실제 수금일": "payment_actual_at",
        "예상 출고일": "shipment_expected_at", "실제 출고일": "shipment_actual_at",
        "예상 선적일": "shipping_expected_at", "실제 선적일": "shipping_actual_at", "비고": "notes",
    }
    business_map = {"에스테틱": "aesthetic", "메디컬": "medical", "메디칼": "medical", "덴탈": "dental"}
    status_map = {"확정": "confirmed", "예정": "scheduled", "추진": "pipeline", "미정": "undecided"}
    timing_map = {
        "당월 추진": "current_new", "당월 신규": "current_new",
        "전월 이월": "previous_carryover", "이월 유입": "previous_carryover",
    }
    customer_history_map = {"신규 거래처": "new", "신규": "new", "기존 거래처": "existing", "기존": "existing"}
    carryover_map = {
        "차월 진행": "yes", "차월": "yes", "예": "yes", "Y": "yes",
        "당월 유지": "no", "차월 아님": "no", "아니오": "no", "N": "no",
    }
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return [], [{"row": 1, "message": "Excel 파일에 데이터가 없습니다."}]
    headers = [str(value or "").strip() for value in rows[0]]
    required = {"거래처명", "국가", "담당자", "사업분야", "최초 추진월", "진행상태"}
    if "매출 대상월" not in headers and "대상월" not in headers:
        required.add("매출 대상월")
    missing = sorted(required - set(headers))
    if "거래금액" not in headers and "매출액(USD)" not in headers:
        missing.append("거래금액")
    if missing:
        return [], [{"row": 1, "message": f"필수 열이 없습니다: {', '.join(missing)}"}]
    parsed = []
    errors = []
    round_marker = {"initial": "I", "round1": "R1", "round2": "R2", "round3": "R3", "final": "F"}[round_key]
    effective_timestamp = cutoff_utc(cutoff_date)
    for row_number, row_values in enumerate(rows[1:], 2):
        if not any(value not in (None, "") for value in row_values):
            continue
        source = {
            header_alias.get(headers[index], headers[index]): value
            for index, value in enumerate(row_values) if index < len(headers)
        }
        supplied_sales_no = clean_text(source.pop("sales_no", ""), "관리번호", 100)
        source["business_unit"] = business_map.get(str(source.get("business_unit") or "").strip(), str(source.get("business_unit") or "").strip())
        source["sales_status"] = status_map.get(str(source.get("sales_status") or "").strip(), str(source.get("sales_status") or "").strip())
        source["customer_history"] = customer_history_map.get(
            str(source.get("customer_history") or "기존 거래처").strip(),
            str(source.get("customer_history") or "existing").strip(),
        )
        source["original_month"] = excel_month_value(source.get("original_month"))
        source["target_month"] = excel_month_value(source.get("target_month") or target_month)
        source["timing_type"] = inferred_timing_type(
            timing_map.get(str(source.get("timing_type") or "").strip(), source.get("timing_type")),
            source["original_month"], source["target_month"],
        )
        source["transaction_currency"] = str(source.get("transaction_currency") or "USD").strip().upper()
        raw_carryover = str(source.get("carryover_decision") or "당월 유지").strip()
        source["carryover_decision"] = carryover_map.get(raw_carryover, raw_carryover.lower())
        if isinstance(source.get("carryover_decided_at"), (datetime, date)):
            source["carryover_decided_at"] = source["carryover_decided_at"].strftime("%Y-%m-%d")
        source["carryover_target_month"] = excel_month_value(source.get("carryover_target_month"))
        milestones = {}
        for field in MILESTONE_FIELDS:
            value = source.pop(field, None)
            if isinstance(value, (datetime, date)):
                value = value.strftime("%Y-%m-%d")
            milestones[field] = value
        source["milestones"] = milestones
        source["apply_status_suggestion"] = False
        try:
            if source["target_month"] != target_month:
                raise ValueError(f"선택한 대상월 {target_month}과 Excel 매출 대상월이 다릅니다.")
            owner_name = str(source.get("owner_name") or "").strip()
            if owner_name not in MONTHLY_SALES_OWNERS:
                raise ValueError(f"담당자는 {', '.join(MONTHLY_SALES_OWNERS)} 중에서 선택하세요.")
            normalized = normalize_sale_input(
                source, db, now, actor=actor, allow_legacy_confirmation=True
            )
            linked_sales_id = strict_round_row_link(db, normalized, supplied_sales_no)
            archive_id = uuid.uuid4().hex
            sales_no = supplied_sales_no or f"HIST-{target_month[2:4]}{target_month[5:7]}-{round_marker}-{row_number - 1:04d}"
            snapshot = {
                "id": archive_id, "sales_no": sales_no,
                "customer_name": normalized["customer_name"], "original_customer_name": normalized["original_customer_name"],
                "customer_master_id": normalized["customer_master_id"], "customer_link_status": normalized["customer_link_status"],
                "country": normalized["country"], "owner_name": normalized["owner_name"],
                "owner_user_id": normalized["owner_user_id"],
                "business_unit": normalized["business_unit"], "currency": "USD",
                "transaction_currency": normalized["transaction_currency"],
                "transaction_currency_standard": normalized["transaction_currency"],
                "transaction_amount": normalized["transaction_amount"], "customer_history": normalized["customer_history"],
                "original_month": normalized["original_month"], "target_month": target_month,
                "timing_type": normalized["timing_type"], "amount_usd": normalized["amount_usd"],
                "plan_rate": normalized["plan_rate"], "plan_usd_rate": normalized["plan_usd_rate"],
                "actual_rate": normalized["actual_rate"],
                "applied_rate": normalized["applied_rate"], "rate_type": normalized["rate_type"],
                "rate_base_date": normalized["rate_base_date"], "rate_source": normalized["rate_source"] or "과거 차수자료 이관",
                "krw_amount": normalized["krw_amount"], "sales_status": normalized["sales_status"],
                "confirmation_basis": normalized["confirmation_basis"],
                "confirmation_note": normalized["confirmation_note"],
                "record_status": "active", "split_role": "none", "carryover_role": "none",
                "carryover_decision": normalized["carryover_decision"],
                "carryover_decided_at": normalized["carryover_decided_at"],
                "carryover_target_month": normalized["carryover_target_month"],
                "carryover_decision_reason": normalized["carryover_decision_reason"],
                "original_sales_id": None, "split_group_id": None, "notes": normalized["notes"],
                "version": 1, "created_by": actor["id"], "updated_by": actor["id"],
                "created_at": effective_timestamp, "updated_at": effective_timestamp,
                "milestones": normalized["milestones"],
            }
            parsed.append({
                "source_row": row_number, "linked_sales_id": linked_sales_id,
                "snapshot": snapshot,
            })
        except ValueError as exc:
            errors.append({"row": row_number, "message": str(exc)})
    if not parsed and not errors:
        errors.append({"row": 2, "message": "이관할 매출 행이 없습니다."})
    return parsed, errors


def apply_available_actual_rates(db, actor, now):
    pending = db.execute(
        """SELECT s.*, m.shipment_actual_at
           FROM monthly_sales s JOIN monthly_sales_milestones m ON m.sales_id=s.id
           WHERE s.record_status='active' AND s.rate_type='pending'
             AND m.shipment_actual_at IS NOT NULL"""
    ).fetchall()
    applied = []
    timestamp = now()
    for sale in pending:
        currency = (
            sale["preliminary_actual_currency"]
            if sale["preliminary_actual_currency"] else canonical_sale_currency(sale)
        )
        rate = daily_rate_on_or_before(db, sale["shipment_actual_at"], currency)
        if not rate:
            continue
        preliminary_amount = sale["preliminary_actual_amount"] or sale["transaction_amount"]
        preliminary_krw = round(
            preliminary_amount if currency == "KRW" else preliminary_amount * float(rate["rate"]), 2
        )
        db.execute(
            """UPDATE monthly_sales SET actual_rate=?, rate_type='shipment',
               rate_base_date=?, rate_source=?, preliminary_actual_currency=?,
               preliminary_actual_amount=?,preliminary_actual_krw_amount=?,version=version+1,
               updated_by=?, updated_at=? WHERE id=?""",
            (rate["rate"], rate["rate_base_date"], rate["source"], currency,
             preliminary_amount, preliminary_krw, actor["id"], timestamp, sale["id"]),
        )
        add_history(
            db, sale["id"], "PRELIMINARY_RATE_APPLY", "운영 환율 Cache 재조회 적용", actor, now,
            effective_at=sale["shipment_actual_at"], effective_at_source="shipment_business_date",
        )
        applied.append(sale["id"])
    return applied


def monthly_customer_link_groups(db):
    rows = db.execute(
        """SELECT id,customer_name,original_customer_name,country,owner_name,business_unit,target_month,
                  amount_usd,customer_master_id,customer_link_status
           FROM monthly_sales ORDER BY customer_name,country,target_month"""
    ).fetchall()
    grouped = {}
    for row in rows:
        item = row_dict(row)
        source_name = item.get("original_customer_name") or item["customer_name"]
        key = customer_link_group_key(source_name, item["country"])
        group = grouped.setdefault(key, {
            "group_key": key, "customer_name": source_name, "country": item["country"],
            "sales_ids": [], "months": set(), "owners": set(), "business_units": set(),
            "amount_usd": 0.0, "linked_ids": set(), "statuses": set(),
        })
        group["sales_ids"].append(item["id"])
        group["months"].add(item["target_month"])
        group["owners"].add(item["owner_name"])
        group["business_units"].add(item["business_unit"])
        group["amount_usd"] += float(item["amount_usd"] or 0)
        if item.get("customer_master_id"):
            group["linked_ids"].add(item["customer_master_id"])
        group["statuses"].add(item.get("customer_link_status") or "unlinked")
    result = []
    for group in grouped.values():
        strict, review = customer_master_candidates(db, group["customer_name"], group["country"])
        linked_ids = list(group.pop("linked_ids"))
        statuses = group.pop("statuses")
        group.update({
            "sales_count": len(group.pop("sales_ids")),
            "months": sorted(group["months"]), "owners": sorted(group["owners"]),
            "business_units": sorted(group["business_units"]),
            "amount_usd": round(group["amount_usd"], 2),
            "customer_master_id": linked_ids[0] if len(linked_ids) == 1 else None,
            "link_status": (
                "mixed" if len(linked_ids) > 1 else
                "linked" if len(linked_ids) == 1 else
                "excluded" if statuses == {"excluded"} else
                "auto_ready" if len(strict) == 1 else
                "review" if review else "unmatched"
            ),
            "exact_candidates": strict,
            "review_candidates": review[:8],
        })
        group["months"] = sorted(group["months"])
        group["owners"] = sorted(group["owners"])
        group["business_units"] = sorted(group["business_units"])
        result.append(group)
    result.sort(key=lambda item: ({"review": 0, "mixed": 0, "unmatched": 1, "auto_ready": 2, "linked": 3, "excluded": 4}.get(item["link_status"], 9), item["customer_name"].casefold()))
    return result


def monthly_customer_link_summary(groups):
    summary = {"total": len(groups), "linked": 0, "auto_ready": 0, "review": 0, "unmatched": 0, "excluded": 0}
    for group in groups:
        key = group["link_status"] if group["link_status"] in summary else "review"
        summary[key] += 1
    summary["needs_action"] = summary["auto_ready"] + summary["review"] + summary["unmatched"]
    return summary


def register_monthly_sales_fcst(app, get_db, role_required, csrf_required, now, audit, rollback_quietly, sqlite_is_busy, country_catalog=None):
    def requested_commercial_context(db, data, values, existing=None):
        """Resolve commercial terms only when a new source selection is made.

        Existing FCST snapshots are intentionally not re-resolved during an
        unrelated edit: later Master changes must never rewrite past FCST terms.
        Legacy rows without an Item remain editable, while every new row must
        enter through the commercial resolver.
        """
        item_was_sent = "interim_product_code" in data
        promotion_was_sent = "promotion_id" in data or "source_promotion_id" in data
        detach_promotion = data.get("detach_promotion") in (True, 1, "1", "true", "True")
        existing_item = (
            existing["interim_product_code"]
            if existing is not None and "interim_product_code" in existing.keys()
            else None
        )
        item_code = clean_text(
            data.get("interim_product_code") if item_was_sent else existing_item,
            "아이템", 80,
        )
        if existing is None and not item_code:
            customer = db.execute(
                "SELECT contract_policy FROM customer_master WHERE id=?",
                (values.get("customer_master_id"),),
            ).fetchone()
            # Unclassified customers are the explicit legacy-review path.  Once
            # a commercial policy is reviewed, no API caller can bypass the
            # Item-based resolver by omitting the field.
            if customer and customer["contract_policy"] in {"contract_required", "no_contract_allowed"}:
                raise ValueError("신규 매출은 거래처와 아이템을 선택해 상업조건을 확인해야 합니다.")
            return None, None
        if not item_code:
            return None, None

        existing_promotion = (
            existing["source_promotion_id"]
            if existing is not None and "source_promotion_id" in existing.keys()
            else None
        )
        promotion_id = data.get("promotion_id", data.get("source_promotion_id")) if promotion_was_sent else existing_promotion
        promotion_id = clean_text(promotion_id, "Promotion", 80)
        if detach_promotion:
            promotion_id = None
            promotion_was_sent = True
        if existing is not None and not item_was_sent and not promotion_was_sent:
            return None, None
        if (
            existing is not None
            and item_code == existing_item
            and (promotion_id or None) == (existing_promotion or None)
            and not detach_promotion
        ):
            # An ordinary edit must retain the original sale-time snapshot.
            # Master changes are picked up only by an explicit source change.
            return None, None
        if (
            existing is not None
            and int(existing["commercial_terms_managed"] or 0)
            and promotion_was_sent
            and not detach_promotion
            and promotion_id != existing_promotion
        ):
            raise ValueError("Promotion 관리 조건은 Promotion에서 수정하세요. 일반 FCST로 전환하려면 연결 해제를 선택하세요.")

        order_date = clean_text(data.get("order_date"), "주문 기준일", 10)
        if not order_date:
            order_date = f"{values['target_month']}-01"
        context = resolve_customer_order_context(
            db,
            values.get("customer_master_id"),
            item_code,
            order_date,
            promotion_id=promotion_id or None,
            country_catalog=country_catalog,
        )
        return context, ("promotion" if promotion_id else "regular")

    """Register the module on the existing Flask application."""

    def operation_error(exc):
        if isinstance(exc, ValueError):
            return jsonify(error=str(exc)), 400
        if isinstance(exc, PermissionError):
            return jsonify(error=str(exc)), 409
        if sqlite_is_busy(exc):
            return jsonify(error="다른 사용자가 월별 매출을 저장 중입니다. 잠시 후 다시 시도하세요.", code="DATABASE_BUSY"), 409
        raise exc

    def maybe_sync(db):
        mark_sheet_dirty(db, now)
        db.commit()
        return sync_sheets(db, now)

    def group_rows(db, group_key):
        rows = db.execute(
            "SELECT * FROM monthly_sales ORDER BY target_month,id"
        ).fetchall()
        return [row for row in rows if customer_link_group_key(
            row["original_customer_name"] or row["customer_name"], row["country"]
        ) == group_key]

    def save_customer_alias(db, customer_id, name, country, source):
        name_key, country_key = customer_link_key(name), customer_link_key(country)
        existing = db.execute(
            "SELECT customer_master_id FROM monthly_sales_customer_aliases WHERE normalized_name=? AND normalized_country=?",
            (name_key, country_key),
        ).fetchone()
        if existing and existing["customer_master_id"] != customer_id:
            raise ValueError("이 거래처명·국가는 이미 다른 거래처에 연결되어 있습니다.")
        db.execute(
            """INSERT INTO monthly_sales_customer_aliases
               (id,customer_master_id,alias_name,alias_country,normalized_name,normalized_country,source,created_by,created_at)
               VALUES (?,?,?,?,?,?,?,?,?)
               ON CONFLICT(normalized_name,normalized_country) DO NOTHING""",
            (uuid.uuid4().hex, customer_id, name, country, name_key, country_key, source,
             g.current_user["id"], now()),
        )

    def link_group(db, rows, customer, link_status, reason):
        timestamp = now()
        for row in rows:
            before = sale_snapshot(db, row["id"])
            original_name = row["original_customer_name"] or row["customer_name"]
            db.execute(
                """UPDATE monthly_sales
                   SET customer_master_id=?,customer_name=?,original_customer_name=?,customer_link_status=?,
                       customer_linked_at=?,customer_linked_by=?,version=version+1,updated_by=?,updated_at=?
                   WHERE id=?""",
                (customer["id"], customer["display_name"], original_name, link_status, timestamp,
                 g.current_user["id"], g.current_user["id"], timestamp, row["id"]),
            )
            add_history(db, row["id"], "CUSTOMER_LINK", reason, g.current_user, now)
            audit("MONTHLY_SALE_CUSTOMER_LINK", "monthly_sale", row["id"],
                  f"{row['sales_no']} 거래처 마스터 연결", before, sale_snapshot(db, row["id"]), connection=db)
        save_customer_alias(db, customer["id"], rows[0]["original_customer_name"] or rows[0]["customer_name"], rows[0]["country"], "auto" if link_status == "auto_linked" else "manual")

    def next_fcst_temp_id(db):
        year = datetime.now(SEOUL).year
        prefix = f"TEMP-{year}-"
        sequences = []
        for row in db.execute("SELECT customer_id FROM customer_master WHERE customer_id LIKE ?", (f"{prefix}%",)).fetchall():
            match = re.fullmatch(re.escape(prefix) + r"(\d+)", row["customer_id"])
            if match:
                sequences.append(int(match.group(1)))
        return f"{prefix}{max(sequences, default=0) + 1:04d}"

    @app.get("/api/monthly-sales-fcst/customer-links")
    @role_required("admin", "manager", "editor", "viewer")
    def monthly_sales_customer_links():
        groups = monthly_customer_link_groups(get_db())
        return jsonify(groups=groups, summary=monthly_customer_link_summary(groups))

    @app.post("/api/monthly-sales-fcst/customer-links/auto-match")
    @role_required("admin", "manager", "editor")
    @csrf_required
    def auto_match_monthly_sales_customers():
        db = get_db(); linked_groups = 0; linked_sales = 0
        try:
            db.execute("BEGIN IMMEDIATE")
            for group in monthly_customer_link_groups(db):
                if group["link_status"] != "auto_ready" or len(group["exact_candidates"]) != 1:
                    continue
                rows = group_rows(db, group["group_key"])
                customer = db.execute("SELECT * FROM customer_master WHERE id=?", (group["exact_candidates"][0]["customer_master_id"],)).fetchone()
                link_group(db, rows, customer, "auto_linked", "거래처명·국가 1:1 일치 자동연결")
                linked_groups += 1; linked_sales += len(rows)
            audit("MONTHLY_SALES_CUSTOMER_AUTO_LINK", "monthly_sale", None,
                  f"기존 FCST 거래처 자동연결 {linked_groups}개 그룹·{linked_sales}건", None,
                  {"groups": linked_groups, "sales": linked_sales}, connection=db)
            db.commit()
        except (ValueError, sqlite3.Error) as exc:
            rollback_quietly(db)
            return operation_error(exc)
        groups = monthly_customer_link_groups(db)
        return jsonify(message=f"{linked_groups}개 거래처 그룹의 FCST {linked_sales}건을 자동 연결했습니다.", summary=monthly_customer_link_summary(groups))

    @app.post("/api/monthly-sales-fcst/customer-links/resolve")
    @role_required("admin", "manager", "editor")
    @csrf_required
    def resolve_monthly_sales_customer():
        data = request.get_json(silent=True) or {}
        group_key = clean_text(data.get("group_key"), "연결 그룹", 80, True)
        action = str(data.get("action") or "link").strip()
        if action not in {"link", "create_temp", "exclude"}:
            return jsonify(error="거래처 연결 처리방식이 올바르지 않습니다."), 400
        db = get_db()
        try:
            db.execute("BEGIN IMMEDIATE")
            rows = group_rows(db, group_key)
            if not rows:
                raise ValueError("연결할 기존 FCST 거래처를 찾을 수 없습니다.")
            if action == "exclude":
                timestamp = now()
                for row in rows:
                    before = sale_snapshot(db, row["id"])
                    db.execute(
                        """UPDATE monthly_sales SET customer_master_id=NULL,customer_link_status='excluded',
                           customer_linked_at=?,customer_linked_by=?,version=version+1,updated_by=?,updated_at=? WHERE id=?""",
                        (timestamp, g.current_user["id"], g.current_user["id"], timestamp, row["id"]),
                    )
                    add_history(db, row["id"], "CUSTOMER_LINK_EXCLUDED", "거래처 마스터 연결 제외", g.current_user, now)
                    audit("MONTHLY_SALE_CUSTOMER_EXCLUDE", "monthly_sale", row["id"],
                          f"{row['sales_no']} 거래처 연결 제외", before, sale_snapshot(db, row["id"]), connection=db)
                message = f"FCST {len(rows)}건을 거래처 연결 제외로 처리했습니다."
            else:
                if action == "link":
                    customer_id = clean_text(data.get("customer_master_id"), "거래처", 80, True)
                    customer = db.execute("SELECT * FROM customer_master WHERE id=?", (customer_id,)).fetchone()
                    if not customer:
                        raise ValueError("연결할 거래처 마스터를 찾을 수 없습니다.")
                    link_status = "linked"
                else:
                    timestamp = now(); customer_id = str(uuid.uuid4())
                    public_id = next_fcst_temp_id(db)
                    first = rows[0]
                    owner = db.execute("SELECT id FROM users WHERE display_name=? AND status='active' AND deleted_at IS NULL ORDER BY id LIMIT 1", (first["owner_name"],)).fetchone()
                    db.execute(
                        """INSERT INTO customer_master
                           (id,customer_id,erp_partner_code,erp_original_name,display_name,headquarters_country,
                            sales_region,source_type,relationship_type,customer_role,primary_owner_id,lifecycle_status,
                            default_currency,created_by,updated_by,created_at,updated_at)
                           VALUES (?,?,NULL,'',?,?,'미분류','temporary','dealer','Distributor',?,'active','USD',?,?,?,?)""",
                        (customer_id, public_id, first["original_customer_name"] or first["customer_name"], first["country"],
                         owner["id"] if owner else None, g.current_user["id"], g.current_user["id"], timestamp, timestamp),
                    )
                    for business_unit in sorted({row["business_unit"] for row in rows}):
                        db.execute("INSERT OR IGNORE INTO customer_business_areas(customer_id,business_unit,created_at) VALUES (?,?,?)", (customer_id, business_unit, timestamp))
                    customer = db.execute("SELECT * FROM customer_master WHERE id=?", (customer_id,)).fetchone()
                    audit("CUSTOMER_CREATE_FROM_FCST", "customer_master", customer_id,
                          f"{public_id} 기존 FCST 거래처 TEMP 생성", None, row_dict(customer), connection=db)
                    link_status = "temp_created"
                link_group(db, rows, customer, link_status, "기존 FCST 거래처 연결 확정")
                message = f"{customer['display_name']}에 기존 FCST {len(rows)}건을 연결했습니다."
            db.commit()
        except (ValueError, sqlite3.Error) as exc:
            rollback_quietly(db)
            return operation_error(exc)
        groups = monthly_customer_link_groups(db)
        return jsonify(message=message, summary=monthly_customer_link_summary(groups))

    @app.get("/api/monthly-sales-fcst")
    @role_required("admin", "manager", "editor", "viewer")
    def get_monthly_sales_fcst():
        target_month = str(request.args.get("month", "")).strip()
        if not valid_month(target_month):
            return jsonify(error="조회월을 YYYY-MM 형식으로 선택하세요."), 400
        as_of = str(request.args.get("as_of", "")).strip()
        as_of_round = str(request.args.get("as_of_round", "")).strip()
        compare_date = str(request.args.get("compare_date", "")).strip()
        if as_of_round and as_of_round not in {key for key, _label, _column in ROUND_DEFINITIONS}:
            return jsonify(error="기준 차수 값이 올바르지 않습니다."), 400
        try:
            if as_of:
                valid_date(as_of)
            if compare_date:
                valid_date(compare_date)
        except ValueError as exc:
            return jsonify(error=str(exc)), 400
        db = get_db()
        customer_link_summary = monthly_customer_link_summary(monthly_customer_link_groups(db))
        all_rows = historical_sales(db, target_month, as_of, g.current_user, as_of_round or None) if as_of else current_sales(db, target_month, g.current_user)
        filtered = filter_rows(all_rows, request.args)
        selected_business_unit = str(request.args.get("business_unit", "")).strip()
        imported_initial = imported_round_rows(db, target_month, round_key="initial")
        filtered_initial = filter_rows(imported_initial, request.args) if imported_initial is not None else None
        summary = summary_payload(
            db, filtered, target_month, as_of or None, selected_business_unit,
            initial_override_rows=filtered_initial,
        )
        # TODAY SALES STATUS must use the same time slice as the rest of the
        # dashboard.  Historical views previously replaced these rows with the
        # current ledger, which made the panel contradict the selected cutoff.
        live_rows = filtered
        live_summary = summary
        compare_summary = None
        if compare_date:
            compare_rows = historical_sales(db, target_month, compare_date, g.current_user)
            compare_rows = filter_rows(compare_rows, request.args)
            compare_summary = summary_payload(
                db, compare_rows, target_month, compare_date, selected_business_unit,
                initial_override_rows=filtered_initial,
            )
        rate_rows = db.execute(
            "SELECT * FROM monthly_sales_rates WHERE target_month = ? ORDER BY currency", (target_month,)
        ).fetchall()
        rates = plan_rates(db, target_month)
        initial_count = db.execute(
            "SELECT COUNT(*) AS c FROM monthly_fcst_snapshots WHERE target_month = ?", (target_month,)
        ).fetchone()["c"]
        initial_count += db.execute(
            "SELECT COUNT(*) AS c FROM monthly_sales_round_imports WHERE target_month=? AND round_key='initial'",
            (target_month,),
        ).fetchone()["c"]
        today = datetime.now(SEOUL).date()
        next_workspace_open = today.day >= 20 or target_month <= f"{today.year:04d}-{today.month:02d}"
        sheet_state = sheet_sync_state(db)
        if g.current_user["role"] not in {"admin", "manager"}:
            sheet_state.pop("service_account_email", None)
            sheet_state.pop("spreadsheet_url", None)
        setting = month_setting(db, target_month, now)
        round_cards = round_cards_payload(
            db, target_month, g.current_user, request.args, setting, summary.get("total", {})
        )
        return jsonify(
            month=target_month,
            as_of=as_of or None,
            as_of_basis="recorded_at_end_of_day" if as_of else "current_recorded_state",
            read_only=bool(as_of) or g.current_user["role"] == "viewer",
            rows=filtered,
            row_count=len(filtered),
            total_row_count=len(all_rows),
            summary=summary,
            live_summary=live_summary,
            live_status=live_status_payload(live_rows),
            live_as_of=as_of or datetime.now(SEOUL).date().isoformat(),
            compare_summary=compare_summary,
            filters=filter_options(db, all_rows, g.current_user),
            business_units=BUSINESS_UNITS,
            statuses=SALES_STATUSES,
            timing_types=TIMING_TYPES,
            customer_histories=CUSTOMER_HISTORIES,
            owner_options=sales_owner_options(db),
            customer_link_summary=customer_link_summary,
            plan_rate=next((row_dict(row) for row in rate_rows if row["currency"] == "USD"), None),
            plan_rates=rates,
            round_cards=round_cards,
            historical_round_imports=round_imports_payload(db, target_month),
            initial_fcst_submitted=initial_count > 0,
            month_setting=setting,
            sheet_sync=sheet_state,
            next_month_workspace_open=next_workspace_open,
            permissions={
                "can_edit": not as_of and g.current_user["role"] in {"admin", "manager", "editor"},
                "can_configure": not as_of and g.current_user["role"] in {"admin", "manager", "editor"},
                "can_manage": not as_of and g.current_user["role"] in {"admin", "manager"},
                "can_admin": not as_of and g.current_user["role"] == "admin",
            },
        )

    @app.post("/api/monthly-sales-fcst/sales")
    @role_required("admin", "manager", "editor")
    @csrf_required
    def create_monthly_sale():
        data = request.get_json(silent=True) or {}
        db = get_db()
        try:
            db.execute("BEGIN IMMEDIATE")
            values = normalize_sale_input(data, db, now, actor=g.current_user)
            assert_month_open(db, values["target_month"], now)
            commercial_context, commercial_source_type = requested_commercial_context(db, data, values)
            duplicate = duplicate_pi(db, values["milestones"].get("pi_no"))
            if duplicate:
                raise ValueError(f"동일 PI 번호가 {duplicate['sales_no']}에 이미 등록되어 있습니다.")
            sale_id = uuid.uuid4().hex
            timestamp = now()
            sales_no = next_sales_no(db, values["target_month"])
            if not values.get("customer_master_id"):
                raise ValueError("신규 매출은 거래처 마스터에서 거래처를 선택하거나 빠른 등록해 주세요.")
            db.execute(
                """
                INSERT INTO monthly_sales
                  (id, sales_no, customer_name, country, owner_name, business_unit, currency,
                   transaction_currency, transaction_currency_standard, transaction_amount, customer_history,
                   original_month, target_month, timing_type, amount_usd, plan_rate, plan_usd_rate, actual_rate,
                   applied_rate, rate_type, rate_base_date, rate_source, krw_amount, sales_status,
                   confirmation_basis,confirmation_note,preliminary_actual_currency,
                   preliminary_actual_amount,preliminary_actual_krw_amount,preliminary_reference_no,preliminary_note,
                   record_status, split_role, carryover_role, carryover_decision, carryover_decided_at,
                   carryover_target_month, carryover_decision_reason, notes, version, created_by, updated_by,
                   owner_user_id,created_at, updated_at)
                VALUES (:id,:sales_no,:customer_name,:country,:owner_name,:business_unit,'USD',
                        :transaction_currency_storage,:transaction_currency,:transaction_amount,:customer_history,
                        :original_month,:target_month,:timing_type,:amount_usd,:plan_rate,:plan_usd_rate,:actual_rate,
                        :applied_rate,:rate_type,:rate_base_date,:rate_source,:krw_amount,:sales_status,
                        :confirmation_basis,:confirmation_note,:preliminary_actual_currency,
                        :preliminary_actual_amount,:preliminary_actual_krw_amount,:preliminary_reference_no,:preliminary_note,
                        'active','none','none',:carryover_decision,:carryover_decided_at,
                        :carryover_target_month,:carryover_decision_reason,:notes,1,:created_by,:updated_by,
                        :owner_user_id,:created_at,:updated_at)
                """,
                {
                    **values, "id": sale_id, "sales_no": sales_no,
                    "created_by": g.current_user["id"], "updated_by": g.current_user["id"],
                    "created_at": timestamp, "updated_at": timestamp,
                },
            )
            apply_sale_customer_link(db, sale_id, values, g.current_user["id"], timestamp)
            if commercial_context is not None:
                apply_monthly_sale_context(
                    db, sale_id, commercial_context, source_type=commercial_source_type
                )
            upsert_milestones(db, sale_id, values["milestones"], timestamp)
            touch_candidate(db, values, timestamp)
            add_history(
                db, sale_id, "CREATE", "매출 건 등록", g.current_user, now,
                effective_at=values["effective_at"], effective_at_source=values["effective_at_source"],
            )
            audit("MONTHLY_SALE_CREATE", "monthly_sale", sale_id, f"{sales_no} 월별 매출 등록", None, sale_snapshot(db, sale_id), connection=db)
            sheet = maybe_sync(db)
        except (ValueError, PermissionError, sqlite3.Error) as exc:
            rollback_quietly(db)
            return operation_error(exc)
        return jsonify(message="월별 매출을 등록했습니다.", sale=sale_snapshot(db, sale_id), sheet_sync=sheet,
                       status_adjusted_to=values["status_adjusted_to"], status_suggestion=values["status_suggestion"]), 201

    @app.put("/api/monthly-sales-fcst/sales/<sale_id>")
    @role_required("admin", "manager", "editor")
    @csrf_required
    def update_monthly_sale(sale_id):
        data = request.get_json(silent=True) or {}
        db = get_db()
        try:
            db.execute("BEGIN IMMEDIATE")
            existing = db.execute("SELECT * FROM monthly_sales WHERE id = ?", (sale_id,)).fetchone()
            if not existing:
                rollback_quietly(db)
                return jsonify(error="수정할 매출 건을 찾을 수 없습니다."), 404
            if existing["record_status"] != "active":
                raise PermissionError("취소 또는 전액차월된 매출은 수정할 수 없습니다.")
            if not can_edit_sale(g.current_user, existing):
                rollback_quietly(db)
                return jsonify(error="본인 담당 매출만 수정할 수 있습니다."), 403
            assert_month_open(db, existing["target_month"], now)
            requested_version = int(data.get("version") or 0)
            if requested_version != existing["version"]:
                rollback_quietly(db)
                return jsonify(error="다른 사용자가 먼저 수정했습니다. 최신 자료를 다시 불러와 주세요.", code="VERSION_CONFLICT"), 409
            values = normalize_sale_input(data, db, now, existing, g.current_user)
            if values["target_month"] != existing["target_month"]:
                raise ValueError("대상월 변경은 차월 기능을 이용하세요.")
            duplicate = duplicate_pi(db, values["milestones"].get("pi_no"), sale_id)
            if duplicate:
                raise ValueError(f"동일 PI 번호가 {duplicate['sales_no']}에 이미 등록되어 있습니다.")
            commercial_context, commercial_source_type = requested_commercial_context(db, data, values, existing)
            before = sale_snapshot(db, sale_id)
            timestamp = now()
            db.execute(
                """
                UPDATE monthly_sales SET customer_name=?, country=?, owner_name=?, business_unit=?,
                  owner_user_id=?,customer_history=?, transaction_currency=?,transaction_currency_standard=?,transaction_amount=?,
                  original_month=?, timing_type=?, amount_usd=?, actual_rate=?, applied_rate=?,
                  rate_type=?, rate_base_date=?, rate_source=?, krw_amount=?, sales_status=?,
                  confirmation_basis=?,confirmation_note=?,preliminary_actual_currency=?,
                  preliminary_actual_amount=?,preliminary_actual_krw_amount=?,preliminary_reference_no=?,preliminary_note=?,
                  carryover_decision=?, carryover_decided_at=?, carryover_target_month=?,
                  carryover_decision_reason=?, notes=?,
                  version=version+1, updated_by=?, updated_at=? WHERE id=? AND version=?
                """,
                (values["customer_name"], values["country"], values["owner_name"], values["business_unit"],
                 values["owner_user_id"],values["customer_history"], values["transaction_currency_storage"],
                 values["transaction_currency"],values["transaction_amount"],
                 values["original_month"], values["timing_type"], values["amount_usd"],
                 values["actual_rate"], values["applied_rate"], values["rate_type"], values["rate_base_date"],
                 values["rate_source"], values["krw_amount"], values["sales_status"],values["confirmation_basis"],
                 values["confirmation_note"],values["preliminary_actual_currency"],values["preliminary_actual_amount"],
                 values["preliminary_actual_krw_amount"],values["preliminary_reference_no"],values["preliminary_note"],
                 values["carryover_decision"], values["carryover_decided_at"],
                 values["carryover_target_month"], values["carryover_decision_reason"], values["notes"],
                 g.current_user["id"], timestamp, sale_id, existing["version"]),
            )
            if db.execute("SELECT changes() AS c").fetchone()["c"] != 1:
                raise PermissionError("다른 사용자가 먼저 수정했습니다. 최신 자료를 다시 불러와 주세요.")
            apply_sale_customer_link(db, sale_id, values, g.current_user["id"], timestamp)
            if commercial_context is not None:
                apply_monthly_sale_context(
                    db, sale_id, commercial_context, source_type=commercial_source_type
                )
            upsert_milestones(db, sale_id, values["milestones"], timestamp)
            touch_candidate(db, values, timestamp)
            reason = clean_text(data.get("change_reason"), "변경사유", 500)
            add_history(
                db, sale_id, "UPDATE", reason or "매출 건 수정", g.current_user, now,
                effective_at=values["effective_at"], effective_at_source=values["effective_at_source"],
            )
            audit("MONTHLY_SALE_UPDATE", "monthly_sale", sale_id, f"{existing['sales_no']} 월별 매출 수정", before, sale_snapshot(db, sale_id), connection=db)
            sheet = maybe_sync(db)
        except (ValueError, PermissionError, sqlite3.Error) as exc:
            rollback_quietly(db)
            return operation_error(exc)
        return jsonify(message="월별 매출을 수정했습니다.", sale=sale_snapshot(db, sale_id), sheet_sync=sheet,
                       status_adjusted_to=values["status_adjusted_to"], status_suggestion=values["status_suggestion"])

    @app.post("/api/monthly-sales-fcst/sales/<sale_id>/replan")
    @role_required("admin", "manager", "editor")
    @csrf_required
    def replan_monthly_sale_rate(sale_id):
        data = request.get_json(silent=True) or {}
        reason = clean_text(data.get("reason"), "공식 재계획 사유", 500, True)
        effective_at = valid_date(data.get("effective_at"))
        if not effective_at:
            return jsonify(error="실제 재계획 적용일을 입력하세요."), 400
        db = get_db()
        try:
            db.execute("BEGIN IMMEDIATE")
            sale = db.execute("SELECT * FROM monthly_sales WHERE id=?", (sale_id,)).fetchone()
            if not sale:
                rollback_quietly(db)
                return jsonify(error="재계획할 매출 건을 찾을 수 없습니다."), 404
            if sale["record_status"] != "active":
                raise PermissionError("취소 또는 차월 완료된 매출은 재계획할 수 없습니다.")
            if not can_edit_sale(g.current_user, sale):
                raise PermissionError("이 매출을 재계획할 권한이 없습니다.")
            requested_version = int(data.get("version") or 0)
            if requested_version != sale["version"]:
                rollback_quietly(db)
                return jsonify(error="다른 사용자가 먼저 수정했습니다. 최신 자료를 다시 불러와 주세요.", code="VERSION_CONFLICT"), 409
            assert_month_open(db, sale["target_month"], now)
            currency = canonical_sale_currency(sale)
            rates = plan_rates(db, sale["target_month"])
            amount_usd, krw_amount, selected_rate = convert_transaction_amount(
                sale["transaction_amount"], currency, rates
            )
            usd_rate = float(rates.get("USD") or 0)
            before = sale_snapshot(db, sale_id)
            timestamp = now()
            cursor = db.execute(
                """UPDATE monthly_sales
                   SET amount_usd=?,krw_amount=?,plan_rate=?,plan_usd_rate=?,applied_rate=?,
                       version=version+1,updated_by=?,updated_at=?
                   WHERE id=? AND version=?""",
                (amount_usd, krw_amount, selected_rate, usd_rate, selected_rate,
                 g.current_user["id"], timestamp, sale_id, sale["version"]),
            )
            if cursor.rowcount != 1:
                raise PermissionError("다른 사용자가 먼저 수정했습니다. 최신 자료를 다시 불러와 주세요.")
            add_history(
                db, sale_id, "PLAN_RATE_REPLAN", reason, g.current_user, now,
                effective_at=effective_at, effective_at_source="user_business_date",
            )
            audit(
                "MONTHLY_SALE_REPLAN", "monthly_sale", sale_id,
                f"{sale['sales_no']} 계획환율 공식 재적용", before, sale_snapshot(db, sale_id), connection=db,
            )
            sheet = maybe_sync(db)
        except (ValueError, PermissionError, sqlite3.Error) as exc:
            rollback_quietly(db)
            return operation_error(exc)
        return jsonify(message="현재 월 기본환율을 이 매출의 새 계획 Snapshot으로 적용했습니다.",
                       sale=sale_snapshot(db, sale_id), sheet_sync=sheet)

    @app.post("/api/monthly-sales-fcst/sales/<sale_id>/cancel")
    @role_required("admin", "manager", "editor")
    @csrf_required
    def cancel_monthly_sale(sale_id):
        data = request.get_json(silent=True) or {}
        reason = clean_text(data.get("reason"), "취소사유", 500, True)
        db = get_db()
        try:
            db.execute("BEGIN IMMEDIATE")
            sale = db.execute("SELECT * FROM monthly_sales WHERE id = ?", (sale_id,)).fetchone()
            if not sale:
                rollback_quietly(db)
                return jsonify(error="취소할 매출 건을 찾을 수 없습니다."), 404
            if not can_edit_sale(g.current_user, sale):
                rollback_quietly(db)
                return jsonify(error="본인 담당 매출만 취소할 수 있습니다."), 403
            assert_month_open(db, sale["target_month"], now)
            if sale["record_status"] != "active":
                raise ValueError("이미 취소 또는 차월 처리된 매출입니다.")
            before = sale_snapshot(db, sale_id)
            timestamp = now()
            db.execute(
                "UPDATE monthly_sales SET record_status='cancelled', cancel_reason=?, cancelled_at=?, version=version+1, updated_by=?, updated_at=? WHERE id=?",
                (reason, timestamp, g.current_user["id"], timestamp, sale_id),
            )
            add_history(db, sale_id, "CANCEL", reason, g.current_user, now)
            audit("MONTHLY_SALE_CANCEL", "monthly_sale", sale_id, f"{sale['sales_no']} 취소", before, sale_snapshot(db, sale_id), connection=db)
            sheet = maybe_sync(db)
        except (ValueError, PermissionError, sqlite3.Error) as exc:
            rollback_quietly(db)
            return operation_error(exc)
        return jsonify(message="매출 건을 삭제하지 않고 취소 이력으로 보존했습니다.", sheet_sync=sheet)

    @app.post("/api/monthly-sales-fcst/sales/<sale_id>/restore")
    @role_required("admin", "manager", "editor")
    @csrf_required
    def restore_monthly_sale(sale_id):
        data = request.get_json(silent=True) or {}
        reason = clean_text(data.get("reason"), "복구사유", 500, True)
        db = get_db()
        try:
            db.execute("BEGIN IMMEDIATE")
            sale = db.execute("SELECT * FROM monthly_sales WHERE id = ?", (sale_id,)).fetchone()
            if not sale:
                rollback_quietly(db)
                return jsonify(error="복구할 매출 건을 찾을 수 없습니다."), 404
            if not can_edit_sale(g.current_user, sale):
                rollback_quietly(db)
                return jsonify(error="매출 건을 복구할 권한이 없습니다."), 403
            assert_month_open(db, sale["target_month"], now)
            if sale["record_status"] != "cancelled":
                raise ValueError("취소 상태인 매출만 복구할 수 있습니다.")
            latest_history = db.execute(
                "SELECT event_type FROM monthly_sales_history WHERE sales_id=? ORDER BY id DESC LIMIT 1",
                (sale_id,),
            ).fetchone()
            if not latest_history or latest_history["event_type"] != "CANCEL":
                raise PermissionError("매출 취소 메뉴로 취소한 건만 복구할 수 있습니다. 분리·차월 취소 건은 원래 이력에서 관리하세요.")
            before = sale_snapshot(db, sale_id)
            timestamp = now()
            db.execute(
                """UPDATE monthly_sales SET record_status='active', cancel_reason='', cancelled_at=NULL,
                   version=version+1, updated_by=?, updated_at=? WHERE id=?""",
                (g.current_user["id"], timestamp, sale_id),
            )
            add_history(db, sale_id, "RESTORE", reason, g.current_user, now)
            audit(
                "MONTHLY_SALE_RESTORE", "monthly_sale", sale_id,
                f"{sale['sales_no']} 취소 매출 복구", before, sale_snapshot(db, sale_id), connection=db,
            )
            sheet = maybe_sync(db)
        except (ValueError, PermissionError, sqlite3.Error) as exc:
            rollback_quietly(db)
            return operation_error(exc)
        return jsonify(message="취소한 매출을 원래 진행상태로 복구했습니다.", sheet_sync=sheet)

    @app.get("/api/monthly-sales-fcst/sales/<sale_id>/history")
    @role_required("admin", "manager", "editor", "viewer")
    def monthly_sale_history(sale_id):
        db = get_db()
        sale = db.execute("SELECT * FROM monthly_sales WHERE id = ?", (sale_id,)).fetchone()
        if not sale:
            return jsonify(error="매출 건을 찾을 수 없습니다."), 404
        rows = [row_dict(row) for row in db.execute(
            """SELECT id,sales_id,sales_no,event_type,reason,actor_name,occurred_at,
                      effective_at,effective_at_source,change_set_json,snapshot_json
               FROM monthly_sales_history WHERE sales_id=? ORDER BY id DESC""",
            (sale_id,),
        ).fetchall()]
        for row in rows:
            row["snapshot"] = json.loads(row.pop("snapshot_json"))
            try:
                row["change_set"] = json.loads(row.pop("change_set_json") or "{}")
            except (TypeError, json.JSONDecodeError):
                row["change_set"] = {}
        return jsonify(sale=sale_snapshot(db, sale_id), history=rows)

    @app.post("/api/monthly-sales-fcst/sales/<sale_id>/split")
    @role_required("admin", "manager", "editor")
    @csrf_required
    def split_monthly_sale(sale_id):
        data = request.get_json(silent=True) or {}
        operation_key = clean_text(data.get("operation_key"), "작업번호", 100, True)
        db = get_db()
        duplicate_operation = db.execute(
            "SELECT child_sales_id FROM monthly_sales_split_links WHERE operation_key = ?", (operation_key,)
        ).fetchone()
        if duplicate_operation:
            return jsonify(message="이미 처리된 매출분리입니다.", sale=sale_snapshot(db, duplicate_operation["child_sales_id"]), idempotent=True)
        try:
            proceed_amount = number(
                data.get("proceed_amount", data.get("proceed_amount_original")),
                "이번 진행금액", 0.01, True,
            )
            split_date = valid_date(data.get("split_date"))
            reason = clean_text(data.get("reason"), "분리사유", 500, True)
            requested_child_month = clean_text(data.get("target_month"), "분리 매출 대상월", 7)
            if not split_date:
                raise ValueError("분리일을 입력하세요.")
            new_pi_no = clean_text(data.get("new_pi_no"), "신규 PI 번호", 100)
            new_pi_issued = valid_date(data.get("new_pi_issued_at"))
            new_pi_confirmed = valid_date(data.get("new_pi_confirmed_at"))
            db.execute("BEGIN IMMEDIATE")
            original = db.execute("SELECT * FROM monthly_sales WHERE id = ?", (sale_id,)).fetchone()
            if not original:
                rollback_quietly(db)
                return jsonify(error="분리할 매출 건을 찾을 수 없습니다."), 404
            if not can_edit_sale(g.current_user, original):
                rollback_quietly(db)
                return jsonify(error="본인 담당 매출만 분리할 수 있습니다."), 403
            child_month = original["target_month"]
            if requested_child_month and requested_child_month != child_month:
                raise ValueError("매출분리는 원본과 같은 매출월에서만 가능합니다. 다음달 이동은 분리 후 차월 기능을 이용하세요.")
            assert_month_open(db, child_month, now)
            if original["record_status"] != "active":
                raise ValueError("취소 또는 전액차월된 매출은 분리할 수 없습니다.")
            original_transaction_amount = original["transaction_amount"] or original["amount_usd"]
            if proceed_amount >= original_transaction_amount:
                raise ValueError("이번 진행금액은 기존 거래금액보다 작아야 합니다.")
            split_transaction_amount = round(original_transaction_amount - proceed_amount, 4)
            ratio = proceed_amount / original_transaction_amount
            proceed_usd = round(original["amount_usd"] * ratio, 4)
            proceed_krw = round(original["krw_amount"] * ratio, 2)
            split_amount = round(original["amount_usd"] - proceed_usd, 4)
            split_krw = round(original["krw_amount"] - proceed_krw, 2)
            root = root_sales(db, original)
            group_id = original["split_group_id"] or f"SPLIT-{uuid.uuid4().hex[:12].upper()}"
            child_id = uuid.uuid4().hex
            child_no = next_child_no(db, root["sales_no"], "S")
            timestamp = now()
            before = sale_snapshot(db, sale_id)
            db.execute(
                """
                UPDATE monthly_sales SET transaction_amount=?, amount_usd=?, krw_amount=?, split_role='original',
                  split_group_id=?, version=version+1, updated_by=?, updated_at=? WHERE id=?
                """,
                (proceed_amount, proceed_usd, proceed_krw, group_id, g.current_user["id"], timestamp, sale_id),
            )
            child_status = "scheduled" if new_pi_confirmed else "pipeline"
            child_currency = canonical_sale_currency(original)
            db.execute(
                """
                INSERT INTO monthly_sales
                  (id, sales_no, customer_name, country, owner_name, business_unit, currency,
                   transaction_currency, transaction_currency_standard, transaction_amount, customer_history,
                   original_month, target_month, timing_type, amount_usd, plan_rate, plan_usd_rate, actual_rate,
                   applied_rate, rate_type, rate_base_date, rate_source, krw_amount, sales_status,
                   confirmation_basis, confirmation_note, preliminary_actual_currency,
                   preliminary_actual_amount, preliminary_actual_krw_amount,
                   preliminary_reference_no, preliminary_note,
                   record_status, split_role, carryover_role, original_sales_id, split_group_id,
                   notes, version, created_by, updated_by, owner_user_id, created_at, updated_at)
                VALUES
                  (:id,:sales_no,:customer_name,:country,:owner_name,:business_unit,'USD',
                   :transaction_currency_storage,:transaction_currency,:transaction_amount,:customer_history,
                   :original_month,:target_month,:timing_type,:amount_usd,:plan_rate,:plan_usd_rate,NULL,
                   :plan_rate,'plan',NULL,'분리 원본 계획환율 Snapshot',:krw_amount,:sales_status,
                   '','',NULL,NULL,NULL,'','',
                   'active','split','none',:original_sales_id,:split_group_id,
                   :notes,1,:created_by,:updated_by,:owner_user_id,:created_at,:updated_at)
                """,
                {
                    "id": child_id, "sales_no": child_no,
                    "customer_name": original["customer_name"], "country": original["country"],
                    "owner_name": original["owner_name"], "owner_user_id": original["owner_user_id"],
                    "business_unit": original["business_unit"],
                    "transaction_currency_storage": legacy_storage_currency(child_currency),
                    "transaction_currency": child_currency,
                    "transaction_amount": split_transaction_amount,
                    "customer_history": original["customer_history"],
                    "original_month": original["original_month"], "target_month": child_month,
                    "timing_type": original["timing_type"], "amount_usd": split_amount,
                    "plan_rate": original["plan_rate"], "plan_usd_rate": original["plan_usd_rate"],
                    "krw_amount": split_krw, "sales_status": child_status,
                    "original_sales_id": root["id"], "split_group_id": group_id,
                    "notes": original["notes"], "created_by": g.current_user["id"],
                    "updated_by": g.current_user["id"], "created_at": timestamp, "updated_at": timestamp,
                },
            )
            inherit_sale_customer_link(db, child_id, original, g.current_user["id"], timestamp)
            inherit_monthly_sale_context(db, sale_id, child_id)
            source_milestones = milestone_dict(db, sale_id)
            child_milestones = {field: None if field != "exception_reason" else "" for field in MILESTONE_FIELDS}
            child_milestones["order_agreed_at"] = source_milestones.get("order_agreed_at")
            child_milestones["po_received_at"] = source_milestones.get("po_received_at")
            child_milestones["pi_no"] = new_pi_no
            child_milestones["pi_sent_at"] = new_pi_issued
            child_milestones["pi_confirmed_at"] = new_pi_confirmed
            upsert_milestones(db, child_id, child_milestones, timestamp)
            link_id = uuid.uuid4().hex
            db.execute(
                """
                INSERT INTO monthly_sales_split_links
                  (id, operation_key, split_group_id, original_sales_id, child_sales_id,
                   amount_before, proceed_amount, split_amount, split_date, reason, new_pi_no,
                   original_version_after, child_version_after, created_by, created_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 1, ?, ?)
                """,
                (link_id, operation_key, group_id, sale_id, child_id, original["amount_usd"], proceed_usd,
                 split_amount, split_date, reason, new_pi_no, original["version"] + 1, g.current_user["id"], timestamp),
            )
            add_history(db, sale_id, "SPLIT_ORIGINAL", reason, g.current_user, now,
                        effective_at=split_date, effective_at_source="user_business_date")
            add_history(db, child_id, "SPLIT_CHILD", reason, g.current_user, now,
                        effective_at=split_date, effective_at_source="user_business_date")
            audit("MONTHLY_SALE_SPLIT", "monthly_sale", sale_id,
                  f"{original['sales_no']} {original_transaction_amount:,.2f} {original['transaction_currency']} → {proceed_amount:,.2f} + {split_transaction_amount:,.2f}",
                  before, {"original": sale_snapshot(db, sale_id), "child": sale_snapshot(db, child_id)}, connection=db)
            sheet = maybe_sync(db)
        except (ValueError, PermissionError, sqlite3.Error) as exc:
            rollback_quietly(db)
            return operation_error(exc)
        return jsonify(message="매출을 분리했습니다. 분리된 매출은 독립 건으로 관리됩니다.",
                       original=sale_snapshot(db, sale_id), child=sale_snapshot(db, child_id),
                       amount_check=round(proceed_usd + split_amount, 4), sheet_sync=sheet), 201

    @app.post("/api/monthly-sales-fcst/sales/<sale_id>/split-cancel")
    @role_required("admin", "manager")
    @csrf_required
    def cancel_monthly_sale_split(sale_id):
        data = request.get_json(silent=True) or {}
        reason = clean_text(data.get("reason"), "분리 취소사유", 500, True)
        db = get_db()
        try:
            db.execute("BEGIN IMMEDIATE")
            link = db.execute(
                "SELECT * FROM monthly_sales_split_links WHERE (original_sales_id=? OR child_sales_id=?) AND status='active' ORDER BY created_at DESC, rowid DESC LIMIT 1",
                (sale_id, sale_id),
            ).fetchone()
            if not link:
                raise ValueError("취소할 활성 매출분리 이력이 없습니다.")
            original = db.execute("SELECT * FROM monthly_sales WHERE id=?", (link["original_sales_id"],)).fetchone()
            child = db.execute("SELECT * FROM monthly_sales WHERE id=?", (link["child_sales_id"],)).fetchone()
            assert_month_open(db, original["target_month"], now)
            assert_month_open(db, child["target_month"], now)
            if original["version"] != link["original_version_after"] or child["version"] != link["child_version_after"]:
                raise PermissionError("분리 후 수정·차월 등 후속처리가 있어 자동 취소할 수 없습니다.")
            if child["record_status"] != "active":
                raise PermissionError("분리 매출에 후속처리가 있어 취소할 수 없습니다.")
            timestamp = now()
            restored_amount = round(original["amount_usd"] + child["amount_usd"], 4)
            restored_transaction_amount = round(
                (original["transaction_amount"] or original["amount_usd"])
                + (child["transaction_amount"] or child["amount_usd"]), 4
            )
            restored_krw = round(original["krw_amount"] + child["krw_amount"], 2)
            db.execute(
                """UPDATE monthly_sales SET transaction_amount=?, amount_usd=?, krw_amount=?,
                   split_role=CASE WHEN original_sales_id IS NULL THEN 'none' ELSE 'split' END,
                   version=version+1, updated_by=?, updated_at=? WHERE id=?""",
                (restored_transaction_amount, restored_amount, restored_krw,
                 g.current_user["id"], timestamp, original["id"]),
            )
            db.execute(
                "UPDATE monthly_sales SET record_status='cancelled', cancel_reason=?, cancelled_at=?, version=version+1, updated_by=?, updated_at=? WHERE id=?",
                (reason, timestamp, g.current_user["id"], timestamp, child["id"]),
            )
            db.execute(
                "UPDATE monthly_sales_split_links SET status='cancelled', cancelled_by=?, cancelled_at=?, cancel_reason=? WHERE id=?",
                (g.current_user["id"], timestamp, reason, link["id"]),
            )
            add_history(db, original["id"], "SPLIT_CANCEL", reason, g.current_user, now)
            add_history(db, child["id"], "SPLIT_CHILD_CANCEL", reason, g.current_user, now)
            audit("MONTHLY_SALE_SPLIT_CANCEL", "monthly_sale", original["id"], f"{original['sales_no']} 매출분리 취소", None,
                  {"original": sale_snapshot(db, original["id"]), "child": sale_snapshot(db, child["id"])}, connection=db)
            sheet = maybe_sync(db)
        except (ValueError, PermissionError, sqlite3.Error) as exc:
            rollback_quietly(db)
            return operation_error(exc)
        return jsonify(message="매출분리를 취소하고 원래 금액을 복원했습니다.", original=sale_snapshot(db, original["id"]), sheet_sync=sheet)

    @app.post("/api/monthly-sales-fcst/sales/<sale_id>/carryover")
    @role_required("admin", "manager", "editor")
    @csrf_required
    def carryover_monthly_sale(sale_id):
        data = request.get_json(silent=True) or {}
        operation_key = clean_text(data.get("operation_key"), "작업번호", 100, True)
        db = get_db()
        prior = db.execute("SELECT target_sales_id FROM monthly_sales_carryovers WHERE operation_key=?", (operation_key,)).fetchone()
        if prior:
            return jsonify(message="이미 처리된 차월입니다.", sale=sale_snapshot(db, prior["target_sales_id"]), idempotent=True)
        try:
            carry_type = str(data.get("carryover_type", "full")).strip()
            if carry_type != "full":
                raise ValueError("일부 금액은 먼저 매출분리한 뒤 해당 건을 전액 차월해 주세요.")
            target_month = clean_text(data.get("target_month"), "차월 대상월", 7, True)
            carry_date = valid_date(data.get("carryover_date"))
            reason = clean_text(data.get("reason"), "차월사유", 500, True)
            if not valid_month(target_month) or not carry_date:
                raise ValueError("차월 대상월과 차월일을 입력하세요.")
            db.execute("BEGIN IMMEDIATE")
            source = db.execute("SELECT * FROM monthly_sales WHERE id=?", (sale_id,)).fetchone()
            if not source:
                rollback_quietly(db)
                return jsonify(error="차월할 매출 건을 찾을 수 없습니다."), 404
            if not can_edit_sale(g.current_user, source):
                rollback_quietly(db)
                return jsonify(error="본인 담당 매출만 차월할 수 있습니다."), 403
            assert_month_open(db, source["target_month"], now)
            assert_month_open(db, target_month, now)
            if target_month <= source["target_month"]:
                raise ValueError("차월 대상월은 현재 대상월보다 이후여야 합니다.")
            if source["record_status"] != "active":
                raise ValueError("취소 또는 전액차월된 매출은 다시 차월할 수 없습니다.")
            source_milestones = milestone_dict(db, sale_id)
            if source_milestones.get("shipment_actual_at") or source["preliminary_actual_amount"] not in (None, 0):
                raise ValueError("출고 입력이 있는 매출은 잔여금액을 먼저 분리한 뒤 잔여 건만 차월하세요.")
            existing_link = db.execute(
                "SELECT target_sales_id FROM monthly_sales_carryovers WHERE source_sales_id=? AND target_month=? AND status='active'",
                (sale_id, target_month),
            ).fetchone()
            if existing_link:
                raise PermissionError("동일 매출에서 같은 대상월로 생성된 차월 건이 이미 있습니다.")
            amount = source["amount_usd"]
            root = root_sales(db, source)
            target_id = uuid.uuid4().hex
            target_no = next_child_no(db, root["sales_no"], "C")
            timestamp = now()
            before = sale_snapshot(db, sale_id)
            db.execute(
                """UPDATE monthly_sales SET record_status='carried_over', carryover_role='outflow',
                   carryover_decision='yes', carryover_decided_at=COALESCE(carryover_decided_at, ?),
                   carryover_target_month=?,
                   carryover_decision_reason=CASE WHEN carryover_decision_reason='' THEN ? ELSE carryover_decision_reason END,
                   version=version+1, updated_by=?, updated_at=? WHERE id=?""",
                (carry_date, target_month, reason, g.current_user["id"], timestamp, sale_id),
            )
            target_rates = plan_rates(db, target_month)
            target_currency = canonical_sale_currency(source)
            target_amount_usd, target_krw, selected_rate = convert_transaction_amount(
                source["transaction_amount"] or source["amount_usd"],
                target_currency, target_rates,
            )
            target_usd_rate = float(target_rates.get("USD") or 0)
            split_role = source["split_role"]
            db.execute(
                """
                INSERT INTO monthly_sales
                  (id, sales_no, customer_name, country, owner_name, business_unit, currency,
                   transaction_currency, transaction_currency_standard, transaction_amount, customer_history,
                   original_month, target_month, timing_type, amount_usd, plan_rate, plan_usd_rate, actual_rate,
                   applied_rate, rate_type, rate_source, krw_amount, sales_status, record_status,
                   confirmation_basis, confirmation_note, preliminary_actual_currency,
                   preliminary_actual_amount, preliminary_actual_krw_amount,
                   preliminary_reference_no, preliminary_note,
                   split_role, carryover_role, original_sales_id, split_group_id, notes, version,
                   created_by, updated_by, owner_user_id, created_at, updated_at)
                VALUES
                  (:id,:sales_no,:customer_name,:country,:owner_name,:business_unit,'USD',
                   :transaction_currency_storage,:transaction_currency,:transaction_amount,:customer_history,
                   :original_month,:target_month,'previous_carryover',:amount_usd,:plan_rate,:plan_usd_rate,NULL,
                   :plan_rate,'plan','차월월 계획환율 Snapshot',:krw_amount,:sales_status,'active',
                   :confirmation_basis,:confirmation_note,NULL,NULL,NULL,'','',
                   :split_role,'inflow',:original_sales_id,:split_group_id,:notes,1,
                   :created_by,:updated_by,:owner_user_id,:created_at,:updated_at)
                """,
                {
                    "id": target_id, "sales_no": target_no,
                    "customer_name": source["customer_name"], "country": source["country"],
                    "owner_name": source["owner_name"], "owner_user_id": source["owner_user_id"],
                    "business_unit": source["business_unit"],
                    "transaction_currency_storage": legacy_storage_currency(target_currency),
                    "transaction_currency": target_currency,
                    "transaction_amount": source["transaction_amount"],
                    "customer_history": source["customer_history"],
                    "original_month": source["original_month"], "target_month": target_month,
                    "amount_usd": target_amount_usd, "plan_rate": selected_rate,
                    "plan_usd_rate": target_usd_rate, "krw_amount": target_krw,
                    "sales_status": source["sales_status"],
                    "confirmation_basis": source["confirmation_basis"],
                    "confirmation_note": source["confirmation_note"],
                    "split_role": split_role, "original_sales_id": root["id"],
                    "split_group_id": source["split_group_id"], "notes": source["notes"],
                    "created_by": g.current_user["id"], "updated_by": g.current_user["id"],
                    "created_at": timestamp, "updated_at": timestamp,
                },
            )
            inherit_sale_customer_link(db, target_id, source, g.current_user["id"], timestamp)
            inherit_monthly_sale_context(db, sale_id, target_id)
            upsert_milestones(db, target_id, carryover_milestones(db, sale_id), timestamp)
            carry_label = f"{SALES_STATUSES[source['sales_status']]}차월"
            link_id = uuid.uuid4().hex
            db.execute(
                """
                INSERT INTO monthly_sales_carryovers
                  (id, operation_key, source_sales_id, target_sales_id, carryover_type, source_month,
                   target_month, amount_usd, status_at_carryover, carryover_label, carryover_date,
                   reason, source_version_after, target_version_after, created_by, created_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 1, ?, ?)
                """,
                (link_id, operation_key, sale_id, target_id, carry_type, source["target_month"], target_month,
                 amount, source["sales_status"], carry_label, carry_date, reason, source["version"] + 1,
                 g.current_user["id"], timestamp),
            )
            add_history(db, sale_id, "CARRYOVER_OUT", reason, g.current_user, now,
                        effective_at=carry_date, effective_at_source="user_business_date")
            add_history(db, target_id, "CARRYOVER_IN", reason, g.current_user, now,
                        effective_at=carry_date, effective_at_source="user_business_date")
            audit("MONTHLY_SALE_CARRYOVER", "monthly_sale", sale_id,
                  f"{source['sales_no']} {carry_label} {amount:,.2f} USD → {target_month}", before,
                  {"source": sale_snapshot(db, sale_id), "target": sale_snapshot(db, target_id)}, connection=db)
            sheet = maybe_sync(db)
        except (ValueError, PermissionError, sqlite3.Error) as exc:
            rollback_quietly(db)
            return operation_error(exc)
        return jsonify(message=f"{carry_label} 처리를 완료했습니다.", source=sale_snapshot(db, sale_id),
                       target=sale_snapshot(db, target_id), sheet_sync=sheet), 201

    @app.post("/api/monthly-sales-fcst/sales/<sale_id>/carryover-cancel")
    @role_required("admin", "manager")
    @csrf_required
    def cancel_monthly_sale_carryover(sale_id):
        data = request.get_json(silent=True) or {}
        reason = clean_text(data.get("reason"), "차월 취소사유", 500, True)
        effective_input = data.get("effective_at")
        try:
            cancel_effective_at = valid_date(effective_input) if effective_input not in (None, "") else datetime.now(SEOUL).date().isoformat()
        except ValueError as exc:
            return operation_error(exc)
        cancel_effective_source = "user_business_date" if effective_input not in (None, "") else "system_default"
        db = get_db()
        try:
            db.execute("BEGIN IMMEDIATE")
            link = db.execute(
                "SELECT * FROM monthly_sales_carryovers WHERE (source_sales_id=? OR target_sales_id=?) AND status='active' ORDER BY created_at DESC, rowid DESC LIMIT 1",
                (sale_id, sale_id),
            ).fetchone()
            if not link:
                raise ValueError("취소할 활성 차월 이력이 없습니다.")
            source = db.execute("SELECT * FROM monthly_sales WHERE id=?", (link["source_sales_id"],)).fetchone()
            target = db.execute("SELECT * FROM monthly_sales WHERE id=?", (link["target_sales_id"],)).fetchone()
            assert_month_open(db, source["target_month"], now)
            assert_month_open(db, target["target_month"], now)
            if source["version"] != link["source_version_after"] or target["version"] != link["target_version_after"]:
                raise PermissionError("차월 후 수정·분리 등 후속처리가 있어 자동 취소할 수 없습니다.")
            if target["record_status"] != "active":
                raise PermissionError("차월 매출에 후속처리가 있어 취소할 수 없습니다.")
            timestamp = now()
            if link["carryover_type"] == "full":
                db.execute(
                    """UPDATE monthly_sales SET record_status='active', carryover_role='none',
                       carryover_decision='no', carryover_decided_at=NULL, carryover_target_month=NULL,
                       carryover_decision_reason='', version=version+1, updated_by=?, updated_at=? WHERE id=?""",
                    (g.current_user["id"], timestamp, source["id"]),
                )
            else:
                restored = round(source["amount_usd"] + target["amount_usd"], 4)
                restored_transaction = round(
                    (source["transaction_amount"] or source["amount_usd"])
                    + (target["transaction_amount"] or target["amount_usd"]), 4
                )
                restored_krw = round(source["krw_amount"] + target["krw_amount"], 2)
                db.execute(
                    """UPDATE monthly_sales SET transaction_amount=?, amount_usd=?, krw_amount=?, carryover_role='none',
                       carryover_decision='no', carryover_decided_at=NULL, carryover_target_month=NULL,
                       carryover_decision_reason='', version=version+1, updated_by=?, updated_at=? WHERE id=?""",
                    (restored_transaction, restored, restored_krw,
                     g.current_user["id"], timestamp, source["id"]),
                )
            db.execute(
                "UPDATE monthly_sales SET record_status='cancelled', cancel_reason=?, cancelled_at=?, version=version+1, updated_by=?, updated_at=? WHERE id=?",
                (reason, timestamp, g.current_user["id"], timestamp, target["id"]),
            )
            db.execute(
                "UPDATE monthly_sales_carryovers SET status='cancelled', cancelled_by=?, cancelled_at=?, cancel_reason=? WHERE id=?",
                (g.current_user["id"], timestamp, reason, link["id"]),
            )
            add_history(
                db, source["id"], "CARRYOVER_CANCEL", reason, g.current_user, now,
                effective_at=cancel_effective_at, effective_at_source=cancel_effective_source,
            )
            add_history(
                db, target["id"], "CARRYOVER_TARGET_CANCEL", reason, g.current_user, now,
                effective_at=cancel_effective_at, effective_at_source=cancel_effective_source,
            )
            audit("MONTHLY_SALE_CARRYOVER_CANCEL", "monthly_sale", source["id"], f"{source['sales_no']} 차월 취소", None,
                  {"source": sale_snapshot(db, source["id"]), "target": sale_snapshot(db, target["id"])}, connection=db)
            sheet = maybe_sync(db)
        except (ValueError, PermissionError, sqlite3.Error) as exc:
            rollback_quietly(db)
            return operation_error(exc)
        return jsonify(message="차월을 취소하고 원래 매출을 복원했습니다.", source=sale_snapshot(db, source["id"]), sheet_sync=sheet)

    @app.put("/api/monthly-sales-fcst/rates/<target_month>")
    @role_required("admin", "manager", "editor")
    @csrf_required
    def update_monthly_sales_rates(target_month):
        if not valid_month(target_month):
            return jsonify(error="설정월 형식이 올바르지 않습니다."), 400
        data = request.get_json(silent=True) or {}
        db = get_db()
        try:
            db.execute("BEGIN IMMEDIATE")
            assert_month_open(db, target_month, now)
            correction_reason = clean_text(data.get("correction_reason"), "환율 변경사유", 500)
            rates, old_rates = apply_plan_rates(
                db, target_month, data.get("rates") or {}, correction_reason, g.current_user, now
            )
            audit(
                "MONTHLY_SALES_RATES", "monthly_sales_rates", target_month,
                f"{target_month} 월 기준환율 설정", old_rates, rates, connection=db,
            )
            sheet = maybe_sync(db)
        except (ValueError, PermissionError, sqlite3.Error) as exc:
            rollback_quietly(db)
            return operation_error(exc)
        return jsonify(message="월별 기준환율을 저장했습니다.", plan_rates=rates, sheet_sync=sheet)

    @app.put("/api/monthly-sales-fcst/settings/<target_month>")
    @role_required("admin", "manager", "editor")
    @csrf_required
    def update_monthly_sales_settings(target_month):
        if not valid_month(target_month):
            return jsonify(error="설정월 형식이 올바르지 않습니다."), 400
        data = request.get_json(silent=True) or {}
        db = get_db()
        try:
            db.execute("BEGIN IMMEDIATE")
            assert_month_open(db, target_month, now)
            timestamp = now()
            rates = None
            if "rates" in data:
                correction_reason = clean_text(data.get("rate_correction_reason"), "환율 변경사유", 500)
                rates, old_rates = apply_plan_rates(
                    db, target_month, data.get("rates") or {}, correction_reason, g.current_user, now
                )
                audit(
                    "MONTHLY_SALES_RATES", "monthly_sales_rates", target_month,
                    f"{target_month} 월 기준환율 설정", old_rates, rates, connection=db,
                )
            targets = data.get("targets") or {}
            for unit in BUSINESS_UNITS:
                unit_values = targets.get(unit) or {}
                target_usd = number(unit_values.get("target_usd"), f"{BUSINESS_UNITS[unit]} 목표 USD", 0)
                target_krw = number(unit_values.get("target_krw"), f"{BUSINESS_UNITS[unit]} 목표 원화", 0)
                db.execute(
                    """
                    INSERT INTO monthly_sales_targets (target_month, business_unit, target_usd, target_krw, updated_by, updated_at)
                    VALUES (?, ?, ?, ?, ?, ?)
                    ON CONFLICT(target_month, business_unit) DO UPDATE SET target_usd=excluded.target_usd,
                      target_krw=excluded.target_krw, updated_by=excluded.updated_by, updated_at=excluded.updated_at
                    """,
                    (target_month, unit, target_usd, target_krw, g.current_user["id"], timestamp),
                )
            cutoffs = data.get("round_cutoffs") or {}
            normalized_cutoffs = {}
            ordered_dates = []
            for key, _label, column in ROUND_DEFINITIONS:
                value = valid_date(cutoffs.get(key))
                normalized_cutoffs[column] = value
                if value:
                    ordered_dates.append(value)
            imported_cutoffs = {
                row["round_key"]: row["cutoff_date"]
                for row in db.execute(
                    "SELECT round_key,cutoff_date FROM monthly_sales_round_imports WHERE target_month=?",
                    (target_month,),
                ).fetchall()
            }
            for key, label, column in ROUND_DEFINITIONS:
                if key in imported_cutoffs and normalized_cutoffs[column] != imported_cutoffs[key]:
                    raise PermissionError(
                        f"{label} 과거자료가 {imported_cutoffs[key]} 기준으로 이관되어 있어 기준일을 변경할 수 없습니다."
                    )
            if ordered_dates != sorted(ordered_dates):
                raise ValueError("마감기준일은 최초 FCST부터 최종마감까지 날짜 순서대로 입력하세요.")
            db.execute(
                """
                INSERT INTO monthly_sales_month_settings
                  (target_month, month_status, initial_cutoff_date, round1_cutoff_date,
                   round2_cutoff_date, round3_cutoff_date, final_cutoff_date, updated_at)
                VALUES (?, 'open', ?, ?, ?, ?, ?, ?)
                ON CONFLICT(target_month) DO UPDATE SET
                  initial_cutoff_date=excluded.initial_cutoff_date,
                  round1_cutoff_date=excluded.round1_cutoff_date,
                  round2_cutoff_date=excluded.round2_cutoff_date,
                  round3_cutoff_date=excluded.round3_cutoff_date,
                  final_cutoff_date=excluded.final_cutoff_date,
                  updated_at=excluded.updated_at
                """,
                (target_month, normalized_cutoffs["initial_cutoff_date"],
                 normalized_cutoffs["round1_cutoff_date"], normalized_cutoffs["round2_cutoff_date"],
                 normalized_cutoffs["round3_cutoff_date"], normalized_cutoffs["final_cutoff_date"], timestamp),
            )
            audit("MONTHLY_SALES_SETTINGS", "monthly_sales_settings", target_month,
                  f"{target_month} 월별 설정", None,
                  {"rates": rates, "round_cutoffs": cutoffs, "targets": targets}, connection=db)
            sheet = maybe_sync(db)
        except (ValueError, PermissionError, sqlite3.Error) as exc:
            rollback_quietly(db)
            return operation_error(exc)
        return jsonify(message="기준환율·목표·차수 기준일을 저장했습니다.", plan_rates=rates or plan_rates(db, target_month), sheet_sync=sheet)

    @app.post("/api/monthly-sales-fcst/forecast/<target_month>/submit")
    @role_required("admin", "manager")
    @csrf_required
    def submit_initial_fcst(target_month):
        if not valid_month(target_month):
            return jsonify(error="FCST 대상월 형식이 올바르지 않습니다."), 400
        db = get_db()
        try:
            db.execute("BEGIN IMMEDIATE")
            assert_month_open(db, target_month, now)
            existing = db.execute("SELECT COUNT(*) AS c FROM monthly_fcst_snapshots WHERE target_month=?", (target_month,)).fetchone()["c"]
            existing += db.execute(
                "SELECT COUNT(*) AS c FROM monthly_sales_round_imports WHERE target_month=? AND round_key='initial'",
                (target_month,),
            ).fetchone()["c"]
            if existing:
                raise PermissionError("최초 FCST가 이미 제출되어 있습니다. 최초 값은 덮어쓸 수 없습니다.")
            rows = db.execute(
                "SELECT * FROM monthly_sales WHERE target_month=? AND record_status='active'", (target_month,)
            ).fetchall()
            if not rows:
                raise ValueError("제출할 FCST 매출 건이 없습니다.")
            timestamp = now()
            for sale in rows:
                db.execute(
                    """
                    INSERT INTO monthly_fcst_snapshots
                      (id, target_month, sales_id, sales_no, business_unit, sales_status,
                       amount_usd, krw_amount, submitted_by, submitted_at)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (uuid.uuid4().hex, target_month, sale["id"], sale["sales_no"], sale["business_unit"],
                     sale["sales_status"], sale["amount_usd"], sale["krw_amount"], g.current_user["id"], timestamp),
                )
            audit("MONTHLY_FCST_SUBMIT", "monthly_fcst", target_month, f"{target_month} 최초 FCST {len(rows)}건 제출", None,
                  {"count": len(rows), "submitted_at": timestamp}, connection=db)
            sheet = maybe_sync(db)
        except (ValueError, PermissionError, sqlite3.Error) as exc:
            rollback_quietly(db)
            return operation_error(exc)
        return jsonify(message=f"{target_month} 최초 FCST {len(rows)}건을 확정했습니다.", sheet_sync=sheet)

    @app.post("/api/monthly-sales-fcst/months/<target_month>/close")
    @role_required("admin")
    @csrf_required
    def close_monthly_sales_month(target_month):
        if not valid_month(target_month):
            return jsonify(error="마감월 형식이 올바르지 않습니다."), 400
        data = request.get_json(silent=True) or {}
        reason = clean_text(data.get("reason"), "마감 확인내용", 500, True)
        db = get_db()
        timestamp = now()
        db.execute(
            """
            INSERT INTO monthly_sales_month_settings (target_month, month_status, closed_by, closed_at, updated_at)
            VALUES (?, 'closed', ?, ?, ?)
            ON CONFLICT(target_month) DO UPDATE SET month_status='closed', closed_by=excluded.closed_by,
              closed_at=excluded.closed_at, updated_at=excluded.updated_at
            """, (target_month, g.current_user["id"], timestamp, timestamp),
        )
        audit("MONTHLY_SALES_CLOSE", "monthly_sales_month", target_month, f"{target_month} 월 마감: {reason}", None,
              {"month_status": "closed", "reason": reason}, connection=db)
        sheet = maybe_sync(db)
        return jsonify(message=f"{target_month} 자료를 마감했습니다.", sheet_sync=sheet)

    @app.post("/api/monthly-sales-fcst/months/<target_month>/reopen")
    @role_required("admin")
    @csrf_required
    def reopen_monthly_sales_month(target_month):
        if not valid_month(target_month):
            return jsonify(error="재개방월 형식이 올바르지 않습니다."), 400
        data = request.get_json(silent=True) or {}
        reason = clean_text(data.get("reason"), "재개방사유", 500, True)
        db = get_db()
        timestamp = now()
        current = month_setting(db, target_month, now)
        if current["month_status"] != "closed":
            return jsonify(error="마감된 월만 재개방할 수 있습니다."), 409
        db.execute(
            """
            UPDATE monthly_sales_month_settings SET month_status='open', reopened_by=?, reopened_at=?,
              reopen_reason=?, updated_at=? WHERE target_month=?
            """, (g.current_user["id"], timestamp, reason, timestamp, target_month),
        )
        audit("MONTHLY_SALES_REOPEN", "monthly_sales_month", target_month, f"{target_month} 재개방: {reason}", current,
              {"month_status": "open", "reason": reason}, connection=db)
        sheet = maybe_sync(db)
        return jsonify(message=f"{target_month} 자료를 재개방했습니다.", sheet_sync=sheet)

    @app.get("/api/monthly-sales-fcst/sheets/status")
    @role_required("admin", "manager")
    def monthly_sales_sheets_status():
        return jsonify(sheet_sync=sheet_sync_state(get_db()), required_tabs=list(SHEET_SCHEMAS))

    @app.post("/api/monthly-sales-fcst/sheets/initialize")
    @role_required("admin")
    @csrf_required
    def initialize_monthly_sales_sheets():
        db = get_db()
        try:
            mirror = GoogleSheetsMirror()
            schema = mirror.ensure_schema()
            result = sync_sheets(db, now)
        except Exception as exc:
            return jsonify(error=str(exc)), 503
        return jsonify(message="Google Sheets 탭 구성과 전체 동기화를 완료했습니다.", schema=schema, sheet_sync=result)

    @app.post("/api/monthly-sales-fcst/sheets/sync")
    @role_required("admin", "manager")
    @csrf_required
    def sync_monthly_sales_sheets():
        result = sync_sheets(get_db(), now)
        status_code = 200 if result["status"] == "success" else 503
        return jsonify(sheet_sync=result), status_code

    def actual_rate_header(value):
        text = re.sub(r"[\s()\[\]_/.-]+", "", str(value or "")).upper()
        if text in {"환율기준일", "기준일", "고시일", "일자", "DATE"}:
            return "rate_base_date"
        if text in {"USD매매기준율", "미화매매기준율", "매매기준율", "USD", "RATE"}:
            return "rate"
        if text in {"출처", "SOURCE"}:
            return "source"
        if text in {"정정사유", "변경사유", "CORRECTIONREASON"}:
            return "correction_reason"
        return None

    @app.post("/api/monthly-sales-fcst/rates/import")
    @role_required("admin", "manager")
    @csrf_required
    def import_monthly_sales_actual_rates():
        upload = request.files.get("file")
        if not upload or not upload.filename.lower().endswith(".xlsx"):
            return jsonify(error="SMBS 조회자료는 .xlsx 파일로 선택하세요."), 400
        try:
            workbook = load_workbook(upload.stream, data_only=True, read_only=True)
            sheet = workbook[workbook.sheetnames[0]]
            raw_rows = list(sheet.iter_rows(values_only=True))
        except Exception:
            return jsonify(error="환율 Excel 파일을 읽을 수 없습니다."), 400
        header_index = None
        header_roles = None
        for index, row in enumerate(raw_rows[:20]):
            roles = [actual_rate_header(value) for value in row]
            if "rate_base_date" in roles and "rate" in roles:
                header_index, header_roles = index, roles
                break
        if header_index is None:
            return jsonify(error="환율 기준일과 USD 매매기준율 열을 찾을 수 없습니다."), 400
        global_reason = clean_text(request.form.get("correction_reason"), "환율 정정사유", 500)
        parsed = []
        errors = []
        db = get_db()
        for row_number, row in enumerate(raw_rows[header_index + 1:], header_index + 2):
            values = {role: row[col] for col, role in enumerate(header_roles) if role and col < len(row)}
            if not any(value not in (None, "") for value in values.values()):
                continue
            try:
                raw_date = values.get("rate_base_date")
                if isinstance(raw_date, (datetime, date)):
                    raw_date = raw_date.strftime("%Y-%m-%d")
                rate_base_date = valid_date(raw_date)
                if not rate_base_date:
                    raise ValueError("환율 기준일을 YYYY-MM-DD 형식으로 입력하세요.")
                rate = number(values.get("rate"), "USD 매매기준율", 0.0001, True)
                source = clean_text(values.get("source") or "SMBS 조회자료 업로드", "환율 출처", 200, True)
                correction_reason = clean_text(values.get("correction_reason") or global_reason, "환율 정정사유", 500)
                existing = db.execute(
                    "SELECT rate FROM monthly_sales_daily_rates WHERE rate_base_date=? AND currency='USD'", (rate_base_date,)
                ).fetchone()
                if existing and abs(existing["rate"] - rate) > 0.0001 and not correction_reason:
                    raise ValueError("기존 환율을 변경하려면 정정사유를 입력하세요.")
                parsed.append((row_number, rate_base_date, rate, source, correction_reason))
            except ValueError as exc:
                errors.append({"row": row_number, "message": str(exc)})
        if errors:
            return jsonify(error="환율자료 검증 오류를 수정한 뒤 다시 업로드하세요.", errors=errors), 400
        if not parsed:
            return jsonify(error="등록할 환율 행이 없습니다."), 400
        try:
            db.execute("BEGIN IMMEDIATE")
            timestamp = now()
            filename = clean_text(upload.filename, "업로드 파일명", 200)
            for _, rate_base_date, rate, source, correction_reason in parsed:
                db.execute(
                    """INSERT INTO monthly_sales_daily_rates
                       (rate_base_date, currency, rate, source, uploaded_file, correction_reason, updated_by, updated_at)
                       VALUES (?, 'USD', ?, ?, ?, ?, ?, ?)
                       ON CONFLICT(rate_base_date, currency) DO UPDATE SET rate=excluded.rate,
                         source=excluded.source, uploaded_file=excluded.uploaded_file,
                         correction_reason=excluded.correction_reason, updated_by=excluded.updated_by,
                         updated_at=excluded.updated_at""",
                    (rate_base_date, rate, source, filename, correction_reason, g.current_user["id"], timestamp),
                )
            applied_ids = apply_available_actual_rates(db, g.current_user, now)
            audit("MONTHLY_ACTUAL_RATE_IMPORT", "monthly_sales_rate", None,
                  f"SMBS 환율자료 {len(parsed)}건 업로드·대기매출 {len(applied_ids)}건 적용", None,
                  {"rates": len(parsed), "applied_ids": applied_ids, "filename": filename}, connection=db)
            sheet_result = maybe_sync(db)
        except (ValueError, PermissionError, sqlite3.Error) as exc:
            rollback_quietly(db)
            return operation_error(exc)
        return jsonify(
            message=f"환율 {len(parsed)}건을 저장하고 확인대기 매출 {len(applied_ids)}건에 적용했습니다.",
            imported_count=len(parsed), applied_count=len(applied_ids), sheet_sync=sheet_result,
        ), 201

    @app.post("/api/monthly-sales-fcst/rates/retry")
    @role_required("admin", "manager")
    @csrf_required
    def retry_monthly_sales_actual_rates():
        db = get_db()
        try:
            db.execute("BEGIN IMMEDIATE")
            applied_ids = apply_available_actual_rates(db, g.current_user, now)
            audit("MONTHLY_ACTUAL_RATE_RETRY", "monthly_sales_rate", None,
                  f"환율 확인대기 재조회 {len(applied_ids)}건 적용", None, {"applied_ids": applied_ids}, connection=db)
            sheet_result = maybe_sync(db)
        except sqlite3.Error as exc:
            rollback_quietly(db)
            return operation_error(exc)
        return jsonify(message=f"확인대기 매출 {len(applied_ids)}건에 저장된 환율을 적용했습니다.",
                       applied_count=len(applied_ids), sheet_sync=sheet_result)

    def style_excel_header_row(sheet, header_row):
        fill = PatternFill("solid", fgColor="073B4C")
        bottom = Side(style="thin", color="A9C9C4")
        for cell in sheet[header_row]:
            cell.fill = fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(bottom=bottom)
        sheet.row_dimensions[header_row].height = 32

    def style_excel_header(sheet, header_row=1):
        style_excel_header_row(sheet, header_row)
        sheet.freeze_panes = f"A{header_row + 1}"
        sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(sheet.max_column)}{max(header_row, sheet.max_row)}"
        sheet.sheet_view.showGridLines = False

    def style_excel_title(sheet, title, subtitle, filter_text, last_column):
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
        sheet.cell(1, 1, title)
        sheet.cell(1, 1).font = Font(size=18, bold=True, color="FFFFFF")
        sheet.cell(1, 1).fill = PatternFill("solid", fgColor="073B4C")
        sheet.cell(1, 1).alignment = Alignment(vertical="center")
        sheet.row_dimensions[1].height = 34
        for row_index, text in ((2, subtitle), (3, filter_text)):
            sheet.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=last_column)
            sheet.cell(row_index, 1, text)
            sheet.cell(row_index, 1).font = Font(size=9, color="567078")
            sheet.cell(row_index, 1).alignment = Alignment(vertical="center")

    @app.get("/api/monthly-sales-fcst/template")
    @role_required("admin", "manager", "editor")
    def monthly_sales_import_template():
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "월별 매출 FCST 등록"
        headers = [
            "거래처명", "국가", "담당자", "사업분야", "신규/기존", "최초 추진월", "매출 대상월",
            "진행 시점", "환종", "거래금액", "진행상태", "영업 확정근거", "확정근거 메모", "업무 적용일",
            "차월 구분", "차월 결정일", "차월 대상월", "차월 사유", "PI 번호",
            "오더합의일", "PO 수신일", "PI 발송일", "PI 컨펌일", "예상 수금일",
            "실제 수금일", "예상 출고일", "실제 출고일", "예상 선적일", "실제 선적일", "비고",
        ]
        sheet.append(headers)
        style_excel_header(sheet)
        widths = [24, 14, 14, 12, 15, 14, 14, 14, 11, 18, 12, 16, 28, 14, 14, 14, 14, 30, 18, 14, 14, 14, 14, 14, 14, 14, 14, 14, 14, 36]
        for index, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        for validation_range, values in (
            ("C2:C1000", f'"{",".join(MONTHLY_SALES_OWNERS)}"'),
            ("D2:D1000", '"에스테틱,메디컬,덴탈"'),
            ("E2:E1000", '"신규 거래처,기존 거래처"'),
            ("H2:H1000", '"당월 추진,전월 이월"'),
            ("I2:I1000", '"USD,EUR,JPY,CNH,KRW"'),
            ("K2:K1000", '"확정,예정,추진,미정"'),
            ("O2:O1000", '"당월 유지,차월 진행"'),
        ):
            validation = DataValidation(type="list", formula1=values, allow_blank=False)
            validation.error = "목록에 있는 값만 선택하세요."
            validation.errorTitle = "허용되지 않은 값"
            validation.prompt = "목록에서 값을 선택하세요."
            validation.promptTitle = "입력 안내"
            validation.showErrorMessage = True
            validation.showInputMessage = True
            sheet.add_data_validation(validation)
            validation.add(validation_range)
        basis_validation = DataValidation(
            type="list", formula1='"수금완료,L/C Open,승인된 Credit,기타"', allow_blank=True
        )
        basis_validation.error = "목록에 있는 확정근거만 선택하세요."
        basis_validation.errorTitle = "허용되지 않은 값"
        basis_validation.showErrorMessage = True
        sheet.add_data_validation(basis_validation)
        basis_validation.add("L2:L1000")
        guide = workbook.create_sheet("사용안내")
        guide["A1"] = "MedPark 월별 매출 FCST 업로드 안내"
        guide["A1"].font = Font(size=17, bold=True, color="FFFFFF")
        guide["A1"].fill = PatternFill("solid", fgColor="073B4C")
        guide.merge_cells("A1:H1")
        instructions = [
            "1. 첫 번째 탭의 2행부터 매출 건별로 입력합니다.",
            "2. 거래처는 먼저 거래처 마스터에 등록하고, Excel의 거래처명·국가를 마스터와 정확히 맞춥니다.",
            "3. 9월 이전을 포함한 과거월도 업로드할 수 있습니다. 날짜는 YYYY-MM-DD, 월은 Excel 날짜 셀 또는 YYYY-MM 형식으로 입력합니다.",
            "4. 거래금액은 선택한 환종 기준이며 신규 건은 월별 계획환율을 매출건 Snapshot으로 저장합니다.",
            "5. 확정 상태는 확정근거가 필수이며, 업무 적용일은 실제 변경일을 입력합니다(미입력 시 오늘).",
            "6. 차월 진행 선택 시 차월 결정일과 차월 대상월을 함께 입력합니다.",
            "7. 오류가 한 행이라도 있으면 전체 업로드가 취소되며 행별 수정내용이 표시됩니다.",
        ]
        for index, instruction in enumerate(instructions, 3):
            guide.cell(index, 1, instruction)
            guide.merge_cells(start_row=index, start_column=1, end_row=index, end_column=8)
        guide["A10"] = "입력 예시"
        guide["A10"].font = Font(bold=True, color="073B4C")
        sample_owner = g.current_user["display_name"] if g.current_user["display_name"] in MONTHLY_SALES_OWNERS else MONTHLY_SALES_OWNERS[0]
        sample = ["예시거래처", "일본", sample_owner, "덴탈", "신규 거래처", "2026-09", "2026-09",
                  "당월 추진", "EUR", 100000, "추진", "", "", "2026-09-01", "당월 유지", "", "", "", "PI-EXAMPLE", "2026-09-01", "", "", "", "2026-09-20",
                  "", "2026-09-25", "", "", "", "프로모션 협의 중"]
        for col, (header, value) in enumerate(zip(headers, sample), 1):
            guide.cell(11, col, header)
            guide.cell(12, col, value)
        style_excel_header(guide, 11)
        guide.freeze_panes = "A11"
        guide.column_dimensions["A"].width = 26
        for col in range(2, len(headers) + 1):
            guide.column_dimensions[get_column_letter(col)].width = 14
        buffer = io.BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        return send_file(buffer, as_attachment=True, download_name="MedPark_월별매출_FCST_업로드양식.xlsx",
                         mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    @app.post("/api/monthly-sales-fcst/import")
    @role_required("admin", "manager", "editor")
    @csrf_required
    def import_monthly_sales():
        upload = request.files.get("file")
        if not upload or not upload.filename.lower().endswith(".xlsx"):
            return jsonify(error=".xlsx 형식의 Excel 파일을 선택하세요."), 400
        try:
            workbook = load_workbook(upload.stream, data_only=True, read_only=True)
            sheet = workbook[workbook.sheetnames[0]]
        except Exception:
            return jsonify(error="Excel 파일을 읽을 수 없습니다."), 400
        header_alias = {
            "거래처명": "customer_name", "국가": "country", "담당자": "owner_name", "사업분야": "business_unit",
            "신규/기존": "customer_history", "최초 추진월": "original_month", "매출 대상월": "target_month",
            "진행 시점": "timing_type", "매출 추진시점": "timing_type", "환종": "transaction_currency",
            "거래금액": "transaction_amount", "매출액(USD)": "transaction_amount", "진행상태": "sales_status", "PI 번호": "pi_no",
            "영업 확정근거": "confirmation_basis", "확정근거 메모": "confirmation_note", "업무 적용일": "effective_at",
            "차월 구분": "carryover_decision", "차월 여부": "carryover_decision", "차월 결정일": "carryover_decided_at",
            "차월 대상월": "carryover_target_month", "차월 사유": "carryover_decision_reason",
            "오더합의일": "order_agreed_at", "PO 수신일": "po_received_at", "PI 발송일": "pi_sent_at",
            "PI 컨펌일": "pi_confirmed_at", "예상 수금일": "payment_expected_at", "실제 수금일": "payment_actual_at",
            "예상 출고일": "shipment_expected_at", "실제 출고일": "shipment_actual_at", "예상 선적일": "shipping_expected_at",
            "실제 선적일": "shipping_actual_at", "비고": "notes",
        }
        business_map = {"에스테틱": "aesthetic", "메디컬": "medical", "메디칼": "medical", "덴탈": "dental"}
        status_map = {"확정": "confirmed", "예정": "scheduled", "추진": "pipeline", "미정": "undecided"}
        confirmation_map = {
            "수금완료": "payment_complete", "L/C OPEN": "lc_open", "LC OPEN": "lc_open",
            "승인된 CREDIT": "approved_credit", "승인 CREDIT": "approved_credit", "기타": "other",
        }
        timing_map = {"당월 추진": "current_new", "당월 신규": "current_new", "전월 이월": "previous_carryover", "이월 유입": "previous_carryover"}
        customer_history_map = {"신규 거래처": "new", "신규": "new", "기존 거래처": "existing", "기존": "existing"}
        carryover_map = {
            "차월 진행": "yes", "차월": "yes", "예": "yes", "Y": "yes",
            "당월 유지": "no", "차월 아님": "no", "아니오": "no", "N": "no",
        }
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return jsonify(error="Excel 파일에 데이터가 없습니다."), 400
        headers = [str(value or "").strip() for value in rows[0]]
        required = {"거래처명", "국가", "담당자", "사업분야", "최초 추진월", "매출 대상월", "진행상태"}
        missing = sorted(required - set(headers))
        if "거래금액" not in headers and "매출액(USD)" not in headers:
            missing.append("거래금액")
        if missing:
            return jsonify(error=f"필수 열이 없습니다: {', '.join(missing)}"), 400
        parsed = []
        errors = []
        seen_signatures = set()
        seen_pi_numbers = set()
        db = get_db()
        for row_number, values in enumerate(rows[1:], 2):
            source = {header_alias.get(headers[index], headers[index]): value for index, value in enumerate(values) if index < len(headers)}
            if not any(value not in (None, "") for value in values):
                continue
            source["business_unit"] = business_map.get(str(source.get("business_unit") or "").strip(), str(source.get("business_unit") or "").strip())
            source["sales_status"] = status_map.get(str(source.get("sales_status") or "").strip(), str(source.get("sales_status") or "").strip())
            raw_basis = str(source.get("confirmation_basis") or "").strip()
            source["confirmation_basis"] = confirmation_map.get(raw_basis.upper(), raw_basis)
            source["customer_history"] = customer_history_map.get(
                str(source.get("customer_history") or "신규 거래처").strip(),
                str(source.get("customer_history") or "new").strip(),
            )
            source["original_month"] = excel_month_value(source.get("original_month"))
            source["target_month"] = excel_month_value(source.get("target_month"))
            source["timing_type"] = inferred_timing_type(
                timing_map.get(str(source.get("timing_type") or "").strip(), source.get("timing_type")),
                source["original_month"], source["target_month"],
            )
            if "환종" not in headers:
                source["transaction_currency"] = "USD"
            raw_carryover = str(source.get("carryover_decision") or "당월 유지").strip()
            source["carryover_decision"] = carryover_map.get(raw_carryover, raw_carryover.lower())
            if isinstance(source.get("carryover_decided_at"), (datetime, date)):
                source["carryover_decided_at"] = source["carryover_decided_at"].strftime("%Y-%m-%d")
            if isinstance(source.get("carryover_target_month"), (datetime, date)):
                source["carryover_target_month"] = source["carryover_target_month"].strftime("%Y-%m")
            if isinstance(source.get("effective_at"), (datetime, date)):
                source["effective_at"] = source["effective_at"].strftime("%Y-%m-%d")
            milestones = {}
            for field in MILESTONE_FIELDS:
                value = source.pop(field, None)
                if isinstance(value, (datetime, date)):
                    value = value.strftime("%Y-%m-%d")
                milestones[field] = value
            source["milestones"] = milestones
            try:
                owner_name = str(source.get("owner_name") or "").strip()
                if owner_name not in MONTHLY_SALES_OWNERS:
                    raise ValueError(f"담당자는 {', '.join(MONTHLY_SALES_OWNERS)} 중에서 선택하세요.")
                normalized = normalize_sale_input(source, db, now, actor=g.current_user)
                if not normalized.get("customer_master_id"):
                    raise ValueError(
                        "거래처 마스터에 같은 거래처명·국가가 없습니다. 먼저 빠른 등록하거나 마스터 정보를 맞춰 주세요."
                    )
                pi_no = (normalized["milestones"].get("pi_no") or "").strip().lower()
                if pi_no and (pi_no in seen_pi_numbers or duplicate_pi(db, pi_no)):
                    raise ValueError("중복 PI 번호입니다. 기존 행과 업로드 파일을 확인하세요.")
                signature = sale_duplicate_signature(normalized)
                duplicate = duplicate_sale(db, normalized)
                if signature in seen_signatures or duplicate:
                    reference = f" ({duplicate['sales_no']})" if duplicate else ""
                    raise ValueError(f"동일 거래처·국가·담당자·사업분야·대상월·금액의 중복 매출입니다{reference}.")
                if pi_no:
                    seen_pi_numbers.add(pi_no)
                seen_signatures.add(signature)
                parsed.append((row_number, normalized))
            except ValueError as exc:
                errors.append({"row": row_number, "message": str(exc)})
        if errors:
            return jsonify(error="Excel 검증 오류를 수정한 뒤 다시 업로드하세요.", errors=errors, valid_count=len(parsed)), 400
        if not parsed:
            return jsonify(error="등록할 매출 행이 없습니다. 첫 번째 탭의 2행부터 자료를 입력하세요."), 400
        try:
            db.execute("BEGIN IMMEDIATE")
            created_ids = []
            for _, values in parsed:
                assert_month_open(db, values["target_month"], now)
                sale_id = uuid.uuid4().hex
                sales_no = next_sales_no(db, values["target_month"])
                timestamp = now()
                db.execute(
                    """
                    INSERT INTO monthly_sales
                      (id, sales_no, customer_name, country, owner_name, business_unit, currency,
                       transaction_currency, transaction_currency_standard, transaction_amount, customer_history,
                       original_month, target_month, timing_type, amount_usd, plan_rate, plan_usd_rate, actual_rate,
                       applied_rate, rate_type, rate_base_date, rate_source, krw_amount, sales_status,
                       confirmation_basis, confirmation_note, preliminary_actual_currency,
                       preliminary_actual_amount, preliminary_actual_krw_amount, preliminary_reference_no,
                       preliminary_note, record_status, split_role, carryover_role, carryover_decision,
                       carryover_decided_at, carryover_target_month, carryover_decision_reason, notes, version,
                       created_by, updated_by, owner_user_id, created_at, updated_at)
                    VALUES (:id,:sales_no,:customer_name,:country,:owner_name,:business_unit,'USD',
                            :transaction_currency_storage,:transaction_currency,:transaction_amount,:customer_history,
                            :original_month,:target_month,:timing_type,:amount_usd,:plan_rate,:plan_usd_rate,:actual_rate,
                            :applied_rate,:rate_type,:rate_base_date,:rate_source,:krw_amount,:sales_status,
                            :confirmation_basis,:confirmation_note,:preliminary_actual_currency,
                            :preliminary_actual_amount,:preliminary_actual_krw_amount,:preliminary_reference_no,
                            :preliminary_note,'active','none','none',:carryover_decision,
                            :carryover_decided_at,:carryover_target_month,:carryover_decision_reason,:notes,1,
                            :created_by,:updated_by,:owner_user_id,:created_at,:updated_at)
                    """,
                    {
                        **values, "id": sale_id, "sales_no": sales_no,
                        "created_by": g.current_user["id"], "updated_by": g.current_user["id"],
                        "created_at": timestamp, "updated_at": timestamp,
                    },
                )
                apply_sale_customer_link(db, sale_id, values, g.current_user["id"], timestamp)
                upsert_milestones(db, sale_id, values["milestones"], timestamp)
                touch_candidate(db, values, timestamp)
                add_history(
                    db, sale_id, "IMPORT", "Excel 초기자료 업로드", g.current_user, now,
                    effective_at=values["effective_at"], effective_at_source=values["effective_at_source"],
                )
                created_ids.append(sale_id)
            audit("MONTHLY_SALES_IMPORT", "monthly_sale", None, f"월별 매출 Excel {len(created_ids)}건 업로드", None,
                  {"created_ids": created_ids}, connection=db)
            sheet_result = maybe_sync(db)
        except (ValueError, PermissionError, sqlite3.Error) as exc:
            rollback_quietly(db)
            return operation_error(exc)
        return jsonify(message=f"Excel 자료 {len(created_ids)}건을 등록했습니다.", created_count=len(created_ids), sheet_sync=sheet_result), 201

    @app.post("/api/monthly-sales-fcst/history/import")
    @role_required("admin", "manager", "editor")
    @csrf_required
    def import_monthly_sales_round_history():
        upload = request.files.get("file")
        target_month = str(request.form.get("target_month") or "").strip()
        round_key = str(request.form.get("round_key") or "").strip()
        cutoff_date = str(request.form.get("cutoff_date") or "").strip()
        replace_existing = str(request.form.get("replace_existing") or "").strip().lower() in {
            "1", "true", "yes", "on",
        }
        round_definition = next((item for item in ROUND_DEFINITIONS if item[0] == round_key), None)
        if not valid_month(target_month):
            return jsonify(error="대상월 형식이 올바르지 않습니다."), 400
        if not round_definition:
            return jsonify(error="자료 구분을 최초 FCST·1차·2차·3차·최종마감 중에서 선택하세요."), 400
        try:
            valid_date(cutoff_date)
        except ValueError as exc:
            return jsonify(error=str(exc)), 400
        if not upload or not upload.filename.lower().endswith(".xlsx"):
            return jsonify(error=".xlsx 형식의 과거 차수자료를 선택하세요."), 400
        db = get_db()
        existing = db.execute(
            "SELECT * FROM monthly_sales_round_imports WHERE target_month=? AND round_key=?",
            (target_month, round_key),
        ).fetchone()
        if existing and not replace_existing:
            return jsonify(error=f"{target_month} {round_definition[1]} 자료가 이미 이관되어 있습니다. 중복 이관할 수 없습니다."), 409
        try:
            replacement_reason = clean_text(
                request.form.get("replacement_reason"), "교체 사유", 500,
                required=bool(existing and replace_existing),
            )
        except ValueError as exc:
            return jsonify(error=str(exc)), 400
        settings = month_setting(db, target_month, now)
        configured_cutoff = settings.get(round_definition[2])
        if not configured_cutoff:
            return jsonify(error=f"월별 설정에서 {round_definition[1]} 기준일을 먼저 저장하세요."), 400
        if configured_cutoff != cutoff_date:
            return jsonify(error=f"저장된 {round_definition[1]} 기준일은 {configured_cutoff}입니다. 같은 날짜로 이관하세요."), 400
        if round_key == "initial" and not existing:
            submitted = db.execute(
                "SELECT 1 FROM monthly_fcst_snapshots WHERE target_month=? LIMIT 1", (target_month,)
            ).fetchone()
            if submitted:
                return jsonify(error=f"{target_month} 최초 FCST가 이미 제출되어 있어 과거자료를 중복 이관할 수 없습니다."), 409
        try:
            workbook = load_workbook(upload.stream, data_only=True, read_only=True)
            sheet = workbook[workbook.sheetnames[0]]
        except Exception:
            return jsonify(error="과거 차수자료 Excel 파일을 읽을 수 없습니다."), 400
        parsed, errors = parse_round_import_sheet(
            db, sheet, target_month, cutoff_date, round_key, g.current_user, now
        )
        if errors:
            return jsonify(
                error="과거 차수자료 검증 오류를 수정한 뒤 다시 업로드하세요.",
                errors=errors,
                valid_count=len(parsed),
            ), 400
        import_id = uuid.uuid4().hex
        source_file = clean_text(upload.filename, "파일명", 255, True)
        timestamp = now()
        was_replaced = False
        revision_no = 0
        try:
            db.execute("BEGIN IMMEDIATE")
            current_import = db.execute(
                "SELECT * FROM monthly_sales_round_imports WHERE target_month=? AND round_key=?",
                (target_month, round_key),
            ).fetchone()
            if current_import:
                if not replace_existing:
                    raise PermissionError(f"{target_month} {round_definition[1]} 자료가 이미 이관되어 있습니다.")
                if not replacement_reason:
                    raise ValueError("기존 과거자료를 교체하는 사유를 입력하세요.")
                current_rows = db.execute(
                    "SELECT * FROM monthly_sales_round_rows WHERE import_id=? ORDER BY source_row",
                    (current_import["id"],),
                ).fetchall()
                revision_no = db.execute(
                    "SELECT COALESCE(MAX(revision_no), 0) + 1 FROM monthly_sales_round_import_revisions WHERE import_id=?",
                    (current_import["id"],),
                ).fetchone()[0]
                archive = {
                    "import": row_dict(current_import),
                    "rows": [row_dict(row) for row in current_rows],
                }
                db.execute(
                    """INSERT INTO monthly_sales_round_import_revisions
                       (id,import_id,revision_no,source_file,row_count,replaced_by,replaced_at,
                        replacement_reason,snapshot_json)
                       VALUES (?,?,?,?,?,?,?,?,?)""",
                    (
                        uuid.uuid4().hex, current_import["id"], revision_no,
                        current_import["source_file"], current_import["row_count"],
                        g.current_user["id"], timestamp, replacement_reason,
                        json.dumps(archive, ensure_ascii=False, separators=(",", ":"), default=str),
                    ),
                )
                db.execute("DELETE FROM monthly_sales_round_rows WHERE import_id=?", (current_import["id"],))
                import_id = current_import["id"]
                db.execute(
                    """UPDATE monthly_sales_round_imports
                       SET cutoff_date=?,source_file=?,row_count=?,imported_by=?,imported_at=?
                       WHERE id=?""",
                    (cutoff_date, source_file, len(parsed), g.current_user["id"], timestamp, import_id),
                )
                was_replaced = True
            else:
                db.execute(
                    """INSERT INTO monthly_sales_round_imports
                       (id,target_month,round_key,cutoff_date,source_file,row_count,imported_by,imported_at)
                       VALUES (?,?,?,?,?,?,?,?)""",
                    (import_id, target_month, round_key, cutoff_date, source_file, len(parsed),
                     g.current_user["id"], timestamp),
                )
            for item in parsed:
                snapshot = item["snapshot"]
                db.execute(
                    """INSERT INTO monthly_sales_round_rows
                       (id,import_id,source_row,target_month,round_key,cutoff_date,linked_sales_id,
                        sales_no,customer_name,country,owner_name,business_unit,customer_history,
                        original_month,timing_type,transaction_currency,transaction_amount,amount_usd,
                        krw_amount,applied_rate,sales_status,record_status,carryover_decision,
                        carryover_target_month,notes,snapshot_json)
                       VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                    (
                        snapshot["id"], import_id, item["source_row"], target_month, round_key, cutoff_date,
                        item["linked_sales_id"], snapshot["sales_no"], snapshot["customer_name"],
                        snapshot["country"], snapshot["owner_name"], snapshot["business_unit"],
                        snapshot["customer_history"], snapshot["original_month"], snapshot["timing_type"],
                        snapshot["transaction_currency"], snapshot["transaction_amount"], snapshot["amount_usd"],
                        snapshot["krw_amount"], snapshot["applied_rate"], snapshot["sales_status"],
                        snapshot["record_status"], snapshot["carryover_decision"],
                        snapshot["carryover_target_month"], snapshot["notes"],
                        json.dumps(snapshot, ensure_ascii=False, separators=(",", ":"), default=str),
                    ),
                )
            audit(
                "MONTHLY_SALES_ROUND_HISTORY_REPLACE" if was_replaced else "MONTHLY_SALES_ROUND_HISTORY_IMPORT",
                "monthly_sales_round_import", import_id,
                f"{target_month} {round_definition[1]} 과거자료 {len(parsed)}건 "
                f"{'교체' if was_replaced else '이관'}",
                None,
                {"target_month": target_month, "round_key": round_key, "cutoff_date": cutoff_date,
                 "source_file": source_file, "row_count": len(parsed),
                 "replacement_reason": replacement_reason if was_replaced else None,
                 "revision_no": revision_no if was_replaced else None},
                connection=db,
            )
            db.commit()
        except (ValueError, PermissionError, sqlite3.Error) as exc:
            rollback_quietly(db)
            return operation_error(exc)
        linked_count = sum(1 for item in parsed if item["linked_sales_id"])
        return jsonify(
            message=f"{target_month} {round_definition[1]} 과거자료 {len(parsed)}건을 "
                    f"{'교체했습니다. 이전 자료는 이력으로 보존됩니다.' if was_replaced else '이관했습니다.'}",
            import_id=import_id,
            row_count=len(parsed),
            linked_count=linked_count,
            unlinked_count=len(parsed) - linked_count,
            replaced=was_replaced,
            revision_no=revision_no if was_replaced else 0,
        ), 200 if was_replaced else 201

    @app.get("/api/monthly-sales-fcst/export")
    @role_required("admin", "manager", "editor", "viewer")
    def export_monthly_sales():
        target_month = str(request.args.get("month", "")).strip()
        if not valid_month(target_month):
            return jsonify(error="내보낼 조회월을 선택하세요."), 400
        as_of = str(request.args.get("as_of", "")).strip()
        db = get_db()
        rows = historical_sales(db, target_month, as_of, g.current_user) if as_of else current_sales(db, target_month, g.current_user)
        rows = filter_rows(rows, request.args)
        summary = summary_payload(
            db, rows, target_month, as_of or None, str(request.args.get("business_unit", "")).strip()
        )
        generated_at = datetime.now(SEOUL).strftime("%Y-%m-%d %H:%M:%S KST")
        filter_names = {
            "business_unit": "사업분야", "owner": "담당자", "status": "상태", "country": "국가",
            "customer": "거래처", "timing_type": "진행 시점", "customer_history": "신규/기존", "original_month": "최초 추진월",
            "split": "분리", "carryover": "차월", "risk": "위험", "search": "검색",
        }
        applied_filters = [f"{label}={request.args.get(key)}" for key, label in filter_names.items() if request.args.get(key)]
        filter_text = "적용 필터: " + (" · ".join(applied_filters) if applied_filters else "전체")
        if as_of:
            filter_text += f" · 기준일={as_of} 23:59:59 Asia/Seoul"
        if request.args.get("compare_date"):
            filter_text += f" · 비교기준={request.args.get('compare_date')}"
        workbook = Workbook()
        summary_sheet = workbook.active
        summary_sheet.title = "사업분야별 요약"
        summary_headers = [
            "사업분야", "사업계획", "최초 FCST", "현재 FCST", "전월 이월", "당월 추진",
            "확정", "예정", "추진", "미정", "확정+예정", "차월 이관", "목표 대비", "최초 대비", "목표 달성률",
        ]
        summary_sheet.append([])
        summary_sheet.append([])
        summary_sheet.append([])
        summary_sheet.append([])
        summary_sheet.append(summary_headers)
        usd_rows = []
        krw_rows = []
        for key in ("aesthetic", "medical", "dental", "total"):
            item = summary[key]
            achievement = item["achievement_pct"] / 100 if item["achievement_pct"] is not None else None
            usd_rows.append([item["label"], item["target_usd"], item["initial_fcst_usd"], item["current_fcst_usd"],
                             item["carryover_in_usd"], item["new_usd"], item["confirmed_usd"], item["scheduled_usd"],
                             item["pipeline_usd"], item["undecided_usd"], item["high_confidence_usd"],
                             item["carryover_out_usd"], item["target_variance_usd"], item["initial_variance_usd"], achievement])
            krw_rows.append([item["label"], item["target_krw"], item["initial_fcst_krw"], item["current_fcst_krw"],
                             item["carryover_in_krw"], item["new_krw"], item["confirmed_krw"], item["scheduled_krw"],
                             item["pipeline_krw"], item["undecided_krw"], item["confirmed_krw"] + item["scheduled_krw"],
                             item["carryover_out_krw"], item["target_variance_krw"], item["initial_variance_krw"], achievement])
        for values in usd_rows:
            summary_sheet.append(values)
        summary_sheet.merge_cells("A11:O11")
        summary_sheet["A11"] = "원화 환산 요약 (KRW)"
        summary_sheet["A11"].fill = PatternFill("solid", fgColor="DFF1EC")
        summary_sheet["A11"].font = Font(bold=True, color="073B4C")
        summary_sheet["A11"].alignment = Alignment(vertical="center")
        summary_sheet.append(summary_headers)
        for values in krw_rows:
            summary_sheet.append(values)
        style_excel_title(summary_sheet, f"{target_month} 월별 매출 FCST", f"생성: {generated_at} · 생성자: {g.current_user['display_name']}", filter_text, len(summary_headers))
        style_excel_header(summary_sheet, 5)
        summary_sheet.auto_filter.ref = "A5:O9"
        style_excel_header_row(summary_sheet, 12)
        for row_index in range(6, 10):
            for col in range(2, 15):
                summary_sheet.cell(row_index, col).number_format = '$#,##0.00;[Red]($#,##0.00);-'
            summary_sheet.cell(row_index, 15).number_format = "0.0%"
        for row_index in range(13, 17):
            for col in range(2, 15):
                summary_sheet.cell(row_index, col).number_format = '₩#,##0;[Red](₩#,##0);-'
            summary_sheet.cell(row_index, 15).number_format = "0.0%"
        for cell in list(summary_sheet[9]) + list(summary_sheet[16]):
            cell.fill = PatternFill("solid", fgColor="EAF6F2")
            cell.font = Font(bold=True, color="073B4C")
        summary_sheet.column_dimensions["A"].width = 14
        for col in range(2, 16):
            summary_sheet.column_dimensions[get_column_letter(col)].width = 17
        summary_sheet.sheet_view.zoomScale = 90
        summary_sheet.sheet_properties.pageSetUpPr.fitToPage = True
        summary_sheet.page_setup.fitToWidth = 1
        summary_sheet.page_setup.fitToHeight = 0
        summary_sheet.page_setup.orientation = "landscape"
        detail = workbook.create_sheet("매출 세부현황")
        detail_headers = ["매출관리번호", "구분", "사업분야", "거래처명", "국가", "진행 시점", "신규/기존", "담당자",
                          "최초 추진월", "대상월", "환종", "거래금액", "USD 기준", "원화", "환율구분", "적용환율", "환율 기준일",
                          "환율 출처", "상태", "차월 구분", "차월 결정일", "차월 대상월", "차월 사유",
                          "레코드 상태", "PI 번호", "오더합의일", "PO 수신일", "PI 발송일", "PI 컨펌일",
                          "예상수금일", "실제수금일", "예상출고일", "실제출고일", "예상선적일", "실제선적일",
                          "경고", "비고", "수정일"]
        detail.append([])
        detail.append([])
        detail.append([])
        detail.append([])
        detail.append(detail_headers)
        for row in rows:
            badges = ", ".join(value for value in (
                "분리 원본" if row["split_role"] == "original" else "",
                "분리 매출" if row["split_role"] in {"split", "carryover_split"} or row["original_sales_id"] else "",
                "전월 이월" if row["carryover_role"] == "inflow" else "차월 이관" if row["carryover_role"] == "outflow" else "",
            ) if value)
            milestone = row["milestones"]
            detail.append([row["sales_no"], badges, BUSINESS_UNITS[row["business_unit"]], row["customer_name"], row["country"],
                           TIMING_TYPES[row["timing_type"]], CUSTOMER_HISTORIES.get(row.get("customer_history", "existing"), "기존 거래처"),
                           row["owner_name"], row["original_month"], row["target_month"], row.get("transaction_currency", "USD"),
                           row.get("transaction_amount") or row["amount_usd"], row["amount_usd"], row["krw_amount"],
                           row["rate_type"], row["applied_rate"], row["rate_base_date"], row["rate_source"], SALES_STATUSES[row["sales_status"]],
                           "차월 진행" if row.get("carryover_decision") == "yes" else "당월 유지", row.get("carryover_decided_at"),
                           row.get("carryover_target_month"), row.get("carryover_decision_reason"), row["record_status"], milestone.get("pi_no"),
                           milestone.get("order_agreed_at"), milestone.get("po_received_at"), milestone.get("pi_sent_at"), milestone.get("pi_confirmed_at"),
                           milestone.get("payment_expected_at"), milestone.get("payment_actual_at"), milestone.get("shipment_expected_at"),
                           milestone.get("shipment_actual_at"), milestone.get("shipping_expected_at"), milestone.get("shipping_actual_at"),
                           ", ".join(warning["message"] for warning in row.get("warnings", [])), row["notes"], row["updated_at"]])
        style_excel_title(detail, f"{target_month} 매출 세부현황", f"생성: {generated_at} · {len(rows):,}건", filter_text, len(detail_headers))
        style_excel_header(detail, 5)
        usd_col = detail_headers.index("USD 기준") + 1
        transaction_col = detail_headers.index("거래금액") + 1
        krw_col = detail_headers.index("원화") + 1
        rate_col = detail_headers.index("적용환율") + 1
        warning_col = detail_headers.index("경고") + 1
        notes_col = detail_headers.index("비고") + 1
        for row_index in range(6, detail.max_row + 1):
            detail.cell(row_index, usd_col).number_format = '$#,##0.00;[Red]($#,##0.00);-'
            detail.cell(row_index, transaction_col).number_format = '#,##0.00'
            detail.cell(row_index, krw_col).number_format = '₩#,##0;[Red](₩#,##0);-'
            detail.cell(row_index, rate_col).number_format = '#,##0.00'
            detail.cell(row_index, notes_col).alignment = Alignment(wrap_text=True, vertical="top")
        detail.column_dimensions[get_column_letter(detail_headers.index("거래처명") + 1)].width = 28
        detail.column_dimensions[get_column_letter(detail_headers.index("환율 출처") + 1)].width = 32
        detail.column_dimensions[get_column_letter(warning_col)].width = 30
        detail.column_dimensions[get_column_letter(notes_col)].width = 42
        for col in range(1, len(detail_headers) + 1):
            letter = get_column_letter(col)
            if detail.column_dimensions[letter].width == 13.0:
                detail.column_dimensions[letter].width = 15
        detail.sheet_view.zoomScale = 80
        buffer = io.BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        suffix = f"_{as_of}기준" if as_of else ""
        return send_file(buffer, as_attachment=True, download_name=f"MedPark_{target_month}_월별매출_FCST{suffix}.xlsx",
                         mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

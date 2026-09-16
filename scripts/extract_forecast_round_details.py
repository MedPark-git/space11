#!/usr/bin/env python3
"""Extract auditable FCST detail rows from the supplied overseas workbook.

This development-only helper never runs in production.  The generated JSON is
deployed so the application does not require openpyxl or the original workbook.
"""

import argparse
import json
import re
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


BUSINESS_UNITS = {
    "에스테틱": "aesthetic",
    "메디컬": "medical",
    "메디칼": "medical",
    "덴탈": "dental",
}

SECTIONS = (
    ("confirmed", "확정 매출", range(39, 75)),
    ("scheduled", "예정 매출", range(89, 108)),
    ("pipeline", "추진 매출", range(129, 144)),
    ("undecided", "미정", range(160, 164)),
)

ROUND_COLUMNS = {
    1: {"carryover_foreign": "S", "carryover_krw": "T", "current_foreign": "V", "current_krw": "W"},
    2: {"carryover_foreign": "AH", "carryover_krw": "AI", "current_foreign": "AK", "current_krw": "AL"},
    3: {"carryover_foreign": "AN", "carryover_krw": "AO", "current_foreign": "AQ", "current_krw": "AR"},
}


def number(value):
    if isinstance(value, bool) or value in (None, ""):
        return 0.0
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def text(value):
    return str(value).strip() if value not in (None, "") else ""


def iso_date(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    raw = text(value)
    match = re.fullmatch(r"(\d{1,2})/(\d{1,2})", raw)
    if match:
        try:
            return date(2026, int(match.group(1)), int(match.group(2))).isoformat()
        except ValueError:
            return None
    return None


def detail_row(ws, row_number, stage, stage_label, round_no, round_columns):
    business_label = text(ws[f"E{row_number}"].value)
    business_unit = BUSINESS_UNITS.get(business_label)
    if not business_unit:
        return None
    values = {key: number(ws[f"{column}{row_number}"].value) for key, column in round_columns.items()}
    values.update(
        next_foreign=number(ws[f"AT{row_number}"].value),
        next_krw=number(ws[f"AU{row_number}"].value),
        plan_foreign=number(ws[f"P{row_number}"].value),
        plan_krw=number(ws[f"Q{row_number}"].value),
    )
    country_name = text(ws[f"G{row_number}"].value)
    account_name = text(ws[f"H{row_number}"].value)
    notes = text(ws[f"BG{row_number}"].value)
    if not any(values.values()) and not any((country_name, account_name, notes)):
        return None
    return {
        "source_key": f"workbook-row-{row_number}",
        "source_row": row_number,
        "round_no": round_no,
        "sales_stage": stage,
        "sales_stage_label": stage_label,
        "business_unit": business_unit,
        "business_label": business_label,
        "classification": text(ws[f"F{row_number}"].value),
        "country_name": country_name or "국가 미지정",
        "account_name": account_name or "거래처 미지정",
        "item_name": "품목 미지정",
        "owner_name": text(ws[f"I{row_number}"].value),
        "timing_note": text(ws[f"J{row_number}"].value),
        "currency": "USD",
        **values,
        "order_agreed_at": iso_date(ws[f"AW{row_number}"].value),
        "po_received_at": iso_date(ws[f"AX{row_number}"].value),
        "pi_sent_at": iso_date(ws[f"AY{row_number}"].value),
        "payment_expected_at": iso_date(ws[f"AZ{row_number}"].value),
        "payment_completed_at": iso_date(ws[f"BA{row_number}"].value),
        "shipment_expected_at": iso_date(ws[f"BB{row_number}"].value),
        "shipment_completed_at": iso_date(ws[f"BC{row_number}"].value),
        "applied_rate": number(ws[f"BD{row_number}"].value),
        "shipping_completed_at": iso_date(ws[f"BE{row_number}"].value),
        "notes": notes,
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    workbook = load_workbook(args.workbook, data_only=True, read_only=False, keep_links=True)
    sheet_name = next(name for name in workbook.sheetnames if name.strip() == "1. 매출현황")
    ws = workbook[sheet_name]
    rows = []
    for stage, stage_label, row_numbers in SECTIONS:
        for row_number in row_numbers:
            for round_no, columns in ROUND_COLUMNS.items():
                extracted = detail_row(ws, row_number, stage, stage_label, round_no, columns)
                if extracted:
                    rows.append(extracted)
            if stage == "undecided":
                extracted = detail_row(ws, row_number, stage, stage_label, 0, ROUND_COLUMNS[1])
                if extracted:
                    rows.append(extracted)
    workbook.close()
    payload = {
        "version": "2026-08-31-workbook-detail-v1",
        "source_name": args.workbook.name,
        "source_sheet": sheet_name,
        "forecast_month": "2026-08",
        "currency": "KRW",
        "item_note": "원본 엑셀에는 품목명이 없어 '품목 미지정'으로 이관했습니다.",
        "integration_status": "standalone_workbook_only",
        "details": rows,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"output": str(args.output), "details": len(rows)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
